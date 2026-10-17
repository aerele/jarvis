"""export_query and Frappe 16 field masking.

``get_list`` masks a ``mask=1`` field for a user without the "mask" permission, but
only inside its own result handling. A spec export runs a raw query-builder SELECT, so
without its own masking step a user with Read and Export would get the real values in
the file. These tests need Frappe 16 (Frappe 15 has no masking) and skip on 15.
"""

from __future__ import annotations

import io
import unittest
from unittest import mock

import frappe
from frappe.core.doctype.doctype.test_doctype import new_doctype
from frappe.model.meta import Meta
from frappe.tests.utils import FrappeTestCase
from openpyxl import load_workbook

from jarvis.exceptions import InvalidArgumentError
from jarvis.tools._export import masking, resolvers
from jarvis.tools.export_query import export_query

DT = "XQS Mask Test"
ROLE_PLAIN = "XQS Mask Plain"  # read + export, no mask permission
ROLE_MASK = "XQS Mask Allowed"  # read + export + the mask permission
USER_PLAIN = "xqs-mask-plain@example.com"
USER_MASK = "xqs-mask-allowed@example.com"
PHONE = "9876543210"
ROWS = 3


def _spec(**extra) -> dict:
	return {
		"from": DT,
		"alias": "m",
		"select": ["m.public_field", "m.secret_phone", "m.secret_note"],
		"order_by": [{"field": "m.public_field", "dir": "asc"}],
		**extra,
	}


def _content(env: dict):
	content = frappe.get_doc("File", env["name"]).get_content()
	return content if isinstance(content, bytes) else content.encode()


def _cells(env: dict) -> list[tuple]:
	wb = load_workbook(io.BytesIO(_content(env)))
	return list(wb.worksheets[0].iter_rows(values_only=True))


def _field(fieldname: str, **extra) -> dict:
	return {
		"label": fieldname.replace("_", " ").title(),
		"fieldname": fieldname,
		"fieldtype": "Data",
		**extra,
	}


def _ensure_role(role: str) -> None:
	if not frappe.db.exists("Role", role):
		frappe.get_doc({"doctype": "Role", "role_name": role, "desk_access": 1}).insert()


def _ensure_user(user: str, role: str) -> None:
	if not frappe.db.exists("User", user):
		frappe.get_doc(
			{"doctype": "User", "email": user, "first_name": "xqs", "send_welcome_email": 0}
		).insert(ignore_permissions=True)
	frappe.get_doc("User", user).add_roles(role)


class TestExportQueryMasking(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		if not hasattr(Meta, "get_masked_fields"):
			raise unittest.SkipTest("this Frappe has no field masking")
		super().setUpClass()
		frappe.set_user("Administrator")
		for role in (ROLE_PLAIN, ROLE_MASK):
			_ensure_role(role)
		_ensure_user(USER_PLAIN, ROLE_PLAIN)
		_ensure_user(USER_MASK, ROLE_MASK)
		if frappe.db.exists("DocType", DT):
			frappe.delete_doc("DocType", DT, force=True, ignore_permissions=True)
		new_doctype(
			name=DT,
			custom=1,
			fields=[
				_field("public_field"),
				_field("secret_phone", options="Phone", mask=1),
				_field("secret_note", mask=1),
			],
			permissions=[
				{"role": ROLE_PLAIN, "permlevel": 0, "read": 1, "export": 1, "mask": 0},
				{"role": ROLE_MASK, "permlevel": 0, "read": 1, "export": 1, "mask": 1},
			],
		).insert()
		for i in range(ROWS):
			frappe.get_doc(
				{
					"doctype": DT,
					"public_field": f"xqs-mask-{i}",
					"secret_phone": PHONE,
					"secret_note": f"private-{i}",
				}
			).insert()
		frappe.db.commit()
		frappe.clear_cache(doctype=DT)

	@classmethod
	def tearDownClass(cls):
		frappe.set_user("Administrator")
		if frappe.db.exists("DocType", DT):
			frappe.db.delete(DT)
			frappe.delete_doc("DocType", DT, force=True, ignore_permissions=True)
			frappe.db.commit()
		super().tearDownClass()

	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")

	def tearDown(self):
		frappe.set_user("Administrator")
		super().tearDown()

	def test_a_masked_column_is_masked_in_the_file(self):
		frappe.set_user(USER_PLAIN)
		rows = _cells(export_query(spec=_spec()))
		self.assertEqual(len(rows) - 1, ROWS)
		for i, row in enumerate(rows[1:]):
			# The phone keeps its first three digits, the note is fully hidden: the
			# same text get_list gives this user.
			self.assertEqual(row, (f"xqs-mask-{i}", "987XXXXXX", "XXXXXXXX"))

	def test_a_masked_column_is_masked_in_csv_too(self):
		frappe.set_user(USER_PLAIN)
		text = _content(export_query(spec=_spec(), format="csv")).decode()
		self.assertIn("987XXXXXX", text)
		self.assertNotIn(PHONE, text)
		self.assertNotIn("private-", text)

	def test_it_matches_what_get_list_shows(self):
		frappe.set_user(USER_PLAIN)
		listed = frappe.get_list(DT, fields=["secret_phone", "secret_note"], limit=1)[0]
		exported = _cells(export_query(spec=_spec()))[1]
		self.assertEqual((exported[1], exported[2]), (listed["secret_phone"], listed["secret_note"]))

	def test_the_doctype_form_is_masked_the_same_way(self):
		frappe.set_user(USER_PLAIN)
		rows = _cells(export_query(DT, fields=["public_field", "secret_phone"], order_by="public_field asc"))
		self.assertEqual(rows[1], ("xqs-mask-0", "987XXXXXX"))

	def test_background_filter_compiles_sql_then_masks_without_database_reads(self):
		import tempfile
		from pathlib import Path

		from jarvis.tools._export import streaming
		from jarvis.tools._export.database import MariaDBStream, connection_settings

		frappe.set_user(USER_PLAIN)
		prepared = streaming.prepare(
			{"doctype": DT, "fields": ["public_field", "secret_phone"], "order_by": "public_field asc"}, 100
		)
		db = MariaDBStream(connection_settings()).open()
		try:
			with (
				tempfile.TemporaryDirectory() as directory,
				mock.patch.object(frappe.db, "sql", side_effect=AssertionError("STREAM queried DB")),
			):
				path = Path(directory) / "masked.csv"
				streaming.write(
					prepared,
					db.stream_rows(prepared.sql, prepared.params),
					path,
					"csv",
					100,
					1024 * 1024,
					lambda: None,
					lambda n: None,
				)
				text = path.read_text()
				self.assertIn("987XXXXXX", text)
				self.assertNotIn(PHONE, text)
		finally:
			db.close()

	def test_users_who_may_see_the_values_get_them(self):
		for user in ("Administrator", USER_MASK):
			frappe.set_user(user)
			env = export_query(spec=_spec())
			rows = _cells(env)
			self.assertEqual(rows[1], ("xqs-mask-0", PHONE, "private-0"), user)
			# Both users export the same bytes, and save_file reuses a file with the same
			# content: the second save would then hit the first user's PRIVATE File and be
			# refused (File.validate_private_file_access). Drop this user's file first.
			frappe.set_user("Administrator")
			frappe.delete_doc("File", env["name"], ignore_permissions=True, force=True)

	def test_a_computed_column_over_a_masked_field_is_refused(self):
		frappe.set_user(USER_PLAIN)
		shout = {"expr": {"fn": "upper", "args": [{"field": "m.secret_note"}]}, "as": "shout"}
		top = {"agg": "max", "field": "m.secret_phone", "as": "top_phone"}
		for item, field in ((shout, "m.secret_note"), (top, "m.secret_phone")):
			spec = {"from": DT, "alias": "m", "select": ["m.public_field", item]}
			files = frappe.db.count("File")
			with mock.patch.object(resolvers, "_run_bounded") as run:
				with self.assertRaisesRegex(InvalidArgumentError, f"{field}, which is masked for you"):
					export_query(spec=spec)
			# Refused before the SELECT ran, and nothing was written.
			run.assert_not_called()
			self.assertEqual(frappe.db.count("File"), files)

	def test_a_computed_column_over_an_unmasked_field_is_fine(self):
		frappe.set_user(USER_PLAIN)
		spec = {
			"from": DT,
			"alias": "m",
			"select": [
				{"expr": {"fn": "upper", "args": [{"field": "m.public_field"}]}, "as": "shout"},
				{"agg": "count", "field": "m.name", "as": "n"},
			],
			"group_by": ["shout"],
			"order_by": [{"field": "shout", "dir": "asc"}],
		}
		rows = _cells(export_query(spec=spec))
		self.assertEqual(rows[1], ("XQS-MASK-0", 1))

	def test_a_computed_column_over_a_masked_field_is_fine_for_a_user_who_may_see_it(self):
		frappe.set_user(USER_MASK)
		item = {"expr": {"fn": "upper", "args": [{"field": "m.secret_note"}]}, "as": "shout"}
		spec = {
			"from": DT,
			"alias": "m",
			"select": ["m.public_field", item],
			"order_by": [{"field": "m.public_field", "dir": "asc"}],
		}
		self.assertEqual(_cells(export_query(spec=spec))[1], ("xqs-mask-0", "PRIVATE-0"))


class TestMaskingWhereFrappeHasNone(FrappeTestCase):
	def test_it_is_a_no_op_without_the_masking_helper(self):
		# Frappe 15: no mask_field_value to import, so nothing is looked up or masked.
		with mock.patch.object(masking, "mask_field_value", None):
			self.assertEqual(masking.masked_fields({"t": "ToDo"}, {}), {})


class _MetaBeforeParenttype:
	"""Frappe 16 before 2026-07-12: ``get_masked_fields`` takes no argument."""

	def __init__(self):
		self.calls = []

	def get_masked_fields(self):
		self.calls.append("bare")
		return [frappe._dict(fieldname="phone")]


class _MetaWithParenttype:
	"""Frappe 16 from 2026-07-12: a child table is masked by its parent's rules."""

	def __init__(self):
		self.calls = []

	def get_masked_fields(self, parenttype=None):
		self.calls.append(parenttype)
		return [frappe._dict(fieldname="phone")]


class _MetaBehindAKwargsWrapper:
	"""A wrapper that hides the real parameter names behind ``**kwargs``."""

	def __init__(self):
		self.calls = []

	def get_masked_fields(self, **kwargs):
		self.calls.append(kwargs.get("parenttype"))
		return [frappe._dict(fieldname="phone")]


class TestMaskingAcrossFrappe16Builds(FrappeTestCase):
	"""``pyproject.toml`` allows every Frappe 16 build, but ``get_masked_fields`` only
	grew ``parenttype`` in July 2026. Passing it to an older build raised a TypeError,
	so every spec export failed with HTTP 500 there. CI runs the newest Frappe, so a
	stand-in carries the older shape."""

	def _masked(self, meta):
		with (
			mock.patch.object(masking, "mask_field_value", object()),
			mock.patch.object(masking.frappe, "get_meta", return_value=meta),
		):
			return masking.masked_fields(
				{"si": "Sales Invoice", "sii": "Sales Invoice Item"}, {"sii": "Sales Invoice"}
			)

	def test_a_build_without_parenttype_masks_without_raising(self):
		meta = _MetaBeforeParenttype()
		self.assertEqual(sorted(self._masked(meta)), ["si", "sii"])
		self.assertEqual(meta.calls, ["bare", "bare"])

	def test_a_build_with_parenttype_masks_a_child_by_its_parent(self):
		meta = _MetaWithParenttype()
		self.assertEqual(sorted(self._masked(meta)), ["si", "sii"])
		self.assertEqual(meta.calls, [None, "Sales Invoice"])

	def test_a_kwargs_wrapper_still_gets_the_parent(self):
		# Fail closed: a hidden signature must not drop the parent's stricter rule.
		meta = _MetaBehindAKwargsWrapper()
		self.assertEqual(sorted(self._masked(meta)), ["si", "sii"])
		self.assertEqual(meta.calls, [None, "Sales Invoice"])
