"""Validated native conditional formatting shared by both Excel writers."""

import datetime
import math
import unicodedata
from dataclasses import dataclass
from decimal import Decimal
from numbers import Real

from jarvis.exceptions import InvalidArgumentError

_OPERATORS = {"gt": ">", "gte": ">=", "lt": "<", "lte": "<=", "eq": "=", "ne": "<>"}
# Restrained fills paired with dark text; meaning also belongs in column labels
# and values, so color is never the only explanation in the exported workbook.
_STYLES = {
	"danger": ("FCE4D6", "9C0006"),
	"warning": ("FFF2CC", "9C6500"),
	"success": ("E2F0D9", "375623"),
	"info": ("DDEBF7", "1F4E78"),
}


@dataclass(frozen=True)
class ExcelHighlight:
	column: int
	operator: str
	value: str | int | float | Decimal
	style: str
	scope: str


def validate_highlights(highlights, data, *, sheet_name=None):
	"""Resolve exact exported headers and reject unsafe inputs before file save.

	Rules are expressions over existing body rows, not per-cell styles. Numeric
	conditions ignore text/bools/blanks; text conditions use case-sensitive EXACT
	and ignore numbers and empty/ordinary-space-only cells (Excel TRIM semantics).
	Dates are Excel numbers internally, so date targets are deliberately rejected.
	"""
	if highlights is None:
		return []
	if not isinstance(highlights, list) or len(highlights) > 12:
		raise InvalidArgumentError("highlights must be a list of at most 12 rules per sheet")
	if not highlights:
		return []
	if len(data) < 2 or not data[0]:
		raise InvalidArgumentError("A sheet with highlights needs headers and data rows")
	if len(data) > 1048576 or len(data[0]) > 16384:
		raise InvalidArgumentError("Highlight ranges must fit within Excel worksheet limits")
	header = data[0]
	if sheet_name is not None:
		# Match the header actually stored by the active Frappe writer. v16's
		# two import/export template names intentionally preserve HTML.
		from frappe.utils import xlsxutils

		unwrap = not (
			hasattr(xlsxutils, "XLSXStyleBuilder") and sheet_name in {"Data Import Template", "Data Export"}
		)
		header = [
			xlsxutils.ILLEGAL_CHARACTERS_RE.sub("", xlsxutils.handle_html(c) if unwrap else c)
			if isinstance(c, str)
			else c
			for c in header
		]
	result = []
	for spec in highlights:
		if not isinstance(spec, dict) or set(spec) - {"column", "operator", "value", "style", "scope"}:
			raise InvalidArgumentError(
				"Each highlight supports only column, operator, value, style and scope"
			)
		name = spec.get("column")
		if not isinstance(name, str) or not name.strip() or header.count(name) != 1:
			raise InvalidArgumentError("Highlight column must name exactly one header in the exported sheet")
		column = header.index(name)
		operator = spec.get("operator")
		if not isinstance(operator, str) or operator not in _OPERATORS:
			raise InvalidArgumentError("Highlight operator must be gt, gte, lt, lte, eq or ne")
		value = spec.get("value")
		if isinstance(value, str):
			if (
				not value.strip()
				or len(value) > 120
				or any(unicodedata.category(c) in ("Cc", "Cs") or c in "\ufffe\uffff" for c in value)
			):
				raise InvalidArgumentError(
					"Highlight text value must be 1 to 120 characters without controls"
				)
			if operator not in ("eq", "ne"):
				raise InvalidArgumentError("Text highlights support only eq and ne")
		else:
			valid = isinstance(value, (Real, Decimal)) and not isinstance(value, bool)
			try:
				valid = valid and math.isfinite(value)
			except (ValueError, OverflowError):
				valid = False
			if not valid:
				raise InvalidArgumentError("Highlight value must be a finite number or literal text")
		style, scope = spec.get("style", "warning"), spec.get("scope", "cell")
		if not isinstance(style, str) or style not in _STYLES:
			raise InvalidArgumentError("Highlight style must be danger, warning, success or info")
		if scope not in ("cell", "row"):
			raise InvalidArgumentError("Highlight scope must be cell or row")
		if any(
			column < len(row) and isinstance(row[column], (datetime.date, datetime.time)) for row in data[1:]
		):
			raise InvalidArgumentError(
				"Date/time columns are not supported by highlights; use a numeric age or status"
			)
		result.append(ExcelHighlight(column, operator, value, style, scope))
	return result


def _formula_and_range(spec, data):
	"""One formula/range contract for both engines; no caller formulas or refs."""
	from openpyxl.utils import get_column_letter

	column = get_column_letter(spec.column + 1)
	cell = f"${column}2"
	if isinstance(spec.value, str):
		literal = spec.value.replace('"', '""')
		comparison = f'EXACT({cell},"{literal}")'
		if spec.operator == "ne":
			comparison = f"NOT({comparison})"
		formula = f"AND(ISTEXT({cell}),LEN(TRIM({cell}))>0,{comparison})"
	else:
		formula = f'AND(ISNUMBER({cell}),{cell}<>"",{cell}{_OPERATORS[spec.operator]}{spec.value})'
	first, last = ("A", get_column_letter(len(data[0]))) if spec.scope == "row" else (column, column)
	return formula, f"{first}2:{last}{len(data)}"


def add_xlsxwriter_highlights(workbook, worksheet, data, highlights):
	formats = {}
	for spec in highlights:
		if spec.style not in formats:
			fill, font = _STYLES[spec.style]
			formats[spec.style] = workbook.add_format({"bg_color": f"#{fill}", "font_color": f"#{font}"})
		formula, target = _formula_and_range(spec, data)
		worksheet.conditional_format(
			target,
			{"type": "formula", "criteria": formula, "format": formats[spec.style], "stop_if_true": True},
		)


def add_openpyxl_highlights(worksheet, data, highlights):
	from openpyxl.formatting.rule import FormulaRule
	from openpyxl.styles import Font, PatternFill

	for priority, spec in enumerate(highlights, 1):
		formula, target = _formula_and_range(spec, data)
		fill, font = _STYLES[spec.style]
		rule = FormulaRule(
			formula=[formula],
			fill=PatternFill(fill_type="solid", fgColor=fill),
			font=Font(color=font),
			stopIfTrue=True,
		)
		# Priority is worksheet-wide, even when cell and row ranges overlap.
		rule.priority = priority
		worksheet.conditional_formatting.add(target, rule)
