"""export_query(report_name=...): a saved report, every row, straight into a file (#2102)."""

import io
import json
from unittest.mock import patch

import frappe
from frappe.tests.utils import FrappeTestCase

from jarvis.exceptions import ExportTooSlowError, InvalidArgumentError, NoDataError, PermissionDeniedError
from jarvis.tests.test_run_report import (
	_ENQUEUE,
	INLINE_REPORT,
	PREP_REPORT,
	_completed,
	_ensure_report,
	_make_pr,
)
from jarvis.tools import _result_guard
from jarvis.tools._export import reports
from jarvis.tools.export_query import export_query
from jarvis.tools.run_method import _is_denied_path
from jarvis.tools.run_report import run_report

BUILDER_REPORT = "Jarvis Test Builder Report"
COLUMNS = [
	{"fieldname": "name", "label": "ID", "fieldtype": "Link", "options": "ToDo"},
	{"fieldname": "description", "label": "Description", "fieldtype": "Data"},
	{"fieldname": "date", "label": "Date", "fieldtype": "Date"},
	{"fieldname": "secret", "label": "Secret", "fieldtype": "Data", "hidden": 1},
]


def _rows(n):
	return [
		{"name": f"r{i}", "description": f"row {i}", "date": "2026-10-21", "secret": "x"} for i in range(n)
	]


def _file_bytes(env):
	content = frappe.get_doc("File", env["name"]).get_content()
	return content.encode("utf-8") if isinstance(content, str) else content


def _sheet(env):
	"""Header + rows of the saved workbook, as openpyxl reads them."""
	from openpyxl import load_workbook

	ws = load_workbook(io.BytesIO(_file_bytes(env)), read_only=True).active
	return [list(r) for r in ws.iter_rows(values_only=True)]


def _ensure_builder_report():
	if frappe.db.exists("Report", BUILDER_REPORT):
		return
	frappe.get_doc(
		{
			"doctype": "Report",
			"report_name": BUILDER_REPORT,
			"ref_doctype": "ToDo",
			"report_type": "Report Builder",
			"is_standard": "No",
		}
	).insert(ignore_permissions=True)


class TestExportReport(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		_ensure_report(INLINE_REPORT, prepared=0)
		_ensure_report(PREP_REPORT, prepared=1)
		_ensure_builder_report()

	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")
		frappe.db.delete("Prepared Report", {"report_name": PREP_REPORT})
		frappe.db.delete("ToDo", {"description": ["like", "er-%"]})
		for i in range(3):
			frappe.get_doc({"doctype": "ToDo", "description": f"er-{i}"}).insert()

	def tearDown(self):
		frappe.set_user("Administrator")
		super().tearDown()

	# --- the sources --------------------------------------------------------
	def test_inline_report_exports_every_row_frappe_returns(self):
		live = run_report(INLINE_REPORT)
		env = export_query(report_name=INLINE_REPORT, format="xlsx")
		self.assertTrue(env["file_url"].endswith(".xlsx"))
		self.assertEqual(env["total"], len(live["result"]))
		self.assertEqual(env["report_name"], INLINE_REPORT)
		self.assertTrue(env["as_of"])
		sheet = _sheet(env)
		# a Query Report labels its columns by the SQL column name
		self.assertEqual(sheet[0], ["name", "description"])
		self.assertEqual(len(sheet) - 1, len(live["result"]))
		first = live["result"][0]
		self.assertEqual(sheet[1][0], first["name"])
		self.assertEqual(sheet[1][1], first["description"])

	def test_prepared_copy_exports_past_the_model_row_cap(self):
		n = 2500  # over _prepared_reports._MAX_ROWS, which run_report applies
		with _completed(_rows(n), columns=COLUMNS):
			shown = run_report(PREP_REPORT)
			env = export_query(report_name=PREP_REPORT)
		self.assertLess(len(shown["result"]), n)
		self.assertIn("export_query(report_name=", shown["row_note"])
		self.assertEqual(env["total"], n)
		self.assertTrue(env["as_of"])
		sheet = _sheet(env)
		# the hidden helper column is not in the file; a Date string is a real date
		self.assertEqual(sheet[0], ["ID", "Description", "Date"])
		self.assertEqual(len(sheet) - 1, n)
		self.assertEqual(sheet[1][:2], ["r0", "row 0"])
		self.assertEqual(str(sheet[1][2])[:10], "2026-10-21")
		self.assertEqual(sheet[n][:2], [f"r{n - 1}", f"row {n - 1}"])

	def test_prepared_not_generated_returns_started_and_no_file(self):
		files = frappe.db.count("File")
		with patch(_ENQUEUE):
			env = export_query(report_name=PREP_REPORT)
		self.assertTrue(env["prepared_report"])
		self.assertEqual(env["status"], "started")
		self.assertEqual(env["report_name"], PREP_REPORT)
		self.assertEqual(frappe.db.count("File"), files)
		self.assertEqual(frappe.db.count("Prepared Report", {"report_name": PREP_REPORT}), 1)

	def test_prepared_generating_returns_that_status(self):
		_make_pr("Queued")
		with patch(_ENQUEUE):
			env = export_query(report_name=PREP_REPORT)
		self.assertEqual(env["status"], "generating")

	def test_csv(self):
		with _completed(_rows(3), columns=COLUMNS):
			env = export_query(report_name=PREP_REPORT, format="csv")
		text = _file_bytes(env).decode("utf-8")
		self.assertTrue(env["file_url"].endswith(".csv"))
		self.assertEqual(text.splitlines()[0], "ID,Description,Date")
		self.assertEqual(len(text.splitlines()), 4)

	def test_inline_report_gets_the_reports_defaults_filled(self):
		# A Script Report's required filter with a default must reach the script set,
		# as Desk sends it; the prepared path already did this (finding from review).
		live = {"columns": COLUMNS[:2], "result": [{"name": "a", "description": "b"}]}
		with (
			patch.object(reports._prepared_reports, "_report_defaults", return_value={"company": "X"}),
			patch.object(reports, "frappe_run_report", return_value=live) as run,
		):
			export_query(report_name=INLINE_REPORT, filters={"status": "Open"})
		self.assertEqual(run.call_args.kwargs["filters"], {"company": "X", "status": "Open"})

	def test_duplicate_column_names_keep_every_column(self):
		cols = [
			{"fieldname": "amount", "label": "Amount", "fieldtype": "Currency"},
			{"fieldname": "amount", "label": "Amount", "fieldtype": "Currency"},
		]
		with _completed([{"amount": 1}], columns=cols):
			env = export_query(report_name=PREP_REPORT)
		self.assertEqual(_sheet(env)[0], ["Amount", "Amount (2)"])

	def test_only_iso_date_strings_become_dates(self):
		from jarvis.tools import _report_rows

		col = {"fieldtype": "Date"}
		self.assertEqual(str(_report_rows.as_file_value("2026-10-21", col)), "2026-10-21")
		self.assertEqual(_report_rows.as_file_value("21-10-2026", col), "21-10-2026")
		self.assertEqual(_report_rows.as_file_value("n/a", col), "n/a")

	def test_the_scope_resolver_runs_before_the_prepared_lookup(self):
		# AR/AP fill company + report_date before the cache match (run_report does
		# the same), so an export finds the copy run_report generated.
		with (
			patch.object(reports, "resolve_scope", side_effect=lambda r, f: (f, None)) as scope,
			_completed(_rows(2), columns=COLUMNS),
		):
			export_query(report_name=PREP_REPORT, filters={"a": 1})
		scope.assert_called_once()
		self.assertEqual(scope.call_args.args[0], PREP_REPORT)

	# --- refusals -----------------------------------------------------------
	def test_report_builder_refused_with_the_doctype_route(self):
		with self.assertRaises(InvalidArgumentError) as cm:
			export_query(report_name=BUILDER_REPORT)
		self.assertIn("doctype='ToDo'", str(cm.exception))

	def test_over_the_immediate_limit_is_refused_not_queued(self):
		jobs_before = (
			frappe.db.count("Jarvis Export Job") if frappe.db.table_exists("Jarvis Export Job") else 0
		)
		with (
			patch.object(reports, "immediate_row_limit", return_value=10),
			_completed(_rows(11), columns=COLUMNS),
			self.assertRaises(InvalidArgumentError) as cm,
		):
			export_query(report_name=PREP_REPORT)
		self.assertIn("Desk", str(cm.exception))
		self.assertNotIn("export_job", str(cm.exception))
		if frappe.db.table_exists("Jarvis Export Job"):
			self.assertEqual(frappe.db.count("Jarvis Export Job"), jobs_before)

	def test_too_slow_is_refused_not_queued(self):
		with (
			patch("jarvis.tools.export_query._Budget.check", side_effect=ExportTooSlowError("slow")),
			_completed(_rows(2), columns=COLUMNS),
			self.assertRaises(InvalidArgumentError) as cm,
		):
			export_query(report_name=PREP_REPORT)
		self.assertIn("Desk", str(cm.exception))

	def test_no_export_permission_is_refused(self):
		with (
			patch.object(reports.frappe.permissions, "can_export", return_value=False),
			self.assertRaises(PermissionDeniedError) as cm,
		):
			export_query(report_name=INLINE_REPORT)
		self.assertIn("export permission", str(cm.exception))

	def test_missing_required_filter_named_in_one_error(self):
		declared = {
			"filters": [{"fieldname": "company", "reqd": True, "default": None}],
			"required": ["company"],
		}
		with (
			patch.object(reports, "get_report_filters", return_value=declared),
			self.assertRaises(InvalidArgumentError) as cm,
		):
			export_query(report_name=INLINE_REPORT)
		self.assertIn("company", str(cm.exception))

	def test_empty_report_is_no_data(self):
		with _completed([], columns=COLUMNS), self.assertRaises(NoDataError):
			export_query(report_name=PREP_REPORT)

	def test_report_name_excludes_the_other_sources_and_doctype_args(self):
		with self.assertRaises(InvalidArgumentError):
			export_query(doctype="ToDo", report_name=INLINE_REPORT)
		with self.assertRaises(InvalidArgumentError):
			export_query(report_name=INLINE_REPORT, fields=["name"])
		with self.assertRaises(InvalidArgumentError):
			export_query(report_name=INLINE_REPORT, format="csv", highlights=[])

	def test_unknown_report(self):
		with self.assertRaises(InvalidArgumentError):
			export_query(report_name="No Such Report zzz")

	# --- the routes around it -----------------------------------------------
	def test_run_method_denies_frappes_browser_exports(self):
		self.assertTrue(_is_denied_path("frappe.desk.query_report.export_query"))
		self.assertTrue(_is_denied_path("frappe.desk.reportview.export_query"))

	def test_truncated_run_report_points_at_export_query_report_name(self):
		data = {"columns": COLUMNS, "result": [{"description": "x" * 200} for _ in range(5000)]}
		out, event = _result_guard.enforce_result_budget(data, "run_report")
		self.assertEqual(event["kind"], "truncated")
		self.assertIn("report_name", out["note"])
		self.assertLessEqual(_result_guard._size(out), _result_guard.MAX_RESULT_CHARS)

	def test_telemetry_on_success_and_pending(self):
		with patch("jarvis.telemetry.record_export_event") as m, _completed(_rows(2), columns=COLUMNS):
			export_query(report_name=PREP_REPORT)
		self.assertEqual(m.call_args.kwargs.get("outcome"), "ok")
		self.assertEqual(m.call_args.kwargs.get("rows"), 2)
		with patch("jarvis.telemetry.record_export_event") as m, patch(_ENQUEUE):
			export_query(report_name=PREP_REPORT)
		self.assertEqual(m.call_args.kwargs.get("outcome"), "report_pending")

	def test_file_is_private_and_unattached(self):
		with _completed(_rows(2), columns=COLUMNS):
			env = export_query(report_name=PREP_REPORT)
		fdoc = frappe.get_doc("File", env["name"])
		self.assertEqual(fdoc.is_private, 1)
		self.assertFalse(fdoc.attached_to_doctype)
		self.assertEqual(json.loads(json.dumps(env))["total"], 2)
