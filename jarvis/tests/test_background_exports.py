"""Background-export boundary tests, including real files and MariaDB streams."""

import contextlib
import csv
import datetime
import errno
import inspect
import json
import os
import re
import shutil
import tempfile
import threading
import time
from pathlib import Path
from unittest import TestCase
from unittest.mock import MagicMock, patch

import frappe
import pymysql
from frappe.tests.utils import FrappeTestCase

from jarvis._session import impersonate
from jarvis.exceptions import InvalidArgumentError, PermissionDeniedError
from jarvis.tools._export import jobs, masking, resolvers, saved_queries, streaming
from jarvis.tools._export.database import MariaDBStream, Watchdog, connection_settings
from jarvis.tools._export.resolvers import _Column


def prepared(limit=None, rules=None):
	return streaming.Prepared(
		"",
		{},
		[_Column("value", None, None)],
		["Value"],
		{},
		["Data"],
		rules or [],
		"yyyy-mm-dd",
		"hh:mm:ss",
		limit,
	)


# CPU-bound, one-row result: it aggregates INSIDE the derived table the count wraps it in, so it
# never builds a big temp table, and only the statement time limit can end it.
SLOW_COUNT = "SELECT COUNT(*) AS n FROM tabDocField a CROSS JOIN tabDocField b CROSS JOIN tabDocField c"


def spy(owner, name, order):
	"""Patch ``owner.name`` to record each call in ``order`` and then run the real method."""
	original = getattr(owner, name)

	def call(*args, **kwargs):
		order.append(name)
		return original(*args, **kwargs)

	return patch.object(owner, name, autospec=True, side_effect=call)


class TestStreamingWriter(TestCase):
	def test_csv_200000_rows_and_no_database_calls(self):
		with (
			tempfile.TemporaryDirectory() as directory,
			patch.object(frappe.db, "sql", side_effect=AssertionError("STREAM queried DB")),
		):
			path = Path(directory) / "export.part"
			result = streaming.write(
				prepared(),
				({"value": i} for i in range(200_000)),
				path,
				"csv",
				200_000,
				50 * 1024 * 1024,
				lambda: None,
				lambda n: None,
			)
			self.assertEqual(result["rows_written"], 200_000)
			with path.open() as source:
				self.assertEqual(sum(1 for _ in csv.reader(source)), 200_001)

	def test_expected_total_follows_the_rule_write_applies(self):
		for count, limit, cap, expected in (
			(None, None, 10, None),
			(0, None, 10, None),
			(7, None, 10, 7),
			(11, None, 10, 11),
			(6, 5, 10, 5),
			(11, 10, 10, 10),
			(3, 20, 10, 3),
		):
			with self.subTest(count=count, limit=limit):
				self.assertEqual(streaming.expected_total(count, limit, cap), expected)

	def test_finalizing_follows_the_final_progress_and_precedes_close(self):
		import xlsxwriter

		real_close = xlsxwriter.Workbook.close
		for fmt in ("csv", "xlsx"):
			events = []

			def record_close(book, events=events):
				events.append(("close",))
				return real_close(book)

			with (
				self.subTest(fmt=fmt),
				tempfile.TemporaryDirectory() as directory,
				patch.object(xlsxwriter.Workbook, "close", autospec=True, side_effect=record_close),
			):
				result = streaming.write(
					prepared(),
					({"value": i} for i in range(1500)),
					Path(directory) / "export.part",
					fmt,
					10_000,
					10**8,
					lambda: events.append(("check",)),
					lambda n: events.append(("progress", n)),
					finalizing=lambda n: events.append(("finalizing", n)),
				)
				self.assertEqual(result["rows_written"], 1500)
				# A check() legitimately runs after the file is closed; the order of the rest is the rule.
				seen = [e for e in events if e != ("check",)]
				tail = [("progress", 1500), ("finalizing", 1500)] + ([("close",)] if fmt == "xlsx" else [])
				self.assertEqual(seen[-len(tail) :], tail)
				self.assertEqual(seen.count(("finalizing", 1500)), 1)
				after = seen[seen.index(("finalizing", 1500)) :]
				self.assertFalse([e for e in after if e[0] == "progress"])

	def test_finalizing_is_not_reached_by_a_failed_or_empty_export(self):
		for fmt in ("csv", "xlsx"):
			for rows, cap in ((3, 2), (0, 10)):
				events = []
				with self.subTest(fmt=fmt, rows=rows), tempfile.TemporaryDirectory() as directory:
					path = Path(directory) / "export.part"

					def run():
						return streaming.write(
							prepared(),
							iter([{"value": 1}] * rows),
							path,
							fmt,
							cap,
							10**8,
							lambda: None,
							lambda n: events.append(("progress", n)),
							finalizing=lambda n: events.append(("finalizing", n)),
						)

					if rows > cap:
						with self.assertRaises(streaming.ExportLimit):
							run()
					else:
						run()
						self.assertEqual(events, [("progress", 0)])
					self.assertNotIn("finalizing", [e[0] for e in events])

	def test_exact_cap_and_over_cap(self):
		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "export.part"
			with self.assertRaisesRegex(streaming.ExportLimit, "More than 2"):
				streaming.write(
					prepared(), iter([{"value": 1}] * 3), path, "csv", 2, 1000, lambda: None, lambda n: None
				)

	def test_explicit_top_n_and_formula_escape(self):
		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "export.part"
			result = streaming.write(
				prepared(limit=1),
				iter([{"value": "=HYPERLINK(1)"}, {"value": 2}]),
				path,
				"csv",
				2,
				1000,
				lambda: None,
				lambda n: None,
			)
			self.assertEqual(result["limited_by_spec"], 1)
			self.assertIn("'=HYPERLINK", path.read_text())

	def test_file_size_refuses(self):
		with tempfile.TemporaryDirectory() as directory:
			with self.assertRaises(streaming.ExportLimit):
				streaming.write(
					prepared(),
					iter([{"value": "a" * 500}]),
					Path(directory) / "export.part",
					"csv",
					2,
					100,
					lambda: None,
					lambda n: None,
				)

	def test_xlsx_types_sanitization_and_highlights(self):
		from openpyxl import load_workbook

		from jarvis._xlsx_highlights import ExcelHighlight

		p = prepared(rules=[ExcelHighlight(0, "eq", "ok", "success", "cell")])
		with (
			tempfile.TemporaryDirectory() as directory,
			patch.object(frappe.db, "sql", side_effect=AssertionError("STREAM queried DB")),
		):
			path = Path(directory) / "export.part"
			result = streaming.write(
				p,
				iter([{"value": "<p>=1+1</p>"}, {"value": "ok"}]),
				path,
				"xlsx",
				200_000,
				50 * 1024 * 1024,
				lambda: None,
				lambda n: None,
			)
			with path.open("rb") as source:
				workbook = load_workbook(source)
				self.assertTrue(workbook.active["A2"].value.startswith("'"))
				self.assertEqual(workbook.active["A2"].value[1:].strip(), "=1+1")
				self.assertEqual(workbook.active["A2"].data_type, "s")
				self.assertEqual(len(workbook.active.conditional_formatting), 1)
				workbook.close()
			self.assertEqual(result["rows_written"], 2)

	def test_xlsx_dates_from_computed_columns_are_dates(self):
		# A computed column is typed "Data" in prepare; its date values are still dates.
		from openpyxl import load_workbook

		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "export.part"
			streaming.write(
				prepared(),
				iter([{"value": datetime.date(2026, 6, 1)}]),
				path,
				"xlsx",
				10,
				50 * 1024 * 1024,
				lambda: None,
				lambda n: None,
			)
			with path.open("rb") as source:
				cell = load_workbook(source).active["A2"]
		self.assertTrue(cell.is_date)
		self.assertEqual(cell.number_format, "yyyy-mm-dd")

	def test_xlsx_abort_does_not_finalise_the_workbook(self):
		import xlsxwriter

		def too_many():
			streaming.write(
				prepared(), iter([{"value": 1}] * 3), directory / "a.part", "xlsx", 2, 10**8, noop, noop1
			)

		def cancelled():
			calls = []

			def check():
				calls.append(1)
				if len(calls) > 1:  # the first call is the one before any row
					raise streaming.ExportLimit("cancelled", "Cancelled.")

			streaming.write(
				prepared(),
				({"value": i} for i in range(5000)),
				directory / "b.part",
				"xlsx",
				10_000,
				10**8,
				check,
				noop1,
			)

		noop, noop1 = (lambda: None), (lambda n: None)
		with tempfile.TemporaryDirectory() as tmp:
			directory = Path(tmp)
			for call, status in ((too_many, "too large"), (cancelled, "cancelled")):
				with self.subTest(status=status), patch.object(xlsxwriter.Workbook, "close") as close:
					with self.assertRaises(streaming.ExportLimit) as caught:
						call()
					self.assertEqual(caught.exception.status, status)
					close.assert_not_called()
			self.assertFalse(list(directory.glob("*.part")))  # the .part only exists once closed


class TestSavedQueries(FrappeTestCase):
	def test_owner_expiry_and_no_rows(self):
		frappe.set_user("Administrator")
		spec = {"from": "ToDo", "limit": 2, "offset": 3}
		key = saved_queries.save(spec)
		stored = frappe.cache().get_value("jarvis:query:" + key)
		self.assertEqual(stored["spec"], {"from": "ToDo", "select": ["name"]})
		self.assertNotIn("rows", stored)
		self.assertEqual(spec["limit"], 2)
		with (
			impersonate("Guest"),
			self.assertRaisesRegex(InvalidArgumentError, "not found"),
		):
			saved_queries.resolve(key)
		frappe.cache().delete_value("jarvis:query:" + key)
		with self.assertRaisesRegex(InvalidArgumentError, "expired"):
			saved_queries.resolve(key)

	def test_saved_result_read_is_fresh(self):
		frappe.set_user("Administrator")
		key = saved_queries.save({"from": "ToDo"})
		try:
			# Populates any request-local copy of the key.
			saved_queries.resolve(key)
			# Straight to Redis, as another worker or the TTL would: the local copy stays behind.
			cache = frappe.cache()
			cache.delete(cache.make_key("jarvis:query:" + key))
			with self.assertRaisesRegex(InvalidArgumentError, "expired"):
				saved_queries.resolve(key)
		finally:
			frappe.cache().delete_value("jarvis:query:" + key)

	def test_partial_only_ids_at_row_and_character_limits(self):
		from jarvis.tools._result_guard import enforce_result_budget
		from jarvis.tools.query import query

		with patch("jarvis.tools.query.compile_spec") as compile_query:
			q = compile_query.return_value
			q.limit.return_value = q
			q.get_sql.return_value = "SELECT"
			q.run.return_value = [{"name": "one"}]
			self.assertNotIn("result_id", query({"from": "ToDo", "limit": 2}))
			q.run.return_value = [{"name": "one"}] * 3
			with patch("jarvis.tools.query.DEFAULT_LIMIT", 2):
				result = query({"from": "ToDo"})
			self.assertIn("result_id", result)
			self.assertEqual(len(result["rows"]), 2)
			q.run.return_value = [{"name": "x" * 40_000}]
			result = query({"from": "ToDo", "limit": 2})
			guarded, _ = enforce_result_budget(result, "query")
			self.assertEqual(guarded["result_id"], result["result_id"])
			self.assertEqual(saved_queries.resolve(result["result_id"])["limit"], 2)

	def test_explicit_top_n_is_complete_even_when_more_rows_match(self):
		from jarvis.tools.query import query

		with patch("jarvis.tools.query.compile_spec") as compile_query:
			q = compile_query.return_value
			q.limit.return_value = q
			q.get_sql.return_value = "SELECT"
			q.run.return_value = [{"name": str(i)} for i in range(11)]
			result = query({"from": "ToDo", "limit": 10})
		self.assertEqual(len(result["rows"]), 10)
		self.assertTrue(result["more_rows"])
		self.assertNotIn("partial", result)
		self.assertNotIn("result_id", result)

	def test_top_n_saved_id_note_keeps_the_limit(self):
		from jarvis.tools._result_guard import enforce_result_budget
		from jarvis.tools.query import query

		with patch("jarvis.tools.query.compile_spec") as compile_query:
			q = compile_query.return_value
			q.limit.return_value = q
			q.get_sql.return_value = "SELECT"
			q.run.return_value = [{"name": "x" * 40_000}]
			top_n = query({"from": "ToDo", "limit": 5000})
			q.run.return_value = [{"name": "one"}] * 3
			with patch("jarvis.tools.query.DEFAULT_LIMIT", 2):
				uncapped = query({"from": "ToDo"})
		for result, rule in ((top_n, "its limit of 5000 stays"), (uncapped, "For every matching row")):
			guarded, _ = enforce_result_budget(result, "query")
			for shown in (result, guarded):
				with self.subTest(rule=rule, where="query" if shown is result else "guard"):
					self.assertIn(rule, shown["note"])
					self.assertIn(f"result_id='{result['result_id']}'", shown["note"])
		# The guard keeps the query's own sentence instead of restating a rule that is not true.
		self.assertNotIn("For every matching row", enforce_result_budget(top_n, "query")[0]["note"])
		self.assertIn("Never rebuild", enforce_result_budget(top_n, "query")[0]["note"])

	def test_uncapped_budget_event_is_not_partial(self):
		from jarvis.tools.query import query

		for kind, minted in (
			("uncapped", False),
			("measure_failed", False),
			("truncated", True),
			("uncapped_nonrow", True),
		):
			with (
				self.subTest(kind=kind),
				patch("jarvis.tools.query.compile_spec") as compile_query,
				patch(
					"jarvis.tools._result_guard.enforce_result_budget",
					return_value=({}, {"kind": kind}),
				),
			):
				q = compile_query.return_value
				q.limit.return_value = q
				q.get_sql.return_value = "SELECT"
				q.run.return_value = [{"name": "one"}]
				result = query({"from": "ToDo", "limit": 2})
				self.assertEqual("result_id" in result, minted)
				self.assertEqual("partial" in result, minted)

	def test_result_id_survives_site_cache_clear_with_its_ttl(self):
		frappe.set_user("Administrator")
		key = saved_queries.save({"from": "ToDo"})
		try:
			frappe.clear_cache()
			self.assertEqual(saved_queries.resolve(key)["from"], "ToDo")
			ttl = frappe.cache().ttl(frappe.cache().make_key("jarvis:query:" + key))
			self.assertGreater(ttl, 0)
			self.assertLessEqual(ttl, saved_queries.TTL)
		finally:
			frappe.cache().delete_value("jarvis:query:" + key)


class TestBackgroundExportJobs(FrappeTestCase):
	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")
		self.conversation = (
			frappe.get_doc({"doctype": "Jarvis Conversation", "title": "export integration test"})
			.insert(ignore_permissions=True)
			.name
		)
		self.names = []
		frappe.db.commit()

	def tearDown(self):
		frappe.set_user("Administrator")
		for name in self.names:
			if frappe.db.exists(jobs.DOCTYPE, name):
				job = frappe.get_doc(jobs.DOCTYPE, name)
				jobs._delete_file(job)
				frappe.db.delete(jobs.DOCTYPE, {"name": name})
		frappe.db.delete(
			"Notification Log", {"document_type": jobs.DOCTYPE, "document_name": ["in", self.names]}
		)
		frappe.db.delete(
			"Error Log", {"reference_doctype": jobs.DOCTYPE, "reference_name": ["in", self.names]}
		)
		frappe.db.delete("Jarvis Chat Message", {"conversation": self.conversation})
		frappe.db.delete("Jarvis Conversation", {"name": self.conversation})
		frappe.db.commit()
		super().tearDown()

	def error_logs(self, name, doctype=None):
		return frappe.get_all(
			"Error Log",
			filters={"reference_doctype": doctype or jobs.DOCTYPE, "reference_name": name},
			fields=["method", "error"],
		)

	def job(self, **values):
		job = frappe.get_doc(
			{
				"doctype": jobs.DOCTYPE,
				"conversation": self.conversation,
				"source": json.dumps({"spec": {"from": "DocType", "select": ["name"], "limit": 2}}),
				"title": "Export test",
				"format": "csv",
				"status": "running",
				"attempt_token": "abc123",
				"heartbeat": jobs.now_datetime(),
				**values,
			}
		)
		job.flags.jarvis_server_write = True
		job.insert(ignore_permissions=True)
		self.names.append(job.name)
		frappe.db.commit()
		return job

	def test_real_stream_and_publish(self):
		job = self.job()
		with patch.object(jobs, "dispatch"):
			jobs.run(job.name, "abc123")
		job.reload()
		self.assertEqual(job.status, "ready", job.error)
		self.assertEqual(job.rows_written, 2)
		self.assertEqual(job.limited_by_spec, 2)
		self.assertTrue(jobs._file_exists(job))
		self.assertTrue(frappe.get_doc("File", job.file).is_private)
		jobs.announce(job.name)
		self.assertEqual(
			frappe.db.count("Jarvis Chat Message", {"tool_call_id": f"export:{job.name}:ready"}), 1
		)
		self.assertEqual(frappe.db.count("Notification Log", {"document_name": job.name}), 1)

	def test_no_limit_exports_publish_from_saved_id_and_filters(self):
		from openpyxl import load_workbook

		from jarvis.tools._export import resolvers
		from jarvis.tools.export_query import export_query

		# Real handoff, worker, publication, bytes and state; only RQ scheduling is stubbed.
		key = saved_queries.save({"from": "DocType", "select": ["name"], "limit": 1})
		try:
			for fmt, source in (
				("xlsx", {"result_id": key}),
				("csv", {"doctype": "DocType", "fields": ["name"]}),
			):
				with (
					self.subTest(format=fmt),
					patch.object(resolvers, "_HARD_ROW_CEILING", 2),
					patch.object(jobs, "current_conversation", return_value=self.conversation),
					patch.object(jobs, "dispatch"),
				):
					state = export_query(**source, format=fmt)
					name = state["export_job"]
					self.names.append(name)
					frappe.db.set_value(jobs.DOCTYPE, name, {"status": "running", "attempt_token": "abc123"})
					frappe.db.commit()
					jobs.run(name, "abc123")
					job = frappe.get_doc(jobs.DOCTYPE, name)
					self.assertEqual(job.status, "ready", job.error)
					self.assertEqual(job.limited_by_spec, 0)
					self.assertGreater(job.rows_written, 2)
					file = frappe.get_doc("File", job.file)
					self.assertRegex(file.file_name, r"jarvis-export-DocType-[0-9a-f]{6}\." + fmt)
					if fmt == "csv":
						with open(file.get_full_path()) as handle:
							self.assertEqual(sum(1 for _ in csv.reader(handle)), job.rows_written + 1)
					else:
						book = load_workbook(file.get_full_path(), read_only=True)
						try:
							self.assertEqual(sum(1 for _ in book.active.rows), job.rows_written + 1)
						finally:
							book.close()
		finally:
			frappe.cache().delete_value("jarvis:query:" + key)

	def test_cancel_before_worker_start_has_cancelled_chat_and_bell(self):
		for status in ("queued", "paused"):
			with self.subTest(status=status):
				job = self.job(status=status)
				with patch.object(jobs, "dispatch"):
					result = jobs.cancel(job.name)
				self.assertEqual(result["status"], "cancelled")
				self.assertEqual(result["error"], "Export cancelled.")
				self.assertEqual(
					frappe.db.get_value(
						"Jarvis Chat Message", {"tool_call_id": f"export:{job.name}:cancelled"}, "content"
					),
					"Export cancelled.",
				)
				self.assertEqual(
					frappe.db.get_value("Notification Log", {"document_name": job.name}, "subject"),
					"Export cancelled.",
				)

	def test_dispatch_unknown_or_stale_registry_but_not_confirmed_absence(self):
		for count, stale, expected in ((None, False, 1), (0, True, 1), (0, False, 0)):
			with self.subTest(count=count, stale=stale):
				job = self.job(status="queued")
				with (
					patch("jarvis.chat.pump._probe_worker_count", return_value=count),
					patch("jarvis.chat.pump._registry_is_stale", return_value=stale),
					patch("frappe.enqueue") as enqueue,
				):
					jobs.dispatch()
				self.assertEqual(enqueue.call_count, expected)
				self.assertEqual(job.reload().status, "running" if expected else "queued")
				frappe.db.delete(jobs.DOCTYPE, {"name": job.name})
				frappe.db.commit()

	def test_failed_cancellation_commit_rolls_back_state(self):
		job = self.job(status="queued")
		commit = frappe.db.commit
		calls = 0

		def fail_cancellation():
			nonlocal calls
			calls += 1
			if calls == 2:
				raise RuntimeError("commit failed")
			return commit()

		with (
			patch.object(frappe.db, "commit", side_effect=fail_cancellation),
			self.assertRaisesRegex(RuntimeError, "commit failed"),
		):
			jobs.cancel(job.name)
		self.assertEqual(job.reload().status, "queued")
		self.assertFalse(job.cancel_requested)

	def test_failed_notification_rolls_back_message_and_card(self):
		job = self.job()
		with (
			patch.object(jobs, "_commit", side_effect=[None, RuntimeError("commit failed")]),
			self.assertRaisesRegex(RuntimeError, "commit failed"),
		):
			jobs.announce(job.name, initial=True)
		self.assertFalse(frappe.db.exists("Jarvis Chat Message", {"tool_call_id": f"export:{job.name}:card"}))
		self.assertFalse(job.reload().card_message)

	def test_stale_attempt_cannot_change_state(self):
		job = self.job()
		jobs._finish(job.name, "stale", "failed", "wrong worker")
		self.assertEqual(job.reload().status, "running")
		jobs._finish(job.name, "abc123", "paused", "site update")
		self.assertEqual(job.reload().retries, 1)
		frappe.db.set_value(jobs.DOCTYPE, job.name, "status", "running")
		frappe.db.commit()
		jobs._finish(job.name, "abc123", "paused", "site update again")
		self.assertEqual(job.reload().status, "failed")

	def test_cancelled_attempt_cannot_publish(self):
		job = self.job(cancel_requested=1)
		with patch.object(jobs, "dispatch"):
			jobs.run(job.name, "abc123")
		self.assertEqual(job.reload().status, "cancelled")
		self.assertFalse(job.file)

	def test_owner_checks(self):
		job = self.job()
		with impersonate("Guest"):
			with self.assertRaises(frappe.PermissionError):
				jobs.get_jobs([job.name])
			with self.assertRaises(frappe.PermissionError):
				jobs.cancel(job.name)

	def test_jarvis_role_reads_only_owned_jobs(self):
		users = []
		job = self.job()
		try:
			for label in ("owner", "peer"):
				user = frappe.get_doc(
					{
						"doctype": "User",
						"email": f"export-{label}-{frappe.generate_hash(length=8)}@example.com",
						"first_name": label,
						"send_welcome_email": 0,
						"user_type": "System User",
						"roles": [{"role": "Jarvis User"}],
					}
				).insert(ignore_permissions=True)
				users.append(user.name)
			frappe.db.set_value(jobs.DOCTYPE, job.name, "owner", users[0])
			frappe.db.commit()
			job.reload()
			with impersonate(users[0]):
				self.assertTrue(frappe.has_permission(jobs.DOCTYPE, "read", doc=job))
				self.assertEqual(
					frappe.get_list(jobs.DOCTYPE, filters={"name": job.name}, pluck="name"), [job.name]
				)
			with impersonate(users[1]):
				self.assertFalse(frappe.has_permission(jobs.DOCTYPE, "read", doc=job))
				self.assertEqual(frappe.get_list(jobs.DOCTYPE, filters={"name": job.name}, pluck="name"), [])
				self.assertEqual(jobs.get_jobs([job.name]), {job.name: None})
		finally:
			frappe.set_user("Administrator")
			frappe.db.delete(jobs.DOCTYPE, {"name": job.name})
			for user in users:
				frappe.delete_doc("User", user, force=True, ignore_permissions=True)
			frappe.db.commit()

	def test_filter_prepare_runs_real_stream(self):
		p = streaming.prepare({"doctype": "DocType", "fields": ["name"], "filters": {"name": "ToDo"}}, 100)
		db = MariaDBStream(connection_settings()).open()
		try:
			with patch.object(frappe.db, "sql", side_effect=AssertionError("STREAM queried DB")):
				self.assertEqual(list(db.stream_rows(p.sql, p.params)), [{"name": "ToDo"}])
		finally:
			db.close()

	def test_watchdog_sees_cancellation_and_lock_waiters(self):
		job = self.job(cancel_requested=1)
		db = MariaDBStream(connection_settings()).open()
		watchdog = Watchdog(db, job.name, "abc123", frappe.local.site, job.owner, job.conversation)
		try:
			with self.assertRaisesRegex(streaming.ExportLimit, "cancelled"):
				watchdog._tick()
			frappe.db.set_value(jobs.DOCTYPE, job.name, "cancel_requested", 0)
			frappe.db.commit()
			with patch.object(db, "detect_lock_waiters", return_value=True):
				with self.assertRaisesRegex(streaming.ExportLimit, "site update"):
					watchdog._tick()
		finally:
			db.close()

	def test_cancel_kills_only_our_running_query(self):
		job = self.job(started_at=jobs.now_datetime())
		db = MariaDBStream(connection_settings()).open()
		watchdog = Watchdog(db, job.name, "abc123", frappe.local.site, job.owner, job.conversation)
		errors = []

		def slow_query():
			try:
				list(db.stream_rows("SELECT SLEEP(30)", {}))
			except Exception as error:
				errors.append(error)

		thread = threading.Thread(target=slow_query)
		try:
			thread.start()
			watchdog.start()
			frappe.db.set_value(jobs.DOCTYPE, job.name, "cancel_requested", 1)
			frappe.db.commit()
			thread.join(timeout=8)
			self.assertFalse(thread.is_alive(), "Cancellation did not kill the read")
			self.assertTrue(errors)
			self.assertEqual(watchdog.reason.status, "cancelled")
			self.assertEqual(frappe.db.sql("SELECT 1"), ((1,),))
		finally:
			watchdog.stop()
			db.close()
			thread.join(timeout=2)

	def test_over_cap_match_ends_too_large_not_truncated(self):
		# The query asks for one row past the cap; that extra row is how a match
		# larger than the cap is told apart from one that exactly fits.
		job = self.job(source=json.dumps({"spec": {"from": "DocType", "select": ["name"]}}))
		with (
			patch.object(jobs, "limits", return_value=(2, 50 * 1024 * 1024)),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		job.reload()
		self.assertEqual(job.status, "too large", job.error)
		self.assertFalse(job.file)
		self.assertFalse(
			frappe.db.exists("File", {"attached_to_doctype": jobs.DOCTYPE, "attached_to_name": job.name})
		)

	def test_time_cap_kills_our_running_query(self):
		job = self.job(started_at=jobs.now_datetime())
		db = MariaDBStream(connection_settings()).open()
		# No cancel request: only the ten-minute deadline (shortened to one second) may act.
		watchdog = Watchdog(db, job.name, "abc123", frappe.local.site, job.owner, job.conversation, seconds=1)
		errors = []

		def slow_query():
			try:
				list(db.stream_rows("SELECT SLEEP(30)", {}))
			except Exception as error:
				errors.append(error)

		thread = threading.Thread(target=slow_query)
		try:
			thread.start()
			watchdog.start()
			thread.join(timeout=8)
			self.assertFalse(thread.is_alive(), "The time cap did not kill the read")
			self.assertTrue(errors)
			self.assertEqual(watchdog.reason.status, "too slow")
			self.assertEqual(frappe.db.sql("SELECT 1"), ((1,),))
		finally:
			watchdog.stop()
			db.close()
			thread.join(timeout=2)

	def test_stream_session_sets_statement_time_limit(self):
		db = MariaDBStream(connection_settings(), seconds=7).open()
		try:
			with db.connection.cursor() as cursor:
				cursor.execute("SELECT @@SESSION.max_statement_time, @@SESSION.net_write_timeout")
				statement_time, write_timeout = cursor.fetchone()
		finally:
			db.close()
		self.assertEqual((float(statement_time), float(write_timeout)), (7.0, 7.0))

	def streamed(self, db, p):
		return len(list(db.stream_rows(p.sql, p.params)))

	def test_count_matches_the_streamed_rows(self):
		# The legacy v15 filter SQL has its values inlined and params=None; passing () instead would
		# make PyMySQL %-format the LIKE pattern. Both sources must count what they stream.
		sources = (
			{"spec": {"from": "DocType", "select": ["name"]}},
			{"doctype": "DocType", "fields": ["name"], "filters": {"name": ["like", "%Job%"]}},
		)
		for source in sources:
			with self.subTest(source=list(source)):
				p = streaming.prepare(source, 5000)
				db = MariaDBStream(connection_settings()).open()
				try:
					counted = db.count_rows(p.sql, p.params)
					self.assertIsNone(db.count_error)
					self.assertGreater(counted, 0)
					self.assertEqual(counted, self.streamed(db, p))
				finally:
					db.close()

	def test_count_counts_groups_and_distinct_rows(self):
		every = frappe.db.count("DocType")
		for extra in ({"group_by": ["module"]}, {"distinct": True}):
			with self.subTest(extra=extra):
				p = streaming.prepare({"spec": {"from": "DocType", "select": ["module"], **extra}}, 5000)
				db = MariaDBStream(connection_settings()).open()
				try:
					counted = db.count_rows(p.sql, p.params)
					self.assertEqual(counted, self.streamed(db, p))
					self.assertEqual(counted, len(set(frappe.get_all("DocType", pluck="module"))))
					self.assertLess(counted, every)
				finally:
					db.close()

	def test_count_respects_the_cap_and_spec_limit(self):
		db = MariaDBStream(connection_settings()).open()
		try:
			capped = streaming.prepare({"doctype": "DocType", "fields": ["name"]}, 2)
			self.assertEqual(db.count_rows(capped.sql, capped.params), 3)  # cap + 1
			limited = streaming.prepare({"spec": {"from": "DocType", "select": ["name"], "limit": 5}}, 100)
			raw = db.count_rows(limited.sql, limited.params)
			self.assertEqual(raw, 6)  # limit + 1
			self.assertEqual(streaming.expected_total(raw, limited.limit, 100), 5)
		finally:
			db.close()

	def test_count_gives_up_within_its_limit(self):
		db = MariaDBStream(connection_settings(), seconds=10).open()
		try:
			started = time.monotonic()
			self.assertIsNone(db.count_rows(SLOW_COUNT, None, seconds=1))
			self.assertLess(time.monotonic() - started, 4)
			self.assertEqual(db.count_error, 1969)
			# The connection is still good, and the session limit was not touched.
			with db.connection.cursor() as cursor:
				cursor.execute("SELECT 1, @@SESSION.max_statement_time")
				one, session_limit = cursor.fetchone()
			self.assertEqual((one, float(session_limit)), (1, 10.0))
			with self.assertRaises(ValueError):
				db.count_rows("SELECT 1", None, seconds=0)  # 0 would mean no limit at all
		finally:
			db.close()

	def test_count_errors_are_swallowed(self):
		db = MariaDBStream(connection_settings()).open()
		try:
			for sql, code in (("SELECT 1 AS a, 2 AS a", 1060), ("SELEKT nonsense", 1064)):
				with self.subTest(code=code):
					self.assertIsNone(db.count_rows(sql, None))
					self.assertEqual(db.count_error, code)
			self.assertEqual(db.count_rows("SELECT 1", None), 1)
			self.assertIsNone(db.count_error)
			# A real bug must surface, not be mistaken for "no count".
			with patch.object(db.connection, "cursor", side_effect=RuntimeError("bug")):
				with self.assertRaises(RuntimeError):
					db.count_rows("SELECT 1", None)
		finally:
			db.close()

	def test_cancel_during_the_count_ends_without_streaming(self):
		job = self.job(source=json.dumps({"spec": {"from": "DocType", "select": ["name"]}}))
		real_count = MariaDBStream.count_rows
		seen = {}

		def cancel_then_count(db, sql, params):
			frappe.db.set_value(jobs.DOCTYPE, job.name, "cancel_requested", 1)
			frappe.db.commit()
			started = time.monotonic()
			seen["result"] = real_count(db, SLOW_COUNT, None, seconds=8)
			seen["seconds"] = time.monotonic() - started
			return seen["result"]

		order = []
		with (
			patch.object(MariaDBStream, "count_rows", autospec=True, side_effect=cancel_then_count),
			spy(MariaDBStream, "stream_rows", order),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		self.assertEqual(job.reload().status, "cancelled")
		self.assertEqual(order, [])  # the stream never started
		# The watchdog's KILL QUERY reached the count before its own 8 s limit.
		self.assertIsNone(seen["result"])
		self.assertLess(seen["seconds"], 7)

	def test_count_outcome_reaches_telemetry(self):
		def timed_out(db, *args, **kwargs):
			db.count_error = 1969

		for patcher, detail in (
			(contextlib.nullcontext(), "count:ok"),
			(
				patch.object(MariaDBStream, "count_rows", autospec=True, side_effect=timed_out),
				"count:1969",
			),
		):
			with self.subTest(detail=detail):
				job = self.job()
				with (
					patcher,
					patch("jarvis.telemetry.record_export_event") as record,
					patch.object(jobs, "dispatch"),
				):
					jobs.run(job.name, "abc123")
				record.assert_called_once()
				self.assertEqual(record.call_args.kwargs["detail"], detail)
				self.assertNotIn("SELECT", detail)

	def test_run_stores_the_total_while_running(self):
		job = self.job(source=json.dumps({"spec": {"from": "DocType", "select": ["name"]}}))
		real_write = streaming.write
		seen = {}

		def look_then_write(*args, **kwargs):
			seen.update(jobs.get_jobs([job.name])[job.name])
			return real_write(*args, **kwargs)

		order = []
		with (
			patch.object(streaming, "write", look_then_write),
			spy(Watchdog, "start", order),
			spy(MariaDBStream, "count_rows", order),
			spy(MariaDBStream, "stream_rows", order),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		self.assertEqual(order, ["start", "count_rows", "stream_rows"])
		self.assertEqual(seen["status"], "running")
		self.assertEqual(seen["total_rows"], frappe.db.count("DocType"))
		self.assertEqual(job.reload().status, "ready", job.error)
		self.assertIsNone(jobs.public(job)["total_rows"])

	def test_run_without_a_count_streams_as_before(self):
		job = self.job()
		real_write = streaming.write
		seen = {}

		def look_then_write(*args, **kwargs):
			seen.update(jobs.get_jobs([job.name])[job.name])
			return real_write(*args, **kwargs)

		with (
			patch.object(MariaDBStream, "count_rows", return_value=None),
			patch.object(streaming, "write", look_then_write),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		job.reload()
		self.assertEqual((job.status, job.rows_written), ("ready", 2))
		self.assertIsNone(seen["total_rows"])
		self.assertEqual(job.total_rows, 0)

	def test_count_over_the_cap_ends_too_large_without_streaming(self):
		job = self.job(source=json.dumps({"spec": {"from": "DocType", "select": ["name"]}}))
		order = []
		with (
			patch.object(jobs, "limits", return_value=(2, 50 * 1024 * 1024)),
			spy(MariaDBStream, "stream_rows", order),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		job.reload()
		self.assertEqual(job.status, "too large")
		self.assertEqual(job.error, str(streaming.too_many_rows(2)))
		self.assertEqual(order, [])
		self.assertFalse(job.file)

	def test_over_cap_without_a_count_still_ends_too_large(self):
		# What a big production set meets when the count does not answer in time: the stream's own cap.
		job = self.job(source=json.dumps({"spec": {"from": "DocType", "select": ["name"]}}))
		order = []
		with (
			patch.object(jobs, "limits", return_value=(2, 50 * 1024 * 1024)),
			patch.object(MariaDBStream, "count_rows", return_value=None),
			spy(MariaDBStream, "stream_rows", order),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		job.reload()
		self.assertEqual(job.status, "too large")
		self.assertEqual(job.error, str(streaming.too_many_rows(2)))
		self.assertEqual(order, ["stream_rows"])
		self.assertFalse(job.file)

	def phase_seen_by(self, owner, attribute, job_name, seen):
		"""Patch ``owner.attribute`` to record the job's phase and status as the real call starts."""
		original = getattr(owner, attribute)

		def look(*args, **kwargs):
			state = frappe.db.get_value(
				jobs.DOCTYPE, job_name, ["phase", "status", "rows_written"], as_dict=True
			)
			seen.append(dict(state))
			return original(*args, **kwargs)

		return patch.object(owner, attribute, autospec=True, side_effect=look)

	def test_finalizing_is_recorded_before_the_workbook_closes(self):
		import xlsxwriter

		job = self.job(format="xlsx")
		at_close, at_publish = [], []
		with (
			self.phase_seen_by(xlsxwriter.Workbook, "close", job.name, at_close),
			self.phase_seen_by(jobs, "_publish_locked", job.name, at_publish),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		expected = {"phase": "finalizing", "status": "running", "rows_written": 2}
		self.assertEqual(at_close, [expected])
		self.assertEqual(at_publish, [expected])
		job.reload()
		self.assertEqual(job.status, "ready", job.error)
		self.assertFalse(job.phase)
		self.assertIsNone(jobs.public(job)["phase"])

	def test_finalizing_is_recorded_before_a_csv_publish(self):
		job = self.job()
		at_publish = []
		with (
			self.phase_seen_by(jobs, "_publish_locked", job.name, at_publish),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		self.assertEqual(at_publish, [{"phase": "finalizing", "status": "running", "rows_written": 2}])
		self.assertEqual(job.reload().status, "ready", job.error)

	def test_a_failed_phase_write_does_not_fail_the_export(self):
		job = self.job()
		with (
			patch.object(jobs, "_set_progress", side_effect=RuntimeError("write failed")),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		self.assertEqual(job.reload().status, "ready", job.error)
		titles = [log.method for log in self.error_logs(job.name)]
		self.assertIn(f"jarvis.export finalizing: {job.name}", titles)

	def test_run_relies_on_the_ten_minute_defaults(self):
		# run() builds both objects without a `seconds` argument.
		for cls in (Watchdog, MariaDBStream):
			self.assertEqual(inspect.signature(cls).parameters["seconds"].default, 600, cls.__name__)
		self.assertEqual(inspect.signature(MariaDBStream.count_rows).parameters["seconds"].default, 3)

	def test_pause_schedules_retry_five_minutes_out(self):
		job = self.job()
		jobs._finish(job.name, "abc123", "paused", "x")
		job.reload()
		self.assertEqual(job.status, "paused")
		self.assertEqual(job.retries, 1)
		self.assertFalse(job.finished_at)
		due = jobs.now_datetime() + datetime.timedelta(minutes=5)
		self.assertLess(abs((frappe.utils.get_datetime(job.next_attempt) - due).total_seconds()), 15)
		# Re-armed and paused again: the one retry is spent.
		frappe.db.set_value(jobs.DOCTYPE, job.name, {"status": "running", "attempt_token": "abc123"})
		frappe.db.commit()
		jobs._finish(job.name, "abc123", "paused", "y")
		job.reload()
		self.assertEqual(job.status, "failed")
		self.assertIn("interrupted again", job.error)

	def test_dispatch_skips_a_retry_that_is_not_yet_due(self):
		# dispatch() takes the oldest eligible job site-wide, so assert on this job's row and calls.
		job = self.job(
			status="paused",
			attempt_token="",
			next_attempt=jobs.now_datetime() + datetime.timedelta(minutes=5),
		)

		def started():
			return [c for c in enqueue.call_args_list if c.kwargs.get("name") == job.name]

		with (
			patch("jarvis.chat.pump._probe_worker_count", return_value=3),
			patch.object(frappe, "enqueue") as enqueue,
		):
			jobs.dispatch()
			self.assertEqual(job.reload().status, "paused")
			self.assertEqual(started(), [])
			frappe.db.set_value(
				jobs.DOCTYPE,
				job.name,
				"next_attempt",
				jobs.now_datetime() - datetime.timedelta(seconds=1),
			)
			frappe.db.commit()
			jobs.dispatch()
			job.reload()
			self.assertEqual(job.status, "running")
			calls = started()
			self.assertEqual(len(calls), 1)
			self.assertEqual(calls[0].kwargs["job_id"], "export-" + job.attempt_token)

	def ready_job(self):
		"""A job run to ``ready`` with its card in place, as a real publish leaves it."""
		job = self.job()
		with patch.object(jobs, "dispatch"):
			jobs.announce(job.name, initial=True)
			jobs.run(job.name, "abc123")
		return job.reload()

	def stamp(self, job, **columns):
		"""Move lifecycle columns directly, without touching ``modified``."""
		frappe.db.set_value(jobs.DOCTYPE, job.name, columns, update_modified=False)
		frappe.db.commit()
		return job.reload()

	def days_ago(self, days):
		return frappe.utils.add_to_date(jobs.now_datetime(), days=-days)

	def reap(self):
		with patch.object(jobs, "dispatch"):
			jobs.reap()

	def test_published_file_lives_in_jarvis_exports_for_thirty_days(self):
		first = self.ready_job()
		second = self.ready_job()
		file = frappe.get_doc("File", first.file)
		self.assertEqual(file.folder, "Home/Jarvis Exports")
		self.assertEqual(file.attached_to_name, first.name)
		self.assertEqual(file.is_private, 1)
		self.assertTrue(file.file_name.startswith("jarvis-export-"))
		self.assertEqual(frappe.db.get_value("File", second.file, "folder"), "Home/Jarvis Exports")
		for job in (first, second):
			offset = frappe.utils.get_datetime(job.expires_at) - jobs.now_datetime()
			self.assertLess(abs(offset - datetime.timedelta(days=30)), datetime.timedelta(seconds=60))
		self.assertEqual(frappe.db.count("File", {"is_folder": 1, "file_name": "Jarvis Exports"}), 1)

	def test_retention_setting_zero_keeps_files_and_rows(self):
		with patch.dict(frappe.conf, {"jarvis_export_retention_days": 0}):
			job = self.ready_job()
			self.assertIsNone(job.expires_at)
			file_name = job.file
			# A ready row never leaves by age: its File would be swept as an orphan.
			job = self.stamp(job, finished_at=self.days_ago(40), creation=self.days_ago(40))
			self.reap()
			self.assertEqual(job.reload().status, "ready")
			self.assertTrue(frappe.db.exists("File", file_name))
			self.assertIsNone(jobs.public(job)["expires_at"])
			self.assertEqual(jobs.public(job)["status"], "ready")

	def test_retention_change_applies_to_stored_files_only_downward(self):
		# (label, days setting, finished_at age, stored expires_at offset in days, expected status)
		cases = [
			("lowered below the file's age", 7, 8, 22, "expired"),
			("lowered but the file is younger", 7, 3, 27, "ready"),
			("raised never extends", 90, 31, -1, "expired"),
			("zero keeps the stored expiry", 0, 31, -1, "expired"),
			("zero does not expire a kept file", 0, 31, None, "ready"),
		]
		for label, days, age, offset, expected in cases:
			with self.subTest(label), patch.dict(frappe.conf, {"jarvis_export_retention_days": 30}):
				job = self.ready_job()
				expires = None if offset is None else jobs.now_datetime() + datetime.timedelta(days=offset)
				job = self.stamp(job, finished_at=self.days_ago(age), expires_at=expires)
				file_name = job.file
				with patch.dict(frappe.conf, {"jarvis_export_retention_days": days}):
					self.assertEqual(jobs.public(job)["status"], expected)
					self.reap()
				self.assertEqual(job.reload().status, expected)
				self.assertEqual(bool(frappe.db.exists("File", file_name)), expected == "ready")
		# The card shows the lowered date, not the stored one.
		job = self.ready_job()
		job = self.stamp(job, finished_at=self.days_ago(3))
		with patch.dict(frappe.conf, {"jarvis_export_retention_days": 7}):
			shown = frappe.utils.get_datetime(jobs.public(job)["expires_at"].replace("+00:00", ""))
		self.assertLess(
			abs(shown - (frappe.utils.get_datetime(job.finished_at) + datetime.timedelta(days=7))),
			datetime.timedelta(seconds=2),
		)

	def test_expiry_is_quiet(self):
		job = self.ready_job()
		expired_key = f"export:{job.name}:expired"

		def expired_messages():
			return frappe.db.count(
				"Jarvis Chat Message", {"conversation": self.conversation, "tool_call_id": expired_key}
			)

		def bells():
			return frappe.db.count("Notification Log", {"document_name": job.name})

		bells_before = bells()
		file_name = job.file
		# Created 31 days ago with a 30 day file: expires_at is about a day past, as in real life.
		job = self.stamp(job, creation=self.days_ago(31), expires_at=self.days_ago(1))
		self.reap()
		self.assertEqual(job.reload().status, "expired")
		self.assertFalse(frappe.db.exists("File", file_name))
		self.assertTrue(frappe.db.exists("Jarvis Chat Message", job.card_message))
		stored = json.loads(frappe.db.get_value("Jarvis Chat Message", job.card_message, "tool_result"))
		self.assertEqual(stored["data"]["rows_written"], 2)
		self.assertEqual((expired_messages(), bells()), (0, bells_before))
		self.assertEqual(job.announced_status, "expired")
		# expires_at records when it expired (here), not the date it was due, so the 30 days start now.
		self.assertLess(abs(job.expires_at - jobs.now_datetime()), datetime.timedelta(minutes=1))
		# The row stays visible for 30 days after it expired, and a second pass changes nothing.
		self.reap()
		self.assertEqual(job.reload().status, "expired")
		self.assertEqual((expired_messages(), bells()), (0, bells_before))
		self.stamp(job, expires_at=self.days_ago(29))
		self.reap()
		self.assertTrue(frappe.db.exists(jobs.DOCTYPE, job.name))
		self.stamp(job, expires_at=self.days_ago(31))
		self.reap()
		self.assertFalse(frappe.db.exists(jobs.DOCTYPE, job.name))

	def test_lowered_retention_shows_expired_before_the_row_is_purged(self):
		# A file due long ago (setting lowered, or the reaper was down) must not be expired and
		# deleted in one pass: the card has to show "Expired" for the 30 days.
		job = self.ready_job()
		job = self.stamp(job, finished_at=self.days_ago(100), expires_at=self.days_ago(60))
		self.reap()
		self.assertTrue(frappe.db.exists(jobs.DOCTYPE, job.name))
		self.assertEqual(job.reload().status, "expired")
		self.assertEqual(jobs.public(job)["status"], "expired")

	def test_row_retention_spares_ready_rows_and_purges_other_terminal_rows(self):
		kept = self.ready_job()
		kept = self.stamp(kept, creation=self.days_ago(40))
		gone = self.job(status="failed", attempt_token="", announced_status="failed")
		self.stamp(gone, creation=self.days_ago(40))
		self.reap()
		self.assertEqual(kept.reload().status, "ready")
		self.assertFalse(frappe.db.exists(jobs.DOCTYPE, gone.name))

	def test_missing_exports_folder_is_recreated_for_the_publisher(self):
		user = frappe.get_doc(
			{
				"doctype": "User",
				"email": f"export-folder-{frappe.generate_hash(length=8)}@example.com",
				"first_name": "folder",
				"send_welcome_email": 0,
				"user_type": "System User",
				"roles": [{"role": "Jarvis User"}],
			}
		).insert(ignore_permissions=True)
		try:
			frappe.db.delete("File", {"name": jobs.EXPORTS_FOLDER})
			frappe.db.commit()
			with impersonate(user.name):
				jobs.ensure_exports_folder()
				# A second call finds the folder and inserts nothing.
				jobs.ensure_exports_folder()
			folder = frappe.db.get_value(
				"File", jobs.EXPORTS_FOLDER, ["owner", "is_folder", "folder"], as_dict=True
			)
			self.assertEqual((folder.owner, folder.is_folder, folder.folder), ("Administrator", 1, "Home"))
		finally:
			frappe.set_user("Administrator")
			frappe.delete_doc("User", user.name, force=True, ignore_permissions=True)
			frappe.db.commit()

	def test_publish_creates_the_folder_and_survives_a_folder_failure(self):
		frappe.db.delete("File", {"name": jobs.EXPORTS_FOLDER})
		frappe.db.commit()
		job = self.ready_job()
		self.assertEqual(frappe.db.get_value("File", job.file, "folder"), jobs.EXPORTS_FOLDER)
		# A folder that cannot be made never fails a finished export; it files elsewhere.
		frappe.db.delete("File", {"name": jobs.EXPORTS_FOLDER})
		frappe.db.commit()
		with patch.object(jobs, "ensure_exports_folder", side_effect=RuntimeError("no folder")):
			other = self.ready_job()
		self.assertEqual(other.status, "ready", other.error)
		self.assertNotEqual(frappe.db.get_value("File", other.file, "folder"), jobs.EXPORTS_FOLDER)
		self.assertEqual(len(self.error_logs(other.name)), 1)
		jobs.ensure_exports_folder()

	def test_expiry_deletes_only_the_linked_file(self):
		job = self.ready_job()
		file = frappe.get_doc("File", job.file)
		path = Path(file.get_full_path())
		twin = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": file.file_name,
				"is_private": 1,
				"file_url": file.file_url,
				"file_size": file.file_size,
				"attached_to_doctype": "Jarvis Conversation",
				"attached_to_name": self.conversation,
			}
		)
		twin.flags.copy_from_existing_file = True
		twin.insert(ignore_permissions=True)
		frappe.db.commit()
		try:
			self.stamp(job, expires_at=self.days_ago(1))
			self.reap()
			self.assertEqual(job.reload().status, "expired")
			self.assertFalse(frappe.db.exists("File", file.name))
			self.assertTrue(frappe.db.exists("File", twin.name))
			self.assertTrue(path.exists())
		finally:
			frappe.delete_doc("File", twin.name, force=True, ignore_permissions=True)
			path.unlink(missing_ok=True)
			frappe.db.commit()

	def test_reaper_retries_worker_death_only_once(self):
		old = frappe.utils.add_to_date(jobs.now_datetime(), minutes=-2)
		job = self.job(heartbeat=old, started_at=old)
		with patch.object(jobs, "dispatch"):
			jobs.reap()
			self.assertEqual(job.reload().status, "queued")
			self.assertEqual(job.retries, 1)
			frappe.db.set_value(
				jobs.DOCTYPE, job.name, {"status": "running", "heartbeat": old, "attempt_token": "second"}
			)
			frappe.db.commit()
			jobs.reap()
		self.assertEqual(job.reload().status, "failed")

	def test_same_query_joins_job_and_highlights_do_not_collide(self):
		source = {"spec": {"from": "DocType", "select": ["name"], "limit": 2}}
		with (
			patch.object(jobs, "current_conversation", return_value=self.conversation),
			patch.object(jobs, "dispatch"),
		):
			first = jobs.enqueue(source, "csv", "First")
			self.names.append(first["export_job"])
			second = jobs.enqueue(source, "csv", "Second")
		self.assertEqual(first["export_job"], second["export_job"])
		# Joining inserts nothing: the card already lives in this chat.
		self.assertEqual(
			frappe.db.count(
				"Jarvis Chat Message",
				{"conversation": self.conversation, "tool_call_id": f"export:{first['export_job']}:card"},
			),
			1,
		)
		self.assertNotEqual(
			jobs._key(source, "xlsx", None),
			jobs._key(source, "xlsx", [{"column": "name", "operator": "eq", "value": "A"}]),
		)

	def test_conversation_deletion_revokes_attempt(self):
		job = self.job()
		jobs.on_conversation_trash(frappe.get_doc("Jarvis Conversation", self.conversation))
		frappe.db.commit()
		self.assertFalse(jobs._attempt(job.name, "abc123"))

	def test_reaper_removes_registered_file_after_job_disappears(self):
		job = self.job()
		with patch.object(jobs, "dispatch"):
			jobs.run(job.name, "abc123")
			job.reload()
			file_name = job.file
			path = Path(frappe.get_doc("File", file_name).get_full_path())
			frappe.db.delete(jobs.DOCTYPE, {"name": job.name})
			frappe.db.commit()
			jobs.reap()
		self.assertFalse(frappe.db.exists("File", file_name))
		self.assertFalse(path.exists())

	def test_site_capacity_and_one_job_per_owner(self):
		from unittest.mock import MagicMock

		from frappe.utils.background_jobs import generate_qname

		first = self.job(status="queued")
		second = self.job(status="queued")
		other = self.job(status="queued", owner="Guest")
		# Document.insert assigns the current session owner even with ignore_permissions.
		frappe.db.set_value(jobs.DOCTYPE, other.name, "owner", "Guest")
		frappe.db.commit()
		worker = MagicMock()
		worker.queue_names.return_value = [generate_qname("long")]
		with (
			patch("frappe.utils.background_jobs.get_workers", return_value=[worker] * 3),
			patch("frappe.enqueue") as enqueue,
		):
			jobs.dispatch()
		self.assertEqual(enqueue.call_count, 2)
		self.assertEqual(first.reload().status, "running")
		self.assertEqual(second.reload().status, "queued")
		self.assertEqual(other.reload().status, "running")

	def test_ready_reuse_rechecks_permission_scope(self):
		from dataclasses import replace

		source = {"spec": {"from": "DocType", "select": ["name"], "limit": 2}}
		job = self.job(fingerprint=jobs._key(source, "csv", None))
		with patch.object(jobs, "dispatch"):
			jobs.run(job.name, "abc123")
		self.assertEqual(jobs.existing(source, "csv", conversation=job.conversation)["export_job"], job.name)
		# Another chat never reuses it, and neither does a caller with no chat.
		self.assertIsNone(jobs.existing(source, "csv", conversation="another-chat"))
		self.assertIsNone(jobs.existing(source, "csv", conversation=None))
		changed = replace(streaming.prepare(source, 200_000), sql="Changed row-permission scope")
		with patch.object(jobs, "prepare", return_value=changed):
			self.assertIsNone(jobs.existing(source, "csv", conversation=job.conversation))

	def test_export_query_handoff_and_saved_id(self):
		from jarvis.tools._export import resolvers
		from jarvis.tools.export_query import export_query

		key = saved_queries.save({"from": "DocType", "select": ["name"], "limit": 1})
		with (
			patch.object(resolvers, "_HARD_ROW_CEILING", 2),
			patch.object(jobs, "current_conversation", return_value=self.conversation),
			patch.object(jobs, "dispatch"),
		):
			state = export_query(result_id=key, format="csv")
		self.names.append(state["export_job"])
		job = frappe.get_doc(jobs.DOCTYPE, state["export_job"])
		self.assertEqual(job.status, "queued")
		self.assertNotIn("limit", json.loads(job.source)["spec"])
		with self.assertRaises(InvalidArgumentError):
			export_query(doctype="DocType", result_id=key)

	def test_upload_limit_hands_off_to_generated_file_budget(self):
		from jarvis.tools.export_query import export_query

		with (
			patch("frappe.utils.file_manager.get_max_file_size", return_value=1),
			patch.object(jobs, "current_conversation", return_value=self.conversation),
			patch.object(jobs, "dispatch"),
		):
			state = export_query(spec={"from": "DocType", "select": ["name"], "limit": 2}, format="csv")
		self.names.append(state["export_job"])
		self.assertEqual(state["status"], "queued")

	def test_literal_percent_filters_use_real_parameters(self):
		values = ["50% complete", "100%% literal", "%(value)s", "quote ' and slash \\"]
		docs = [frappe.get_doc({"doctype": "ToDo", "description": value}).insert() for value in values]
		frappe.db.commit()
		try:
			for doc, value in zip(docs, values, strict=True):
				for filter_value in (value, ["in", [value]]):
					with self.subTest(value=filter_value):
						p = streaming.prepare(
							{
								"doctype": "ToDo",
								"fields": ["name"],
								"filters": {
									"description": filter_value,
								},
							},
							100,
						)
						db = MariaDBStream(connection_settings()).open()
						try:
							self.assertEqual(list(db.stream_rows(p.sql, p.params)), [{"name": doc.name}])
						finally:
							db.close()
			filters = {"description": ["like", "50%complete"]}
			p = streaming.prepare({"doctype": "ToDo", "fields": ["name"], "filters": filters}, 100)
			db = MariaDBStream(connection_settings()).open()
			try:
				self.assertEqual(
					list(db.stream_rows(p.sql, p.params)), frappe.get_list("ToDo", filters=filters)
				)
			finally:
				db.close()
		finally:
			for doc in docs:
				frappe.delete_doc("ToDo", doc.name, force=True)
			frappe.db.commit()

	def test_heartbeat_uses_utc_even_when_database_and_site_clocks_differ(self):
		job = self.job(started_at=jobs.now_datetime())
		db = MariaDBStream(connection_settings()).open()
		try:
			with db.control.cursor() as cursor:
				cursor.execute("SET SESSION time_zone='+09:00'")
			with patch("frappe.utils.get_system_timezone", return_value="America/New_York"):
				Watchdog(db, job.name, "abc123", frappe.local.site, job.owner, job.conversation)._tick()
				frappe.db.commit()
				job.reload()
				self.assertLess(abs((jobs.now_datetime() - job.heartbeat).total_seconds()), 5)
				self.assertTrue(jobs.public(job)["started_at"].endswith("+00:00"))
				with patch.object(jobs, "dispatch"):
					jobs.reap()
				self.assertEqual(job.reload().status, "running")
				self.assertEqual(job.retries, 0)
		finally:
			db.close()

	def test_expiry_selection_does_not_use_database_clock(self):
		job = self.job(
			status="ready",
			expires_at=frappe.utils.add_to_date(jobs.now_datetime(), seconds=-1),
			announced_status="ready",
		)
		jobs.announce(job.name, initial=True)
		previous = frappe.db.sql("SELECT @@session.time_zone")[0][0]
		try:
			frappe.db.sql("SET SESSION time_zone='-12:00'")
			with patch.object(jobs, "dispatch"):
				jobs.reap()
			self.assertEqual(job.reload().status, "expired")
		finally:
			frappe.db.sql("SET SESSION time_zone=%s", (previous,))

	def test_dispatch_finds_eligible_owner_beyond_200_waiting_jobs(self):
		from unittest.mock import MagicMock

		from frappe.utils.background_jobs import generate_qname

		self.job(status="running")
		template = self.job(status="queued")
		# Insert cheap copies of the validated fixture to exercise the real SQL LIMIT.
		for _ in range(200):
			name = frappe.generate_hash(length=10)
			frappe.db.sql(
				"INSERT INTO `tabJarvis Export Job` (name,owner,creation,status,conversation) "
				"SELECT %s,owner,creation,status,conversation FROM `tabJarvis Export Job` WHERE name=%s",
				(name, template.name),
			)
			self.names.append(name)
		eligible = self.job(status="queued")
		frappe.db.set_value(jobs.DOCTYPE, eligible.name, "owner", "Guest")
		frappe.db.commit()
		worker = MagicMock()
		worker.queue_names.return_value = [generate_qname("long")]
		with (
			patch("frappe.utils.background_jobs.get_workers", return_value=[worker] * 3),
			patch("frappe.enqueue") as enqueue,
		):
			jobs.dispatch()
		self.assertEqual(enqueue.call_count, 1)
		self.assertEqual(eligible.reload().status, "running")
		self.assertEqual(template.reload().status, "queued")

	def test_state_fetch_batches_jobs_and_file_records(self):
		first, second = self.job(), self.job()
		with patch.object(jobs, "dispatch"):
			jobs.run(first.name, "abc123")
			jobs.run(second.name, "abc123")
		with patch.object(frappe, "get_list", wraps=frappe.get_list) as get_list:
			states = jobs.get_jobs([first.name, second.name, "missing"])
		self.assertEqual(states[first.name]["status"], "ready")
		self.assertEqual(states[second.name]["status"], "ready")
		self.assertIsNone(states["missing"])
		self.assertEqual([c.args[0] for c in get_list.call_args_list], [jobs.DOCTYPE, "File"])

	def test_kill_switch_refuses_enqueue_and_holds_the_queue(self):
		source = {"spec": {"from": "DocType", "select": ["name"], "limit": 2}}
		with patch.dict(frappe.conf, {"jarvis_background_export_disabled": 1}):
			with self.assertRaisesRegex(jobs.BackgroundUnavailable, "turned off"):
				jobs.enqueue(source, "csv")
			# A queued job waits: dispatch() reserves nothing while the switch is on.
			queued = self.job(status="queued", attempt_token="")
			with (
				patch("jarvis.chat.pump._probe_worker_count", return_value=3),
				patch.object(frappe, "enqueue") as enqueue,
			):
				jobs.dispatch()
			self.assertEqual(queued.reload().status, "queued")
			self.assertEqual([c for c in enqueue.call_args_list if c.kwargs.get("name") == queued.name], [])
			# An attempt reserved just before the flip goes back to the queue, unspent.
			reserved = self.job()
			with patch.object(jobs, "dispatch"):
				jobs.run(reserved.name, "abc123")
			reserved.reload()
			self.assertEqual(reserved.status, "queued")
			self.assertEqual(reserved.attempt_token or "", "")
			self.assertEqual((reserved.retries or 0, reserved.attempts or 0), (0, 0))
		# Without the switch the same queued job is reserved.
		with (
			patch("jarvis.chat.pump._probe_worker_count", return_value=3),
			patch.object(frappe, "enqueue") as enqueue,
		):
			jobs.dispatch()
		self.assertEqual(len([c for c in enqueue.call_args_list if c.kwargs.get("name") == queued.name]), 1)
		self.assertEqual(queued.reload().status, "running")

	def test_kill_switch_string_zero_is_off(self):
		# bench set-config without -p stores "0" as a string; it must not disable exports.
		source = {"spec": {"from": "DocType", "select": ["name"], "limit": 2}}
		with (
			patch.dict(frappe.conf, {"jarvis_background_export_disabled": "0"}),
			patch.object(jobs, "current_conversation", return_value=self.conversation),
			patch.object(jobs, "dispatch"),
		):
			state = jobs.enqueue(source, "csv")
		self.names.append(state["export_job"])
		self.assertEqual(state["status"], "queued")

	def test_xlsx_needs_xlsxwriter_but_csv_does_not(self):
		source = {"spec": {"from": "DocType", "select": ["name"], "limit": 2}}
		with (
			patch.object(jobs, "_xlsxwriter_available", return_value=False),
			patch.object(jobs, "current_conversation", return_value=self.conversation),
			patch.object(jobs, "dispatch"),
		):
			with self.assertRaisesRegex(jobs.BackgroundUnavailable, "export it as CSV"):
				jobs.enqueue(source, "xlsx")
			state = jobs.enqueue(source, "csv")
			self.names.append(state["export_job"])
			self.assertEqual(state["status"], "queued")
			# A job queued before the package vanished fails with the same text, not a retry loop.
			job = self.job(format="xlsx")
			jobs.run(job.name, "abc123")
		job.reload()
		self.assertEqual(job.status, "failed")
		self.assertIn("export it as CSV", job.error)
		self.assertEqual(len(self.error_logs(job.name)), 1)

	def test_refused_handoff_keeps_the_narrowing_advice(self):
		from jarvis.tools.export_query import export_query

		real_has_permission = frappe.has_permission

		def no_job_access(doctype=None, *args, **kwargs):
			if doctype == jobs.DOCTYPE:
				return False
			return real_has_permission(doctype, *args, **kwargs)

		busy = InvalidArgumentError("Exports are busy; please try again shortly.")
		unsupported = InvalidArgumentError("x does not support background streaming")
		# (label, a patch that makes the background route refuse, the reason the agent must see)
		blockers = [
			(
				"no job access",
				lambda: patch.object(frappe, "has_permission", side_effect=no_job_access),
				"role cannot",
			),
			(
				"no chat",
				lambda: patch.object(jobs, "current_conversation", return_value=None),
				"Start it from a Jarvis chat",
			),
			(
				"kill switch",
				lambda: patch.dict(frappe.conf, {"jarvis_background_export_disabled": 1}),
				"turned off",
			),
			("busy", lambda: patch.object(jobs, "_lock", side_effect=busy), "Exports are busy"),
			(
				"prepare",
				lambda: patch.object(jobs, "prepare", side_effect=unsupported),
				"background streaming",
			),
		]
		spec = {"from": "DocType", "select": ["name"]}
		for label, blocker, reason in blockers:
			with (
				self.subTest(blocker=label),
				patch.object(resolvers, "_HARD_ROW_CEILING", 1),
				patch.object(jobs, "current_conversation", return_value=self.conversation),
				patch.object(jobs, "dispatch"),
				blocker(),
				self.assertRaises(InvalidArgumentError) as ctx,
			):
				export_query(spec=spec, format="csv")
			self.assertNotIsInstance(ctx.exception, frappe.PermissionError)
			self.assertIn("Narrow", str(ctx.exception))
			self.assertIn(reason, str(ctx.exception))
		# The upload-size route carries its own advice too.
		with (
			patch("frappe.utils.file_manager.get_max_file_size", return_value=1),
			patch.object(jobs, "current_conversation", return_value=None),
			self.assertRaises(InvalidArgumentError) as ctx,
		):
			export_query(spec={"from": "DocType", "select": ["name"], "limit": 2}, format="csv")
		self.assertIn("Narrow", str(ctx.exception))
		self.assertIn("Start it from a Jarvis chat", str(ctx.exception))

	def test_cancel_after_ready_does_not_flag_the_job(self):
		job = self.ready_job()
		real = jobs._get_job
		stale = real(job.name)
		stale.status, stale.started_at = "running", jobs.now_datetime()
		seen = []

		def lagging(name, **kwargs):
			# The first two reads (owner check, locked re-read) see the job as still running:
			# the publish committed in between. The fenced write must not flag it.
			seen.append(kwargs)
			return stale if len(seen) <= 2 else real(name, **kwargs)

		with patch.object(jobs, "_get_job", side_effect=lagging), patch.object(jobs, "dispatch"):
			state = jobs.cancel(job.name)
		self.assertTrue(seen[1].get("for_update"))
		self.assertEqual(state["status"], "ready")
		self.assertFalse(job.reload().cancel_requested)
		self.assertEqual(job.status, "ready")

	def test_cancel_still_flags_a_running_job(self):
		job = self.job(started_at=jobs.now_datetime())
		with patch.object(jobs, "dispatch"):
			jobs.cancel(job.name)
		self.assertTrue(job.reload().cancel_requested)

	def test_publish_falls_back_to_copy_across_devices(self):
		files = Path(frappe.get_site_path("private", "files"))
		spool = Path(frappe.get_site_path("private", "jarvis_exports"))
		real_link, real_copy, real_get = os.link, shutil.copyfile, frappe.db.get_value
		events = []

		def across_devices(source, target, *args, **kwargs):
			if Path(source).is_relative_to(spool):
				events.append("link-refused")
				raise OSError(errno.EXDEV, "Invalid cross-device link")
			return real_link(source, target, *args, **kwargs)

		def copy(source, target, *args, **kwargs):
			events.append("copy")
			return real_copy(source, target, *args, **kwargs)

		def get_value(*args, **kwargs):
			if kwargs.get("for_update"):
				events.append("lock")
			return real_get(*args, **kwargs)

		job = self.job()
		with (
			patch.object(jobs.os, "link", side_effect=across_devices),
			patch.object(jobs.shutil, "copyfile", side_effect=copy),
			patch.object(frappe.db, "get_value", side_effect=get_value),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		job.reload()
		self.assertEqual(job.status, "ready", job.error)
		self.assertEqual(events[:2], ["link-refused", "copy"])
		self.assertLess(events.index("copy"), events.index("lock"))
		with open(frappe.get_doc("File", job.file).get_full_path(), newline="") as source:
			self.assertEqual(len(list(csv.reader(source))), 3)
		self.assertEqual(list(files.glob("jarvis-export-tmp-*")), [])

	def test_public_state_carries_the_card_fields(self):
		before = jobs.now_datetime() - datetime.timedelta(seconds=2)
		with (
			patch.object(jobs, "current_conversation", return_value=self.conversation),
			patch.object(jobs, "dispatch"),
		):
			state = jobs.enqueue({"spec": {"from": "ToDo", "select": ["name", "owner"]}}, "csv")
		self.names.append(state["export_job"])
		fresh = jobs.get_jobs([state["export_job"]])[state["export_job"]]
		for shown in (state, fresh):
			self.assertEqual(shown["source_doctype"], "ToDo")
			self.assertEqual(shown["column_count"], 2)
			self.assertEqual((shown["attempts"], shown["retries"]), (0, 0))
			self.assertTrue(shown["asked_at"].endswith("+00:00"))
			asked = frappe.utils.get_datetime(shown["asked_at"].replace("+00:00", ""))
			self.assertGreaterEqual(asked, before)
			self.assertIn("last 10 minutes", shown["note"])
			self.assertNotIn("check its state", shown["note"])
			self.assertNotIn("\u2014", shown["note"])
		# Rows made before the fields existed stay readable: nothing raises, the card hides the parts.
		old = self.job(status="queued", attempt_token="")
		shown = jobs.public(old)
		self.assertIsNone(shown["asked_at"])
		self.assertEqual((shown["source_doctype"], shown["column_count"]), (None, 0))

	def test_public_shows_progress_only_while_running(self):
		running = self.job(total_rows=10, phase="finalizing")
		shown = jobs.get_jobs([running.name])[running.name]
		self.assertEqual((shown["total_rows"], shown["phase"]), (10, "finalizing"))
		# Meaningful only while running: the stale values on any other status stay hidden.
		for status in ("queued", "paused", "failed", "cancelled"):
			with self.subTest(status=status):
				frappe.db.set_value(jobs.DOCTYPE, running.name, "status", status)
				frappe.db.commit()
				shown = jobs.get_jobs([running.name])[running.name]
				self.assertEqual((shown["total_rows"], shown["phase"]), (None, None))
		ready = self.ready_job()
		self.stamp(ready, total_rows=10, phase="finalizing")
		shown = jobs.get_jobs([ready.name])[ready.name]
		self.assertEqual(shown["status"], "ready")
		self.assertEqual((shown["total_rows"], shown["phase"]), (None, None))
		unknown = self.job(total_rows=0)
		self.assertIsNone(jobs.get_jobs([unknown.name])[unknown.name]["total_rows"])

	def test_dispatch_resets_progress_fields(self):
		job = self.job(status="queued", attempt_token="", total_rows=500, phase="finalizing")
		with (
			patch("jarvis.chat.pump._probe_worker_count", return_value=3),
			patch.object(frappe, "enqueue"),
		):
			jobs.dispatch()
		job.reload()
		self.assertEqual(job.status, "running")
		self.assertEqual(job.total_rows, 0)
		self.assertFalse(job.phase)

	def test_set_progress_is_fenced(self):
		job = self.job()
		jobs._set_progress(job.name, "stale", phase="finalizing", total_rows=9)
		self.assertEqual(frappe.db.get_value(jobs.DOCTYPE, job.name, ["phase", "total_rows"]), ("", 0))
		for status in ("cancelled", "paused"):
			with self.subTest(status=status):
				frappe.db.set_value(jobs.DOCTYPE, job.name, "status", status)
				frappe.db.commit()
				jobs._set_progress(job.name, "abc123", phase="finalizing", total_rows=9)
				self.assertEqual(
					frappe.db.get_value(jobs.DOCTYPE, job.name, ["phase", "total_rows"]), ("", 0)
				)
		frappe.db.set_value(jobs.DOCTYPE, job.name, "status", "running")
		frappe.db.commit()
		jobs._set_progress(job.name, "abc123", phase="finalizing", total_rows=9)
		self.assertEqual(
			frappe.db.get_value(jobs.DOCTYPE, job.name, ["phase", "total_rows"]), ("finalizing", 9)
		)

	def test_background_title_defaults_to_source_doctype(self):
		with (
			patch.object(jobs, "current_conversation", return_value=self.conversation),
			patch.object(jobs, "dispatch"),
		):
			first = jobs.enqueue({"spec": {"from": "ToDo", "select": ["name"]}}, "csv")
			second = jobs.enqueue({"doctype": "DocType", "fields": ["name"]}, "csv")
			third = jobs.enqueue({"spec": {"from": "ToDo", "select": ["name", "owner"]}}, "csv", "Open tasks")
		self.names.extend(state["export_job"] for state in (first, second, third))
		self.assertEqual(first["title"], "ToDo")
		self.assertEqual(second["title"], "DocType")
		self.assertEqual(third["title"], "Open tasks")

	def test_failure_is_logged_with_the_job_before_finishing(self):
		job = self.job()
		with (
			patch.object(streaming, "write", side_effect=RuntimeError("disk full")),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		self.assertEqual(job.reload().status, "failed")
		logs = self.error_logs(job.name)
		self.assertEqual(len(logs), 1)
		self.assertTrue(logs[0].method.startswith("jarvis.export"), logs[0].method)
		self.assertIn("RuntimeError", logs[0].error)
		self.assertIn("disk full", logs[0].error)
		self.assertNotIn(job.source, logs[0].error)
		# The log survives a failing terminal write: it is committed before _finish runs.
		second = self.job()
		with (
			patch.object(streaming, "write", side_effect=RuntimeError("disk full again")),
			patch.object(jobs, "_finish", side_effect=RuntimeError("db gone")),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(second.name, "abc123")
		frappe.db.rollback()
		self.assertTrue(any("disk full again" in log.error for log in self.error_logs(second.name)))

	def test_refusals_shown_to_the_user_are_not_logged(self):
		# Permission and invalid-argument refusals are user outcomes, shown verbatim: no Error Log.
		for error in (
			PermissionDeniedError("No read permission on ToDo."),
			InvalidArgumentError("Bad field."),
		):
			with self.subTest(error=type(error).__name__):
				job = self.job()
				with (
					patch.object(streaming, "write", side_effect=error),
					patch.object(jobs, "dispatch"),
				):
					jobs.run(job.name, "abc123")
				self.assertEqual(job.reload().status, "failed")
				self.assertEqual(job.error, str(error))
				self.assertEqual(self.error_logs(job.name), [])

	def test_failed_outcomes_are_logged_but_zero_rows_is_not(self):
		user = frappe.get_doc(
			{
				"doctype": "User",
				"email": f"export-off-{frappe.generate_hash(length=8)}@example.com",
				"first_name": "off",
				"send_welcome_email": 0,
				"enabled": 0,
				"user_type": "System User",
			}
		).insert(ignore_permissions=True)
		try:
			inactive = self.job()
			frappe.db.set_value(jobs.DOCTYPE, inactive.name, "owner", user.name)
			empty = self.job(
				source=json.dumps(
					{
						"spec": {
							"from": "ToDo",
							"select": ["name"],
							"where": [{"field": "description", "op": "=", "value": "jv-no-such-todo-xyz"}],
						}
					}
				)
			)
			frappe.db.commit()
			with patch.object(jobs, "dispatch"):
				jobs.run(inactive.name, "abc123")
				jobs.run(empty.name, "abc123")
			self.assertEqual(inactive.reload().status, "failed")
			# (The chat belongs to someone else here, so the bell step may log as well.)
			self.assertTrue(any("active account" in log.error for log in self.error_logs(inactive.name)))
			self.assertEqual(empty.reload().status, "failed")
			self.assertIn("No records match", empty.error)
			self.assertEqual(self.error_logs(empty.name), [])
		finally:
			frappe.delete_doc("User", user.name, force=True, ignore_permissions=True)
			frappe.db.commit()

	def test_enqueue_failure_returns_the_reservation(self):
		first = self.job(status="queued", attempt_token="")
		second = self.job(status="queued", attempt_token="")
		# A different owner, so only the break after the failure keeps dispatch from trying it.
		frappe.db.set_value(jobs.DOCTYPE, second.name, "owner", "Guest")
		frappe.db.commit()
		with (
			patch("jarvis.chat.pump._probe_worker_count", return_value=3),
			patch("frappe.enqueue", side_effect=RuntimeError("redis down")) as enqueue,
		):
			jobs.dispatch()
		self.assertEqual(enqueue.call_count, 1)
		for job in (first, second):
			job.reload()
			self.assertEqual(job.status, "queued")
			self.assertEqual(job.attempt_token or "", "")
			self.assertEqual(job.retries or 0, 0)
		self.assertEqual(len(self.error_logs(first.name)), 1)

	def test_run_tail_survives_a_stuck_watchdog(self):
		job = self.job()

		def stuck(watchdog):
			watchdog.stop_event.set()
			raise RuntimeError("Export watchdog did not stop")

		with (
			patch.object(Watchdog, "stop", autospec=True, side_effect=stuck),
			patch.object(jobs, "announce") as announce,
			patch.object(jobs, "dispatch") as dispatch,
		):
			jobs.run(job.name, "abc123")
		announce.assert_called_once()
		dispatch.assert_called_once()

	def test_publish_phase_refreshes_the_heartbeat(self):
		# The watchdog stops before publishing; a slow copy or lock wait must not look like a dead worker.
		job = self.job()
		stale = jobs.now_datetime() - datetime.timedelta(hours=2)
		seen = {}
		real_stop = Watchdog.stop

		def stop_then_age_the_heartbeat(watchdog):
			real_stop(watchdog)
			frappe.db.set_value(jobs.DOCTYPE, job.name, "heartbeat", stale, update_modified=False)
			frappe.db.commit()

		def publish(row, token, *args):
			seen["heartbeat"] = frappe.db.get_value(jobs.DOCTYPE, row.name, "heartbeat")

		with (
			patch.object(Watchdog, "stop", autospec=True, side_effect=stop_then_age_the_heartbeat),
			patch.object(jobs, "_publish_file", side_effect=publish),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(job.name, "abc123")
		self.assertGreater(
			frappe.utils.get_datetime(seen["heartbeat"]), jobs.now_datetime() - datetime.timedelta(minutes=1)
		)

		# Fenced: a stale attempt's heartbeat touches nothing.
		frappe.db.set_value(jobs.DOCTYPE, job.name, "heartbeat", stale, update_modified=False)
		frappe.db.commit()
		jobs._heartbeat(job.name, "not-the-token")
		self.assertLess(
			frappe.utils.get_datetime(job.reload().heartbeat),
			jobs.now_datetime() - datetime.timedelta(hours=1),
		)

	def test_watchdog_keeps_the_original_cause(self):
		job = self.job(started_at=jobs.now_datetime())
		lost = pymysql.err.OperationalError(2013, "Lost connection to server")
		db = MariaDBStream(connection_settings()).open()
		watchdog = Watchdog(db, job.name, "abc123", frappe.local.site, job.owner, job.conversation)
		try:
			with patch.object(db, "detect_lock_waiters", side_effect=lost):
				watchdog.start()
				deadline = time.monotonic() + 8
				while not watchdog.reason and time.monotonic() < deadline:
					time.sleep(0.1)
		finally:
			watchdog.stop()
			db.close()
		self.assertEqual(watchdog.reason.status, "failed")
		self.assertIs(watchdog.cause, lost)

		# run() hands that cause to the log; re-raising the replacement reason erases __cause__.
		second = self.job()

		def wait_for_the_watchdog(prepared, rows, path, fmt, cap, size_cap, check, progress, **kwargs):
			for _ in range(100):
				check()
				time.sleep(0.1)
			raise AssertionError("the watchdog never fired")

		with (
			patch.object(MariaDBStream, "detect_lock_waiters", side_effect=lost),
			patch.object(streaming, "write", wait_for_the_watchdog),
			patch.object(jobs, "dispatch"),
		):
			jobs.run(second.name, "abc123")
		self.assertEqual(second.reload().status, "failed")
		self.assertTrue(any("OperationalError" in log.error for log in self.error_logs(second.name)))

	def test_run_tail_errors_do_not_skip_dispatch(self):
		job = self.job()
		dispatch = MagicMock()
		with (
			patch.object(jobs, "announce", side_effect=RuntimeError("announce failed")),
			patch.object(jobs, "dispatch", dispatch),
		):
			jobs.run(job.name, "abc123")
		dispatch.assert_called_once()
		self.assertEqual(job.reload().status, "ready")

	def test_reaper_tail_is_isolated_and_still_dispatches(self):
		job = self.job()
		with patch.object(jobs, "dispatch"):
			jobs.run(job.name, "abc123")
		file_name = job.reload().file
		frappe.db.delete(jobs.DOCTYPE, {"name": job.name})
		frappe.db.commit()
		try:
			with (
				patch.object(jobs, "dispatch") as dispatch,
				patch.object(frappe, "delete_doc", side_effect=RuntimeError("cannot delete")),
			):
				jobs.reap()
			dispatch.assert_called_once()
			self.assertEqual(len(self.error_logs(file_name, "File")), 1)
		finally:
			jobs._delete_generated_file(file_name)
			frappe.db.delete("Error Log", {"reference_doctype": "File", "reference_name": file_name})
			frappe.db.commit()

	def test_reaper_purges_old_rows_even_if_announce_fails(self):
		job = self.job(status="failed", attempt_token="", announced_status="")
		frappe.db.set_value(
			jobs.DOCTYPE,
			job.name,
			"creation",
			frappe.utils.add_to_date(jobs.now_datetime(), days=-31),
			update_modified=False,
		)
		frappe.db.commit()
		with (
			patch.object(jobs, "dispatch"),
			patch.object(jobs, "announce", side_effect=RuntimeError("owner is gone")),
		):
			jobs.reap()
		self.assertFalse(frappe.db.exists(jobs.DOCTYPE, job.name))

	def test_same_request_in_another_chat_starts_its_own_export(self):
		from jarvis.tools.export_query import export_query

		other = frappe.get_doc({"doctype": "Jarvis Conversation", "title": "second export chat"}).insert()
		spec = {"from": "DocType", "select": ["name"]}
		try:
			with patch.object(jobs, "dispatch"):
				with patch.object(jobs, "current_conversation", return_value=self.conversation):
					first = jobs.enqueue({"spec": spec}, "csv")
				with (
					patch.object(resolvers, "_HARD_ROW_CEILING", 2),
					patch.object(jobs, "current_conversation", return_value=other.name),
				):
					second = export_query(spec=spec, format="csv")
			a, b = first["export_job"], second["export_job"]
			self.names.extend([a, b])
			self.assertNotEqual(a, b)
			self.assertEqual(frappe.db.get_value(jobs.DOCTYPE, b, "conversation"), other.name)
			for job_name, chat in ((a, self.conversation), (b, other.name)):
				filters = {"conversation": chat, "tool_call_id": f"export:{job_name}:card"}
				self.assertEqual(frappe.db.count("Jarvis Chat Message", filters), 1)
			self.assertEqual(
				frappe.db.count(
					"Jarvis Chat Message", {"conversation": other.name, "tool_call_id": f"export:{a}:card"}
				),
				0,
			)
			# One export per user at a time: the second queues behind the first.
			frappe.db.set_value(jobs.DOCTYPE, a, {"status": "running", "attempt_token": "abc123"})
			frappe.db.commit()
			with (
				patch("jarvis.chat.pump._probe_worker_count", return_value=3),
				patch.object(frappe, "enqueue") as enqueue,
			):
				jobs.dispatch()
			self.assertEqual(frappe.db.get_value(jobs.DOCTYPE, b, "status"), "queued")
			self.assertEqual([c for c in enqueue.call_args_list if c.kwargs.get("name") == b], [])
			# Deleting the second chat deletes its job and file and leaves the first job alone.
			frappe.db.set_value(jobs.DOCTYPE, b, {"status": "running", "attempt_token": "abc123"})
			frappe.db.commit()
			snapshot = frappe.get_doc(jobs.DOCTYPE, b)
			with tempfile.TemporaryDirectory() as directory:
				path = Path(directory) / "export.part"
				path.write_text("Name\nToDo\n")
				jobs._publish_file(snapshot, "abc123", path, {"rows_written": 1}, "permission")
			file_name = frappe.db.get_value(jobs.DOCTYPE, b, "file")
			self.assertTrue(file_name)
			frappe.delete_doc("Jarvis Conversation", other.name, force=True, ignore_permissions=True)
			frappe.db.commit()
			self.assertFalse(frappe.db.exists(jobs.DOCTYPE, b))
			self.assertFalse(frappe.db.exists("File", file_name))
			self.assertTrue(frappe.db.exists(jobs.DOCTYPE, a))
		finally:
			frappe.db.delete("Jarvis Chat Message", {"conversation": other.name})
			frappe.db.delete("Jarvis Conversation", {"name": other.name})
			frappe.db.commit()

	def test_state_messages_post_once_to_the_jobs_own_chat(self):
		job = self.job()
		key = f"export:{job.name}:ready"

		def ready_messages():
			return frappe.get_all(
				"Jarvis Chat Message",
				filters={"conversation": self.conversation, "tool_call_id": key},
				fields=["content"],
			)

		def bells():
			return frappe.get_all("Notification Log", filters={"document_name": job.name}, fields=["subject"])

		with patch.object(jobs, "dispatch"):
			jobs.announce(job.name, initial=True)
			self.assertEqual(
				frappe.db.get_value(
					"Jarvis Chat Message",
					{"conversation": self.conversation, "tool_call_id": f"export:{job.name}:card"},
					"content",
				),
				"Export started: Export test.",
			)
			jobs.run(job.name, "abc123")
			jobs.announce(job.name)
			jobs.announce(job.name)
			messages = ready_messages()
			self.assertEqual(len(messages), 1)
			self.assertIn("Export test", messages[0].content)
			self.assertIn("rows", messages[0].content)
			self.assertEqual([b.subject for b in bells()], [messages[0].content])
			self.assertEqual(job.reload().announced_status, "ready")
			# The reaper re-announces what it believes undelivered: nothing is posted twice, and
			# the marker is set again even though the keyed message already existed.
			frappe.db.set_value(jobs.DOCTYPE, job.name, "announced_status", "")
			frappe.db.commit()
			jobs.reap()
		self.assertEqual(len(ready_messages()), 1)
		self.assertEqual(len(bells()), 1)
		self.assertEqual(job.reload().announced_status, "ready")

	def test_bell_subject_escapes_the_model_supplied_title(self):
		job = self.job(title="<b>Q3</b> suppliers")
		with patch.object(jobs, "dispatch"):
			jobs.run(job.name, "abc123")
		subjects = [
			b.subject
			for b in frappe.get_all(
				"Notification Log", filters={"document_name": job.name}, fields=["subject"]
			)
		]
		self.assertEqual(len(subjects), 1)
		self.assertNotIn("<b>", subjects[0])
		self.assertIn("&lt;b&gt;Q3", subjects[0])

	def test_publish_refuses_when_the_chat_was_deleted(self):
		job = self.job()
		files = Path(frappe.get_site_path("private", "files"))
		before = set(files.glob("jarvis-export-*"))
		# Straight to the table: the job row must survive to show what happened to it.
		frappe.db.delete("Jarvis Conversation", {"name": self.conversation})
		frappe.db.commit()
		with tempfile.TemporaryDirectory() as directory:
			path = Path(directory) / "export.part"
			path.write_text("Name\nToDo\n")
			jobs._publish_file(job, "abc123", path, {"rows_written": 1}, "permission")
		job.reload()
		self.assertEqual(job.status, "cancelled")
		self.assertIn("conversation was deleted", job.error)
		self.assertFalse(job.file)
		self.assertFalse(
			frappe.db.exists("File", {"attached_to_doctype": jobs.DOCTYPE, "attached_to_name": job.name})
		)
		self.assertEqual(set(files.glob("jarvis-export-*")), before)

	def test_conversation_trash_before_migrate_is_a_no_op(self):
		doc = frappe.get_doc("Jarvis Conversation", self.conversation)
		with (
			patch.object(frappe.db, "table_exists", return_value=False),
			patch.object(frappe, "get_all", side_effect=AssertionError("queried a missing table")),
		):
			jobs.on_conversation_trash(doc)


class TestBackgroundFilterPermissions(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		from jarvis.tests.test_query import TestQueryPermlevelFieldACL

		cls.fixture = TestQueryPermlevelFieldACL
		cls.fixture.setUpClass()

	def test_child_rows_obey_parent_record_scope_and_parenttype(self):
		f = self.fixture
		for parent, role, user, expected in (
			(f.SCOPE_PARENT_A, f.ROLE_SCOPE_A, f.USER_SCOPE_A, [10, 20, 30]),
			(f.SCOPE_SINGLE, f.ROLE_SCOPE_SINGLE, f.USER_SCOPE_SINGLE, [f.SINGLE_VAL]),
		):
			frappe.set_user("Administrator")
			perm = frappe.db.get_value(
				"DocPerm", {"parent": parent, "role": role}, ["name", "export"], as_dict=True
			)
			try:
				frappe.db.set_value("DocPerm", perm.name, "export", 1)
				frappe.db.commit()
				frappe.clear_cache(doctype=parent)
				frappe.set_user(user)
				p = streaming.prepare(
					{"doctype": f.SCOPE_CHILD_DT, "parent_doctype": parent, "fields": ["scope_val"]}, 200_000
				)
				db = MariaDBStream(connection_settings()).open()
				try:
					self.assertEqual(
						sorted(r["scope_val"] for r in db.stream_rows(p.sql, p.params)), expected
					)
				finally:
					db.close()
			finally:
				frappe.set_user("Administrator")
				frappe.db.set_value("DocPerm", perm.name, "export", perm.export)
				frappe.db.commit()
				frappe.clear_cache(doctype=parent)


class TestPreparedSqlFetchesOnePastTheCap(FrappeTestCase):
	"""A match larger than the cap is told apart from one that exactly fits by one extra row."""

	CAP = 2

	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")

	def assert_fetches_cap_plus_one(self, sql):
		# Frappe 15's legacy compiler writes lowercase "limit 3"; Frappe 16's builder "LIMIT 3".
		self.assertRegex(sql, re.compile(rf"\blimit\s+{self.CAP + 1}\b", re.I))

	def test_spec_sql_fetches_one_row_past_the_cap(self):
		p = streaming.prepare({"spec": {"from": "DocType", "select": ["name"]}}, self.CAP)
		self.assert_fetches_cap_plus_one(p.sql)

	def test_filter_sql_fetches_one_row_past_the_cap(self):
		p = streaming.prepare({"doctype": "DocType", "fields": ["name"]}, self.CAP)
		self.assert_fetches_cap_plus_one(p.sql)


class TestPreparedSheetName(FrappeTestCase):
	"""The background workbook's tab is named after the export, as the immediate one is."""

	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")

	def test_sheet_name_comes_from_the_title(self):
		source = {"doctype": "ToDo", "fields": ["name"]}
		self.assertEqual(streaming.prepare(source, 10).sheet_name, "Export")
		self.assertEqual(streaming.prepare(source, 10, title="Q1/Q2: [x]").sheet_name, "Q1 Q2   x ")

	def test_sheet_name_never_starts_or_ends_with_an_apostrophe(self):
		# Excel rejects such a tab name; xlsxwriter raises InvalidWorksheetName for it.
		source = {"doctype": "ToDo", "fields": ["name"]}
		cases = (("'Q3' suppliers", "Q3' suppliers"), ("Suppliers'", "Suppliers"), ("'" * 3, "Export"))
		for title, expected in cases:
			with self.subTest(title=title):
				self.assertEqual(streaming.sheet_name(title), expected)
				self.assertEqual(streaming.prepare(source, 10, title=title).sheet_name, expected)


class TestExportGateParity(FrappeTestCase):
	"""The immediate and the background export must refuse the same input in the same words.

	Each path is asserted on its own, against the text the immediate path gives today, so
	deleting a check from one path cannot hide behind the other (a deleted DocType check makes
	both fail later and alike, which an equality test would accept). Both paths run the same
	gate (``resolvers.authorize_spec`` / ``authorize_filter``)."""

	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")

	def immediate(self, **source):
		return resolvers.from_query(**source)

	def background(self, **source):
		return streaming.prepare(source, 10)

	def refused(self, call, error, text):
		with self.assertRaisesRegex(error, text):
			call()

	# Unknown DocType.
	def test_unknown_doctype_immediate(self):
		self.refused(
			lambda: self.immediate(doctype="JV No Such", fields=["name"]),
			InvalidArgumentError,
			"unknown DocType: JV No Such",
		)

	def test_unknown_doctype_background(self):
		self.refused(
			lambda: self.background(doctype="JV No Such", fields=["name"]),
			InvalidArgumentError,
			"unknown DocType: JV No Such",
		)

	# fields=["*"].
	def test_star_fields_immediate(self):
		self.refused(
			lambda: self.immediate(doctype="ToDo", fields=["*"]),
			InvalidArgumentError,
			"name the columns you want",
		)

	def test_star_fields_background(self):
		self.refused(
			lambda: self.background(doctype="ToDo", fields=["*"]),
			InvalidArgumentError,
			"name the columns you want",
		)

	# A child table without its parent: the text names a parent to use.
	def test_child_table_without_parent_immediate(self):
		self.refused(
			lambda: self.immediate(doctype="DocField", fields=["fieldname"]),
			InvalidArgumentError,
			"is a child table; pass parent_doctype",
		)

	def test_child_table_without_parent_background(self):
		self.refused(
			lambda: self.background(doctype="DocField", fields=["fieldname"]),
			InvalidArgumentError,
			"is a child table; pass parent_doctype",
		)

	# Permission and masking refusals.
	def test_no_export_permission_on_both_paths(self):
		with patch("frappe.permissions.can_export", return_value=False):
			for call in (
				lambda: self.immediate(doctype="ToDo", fields=["name"]),
				lambda: self.background(doctype="ToDo", fields=["name"]),
			):
				self.refused(call, PermissionDeniedError, "no export permission on ToDo")

	def test_child_export_checks_the_parent_on_both_paths(self):
		# The Export permission of a child table lives on its parent.
		def only_user_is_denied(doctype, *args, **kwargs):
			return doctype != "User"

		with patch("frappe.permissions.can_export", side_effect=only_user_is_denied):
			for call in (
				lambda: self.immediate(doctype="Has Role", parent_doctype="User", fields=["role"]),
				lambda: self.background(doctype="Has Role", parent_doctype="User", fields=["role"]),
			):
				self.refused(call, PermissionDeniedError, "no export permission on User")

	def test_no_read_permission_on_both_paths(self):
		with patch.object(frappe, "has_permission", return_value=False):
			for call in (
				lambda: self.immediate(doctype="ToDo", fields=["name"]),
				lambda: self.background(doctype="ToDo", fields=["name"]),
			):
				self.refused(call, PermissionDeniedError, "no read permission on ToDo")

	def test_spec_without_export_permission_on_both_paths(self):
		spec = {"from": "ToDo", "select": ["name"]}
		with patch("frappe.permissions.can_export", return_value=False):
			self.refused(
				lambda: resolvers.from_spec(spec), PermissionDeniedError, "no export permission on ToDo"
			)
			self.refused(
				lambda: streaming.prepare({"spec": spec}, 10),
				PermissionDeniedError,
				"no export permission on ToDo",
			)

	def test_computed_column_over_a_masked_field_on_both_paths(self):
		# Patched, so it runs on Frappe 15 (no masking there) as well as 16.
		spec = {
			"from": "ToDo",
			"alias": "m",
			"select": [{"expr": {"fn": "upper", "args": [{"field": "m.description"}]}, "as": "shout"}],
		}
		masked = {"m": {"description": object()}}
		with patch.object(masking, "masked_fields", return_value=masked):
			self.refused(
				lambda: resolvers.from_spec(spec), InvalidArgumentError, r"computes over m\.description"
			)
			self.refused(
				lambda: streaming.prepare({"spec": spec}, 10),
				InvalidArgumentError,
				r"computes over m\.description",
			)
