"""Inspect real workbook chart parts and cell references, without saving Files."""

import datetime
import io
import unittest
import zipfile
from decimal import Decimal
from unittest.mock import patch
from xml.etree import ElementTree as ET

import frappe
from openpyxl import load_workbook

from jarvis.exceptions import InvalidArgumentError
from jarvis.tools.export_excel import export_excel

NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
ROWS = [{"Month": "Jan", "Sales": 100, "Costs": 40}, {"Month": "Feb", "Sales": 200, "Costs": 80}]
CHART = {"type": "column", "categories": "Month", "values": ["Sales"], "title": "Monthly Sales"}


class TestExcelCharts(unittest.TestCase):
	def setUp(self):
		p = patch("frappe.utils.file_manager.save_file")
		self.save = p.start()
		self.addCleanup(p.stop)
		self.save.return_value = frappe._dict(
			file_url="/private/files/chart.xlsx", file_name="chart.xlsx", file_size=0, name="MOCK"
		)

	def workbook(self):
		return load_workbook(io.BytesIO(self.save.call_args.args[1]))

	def chart_xml(self, index=1):
		with zipfile.ZipFile(io.BytesIO(self.save.call_args.args[1])) as z:
			return ET.fromstring(z.read(f"xl/charts/chart{index}.xml"))

	def test_real_native_chart_types_and_exact_ranges(self):
		for kind, element in [
			("column", "barChart"),
			("bar", "barChart"),
			("line", "lineChart"),
			("pie", "pieChart"),
		]:
			with self.subTest(kind=kind):
				result = export_excel(rows=ROWS, title="Sales", charts=[{**CHART, "type": kind}])
				self.assertEqual(result["chart_count"], 1)
				self.assertEqual(self.save.call_args.kwargs["is_private"], 1)
				wb = self.workbook()
				self.assertEqual(len(wb["Sales"]._charts), 1)
				self.assertEqual(
					list(wb["Sales"].values),
					[("Month", "Sales", "Costs"), ("Jan", 100, 40), ("Feb", 200, 80)],
				)
				root = self.chart_xml()
				self.assertIsNotNone(root.find(f".//c:{element}", NS))
				refs = [node.text.replace("'", "") for node in root.findall(".//c:f", NS)]
				self.assertIn("Sales!$A$2:$A$3", refs)
				self.assertIn("Sales!$B$2:$B$3", refs)
				self.assertIn("Sales!B1", [r.replace("$", "") for r in refs])
				self.assertIn("Monthly Sales", ET.tostring(root, encoding="unicode"))
				if kind in ("column", "bar"):
					self.assertEqual(
						root.find(".//c:barDir", NS).attrib["val"], "col" if kind == "column" else "bar"
					)

	def test_multiple_series_charts_and_correct_sheet(self):
		out = export_excel(
			sheets=[
				{"title": "Empty", "rows": []},
				{
					"title": "Sales",
					"rows": ROWS,
					"charts": [CHART, {**CHART, "type": "line", "values": ["Sales", "Costs"]}],
				},
				{"title": "Raw", "rows": ROWS},
			]
		)
		self.assertEqual(out["chart_count"], 2)
		wb = self.workbook()
		self.assertEqual(wb.sheetnames, ["Sales", "Raw"])
		self.assertEqual([len(ws._charts) for ws in wb], [2, 0])
		self.assertEqual(len(wb["Sales"]._charts[1].series), 2)
		self.assertEqual([c.anchor._from.row for c in wb["Sales"]._charts], [1, 21])
		self.assertEqual([c.anchor._from.col for c in wb["Sales"]._charts], [4, 4])

	def test_duplicate_sheet_titles_and_sanitized_reference(self):
		export_excel(
			sheets=[
				{"title": "Kavin's: Sales", "rows": ROWS, "charts": [CHART]},
				{"title": "Kavin's: Sales", "rows": ROWS, "charts": [CHART]},
			]
		)
		wb = self.workbook()
		for i, ws in enumerate(wb, 1):
			refs = [n.text for n in self.chart_xml(i).findall(".//c:f", NS)]
			quoted = "'" + ws.title.replace("'", "''") + "'"
			self.assertIn(quoted + "!$A$2:$A$3", refs)

	def test_selected_columns_and_list_rows(self):
		for rows in (ROWS, [[100, "Jan"], [200, "Feb"]]):
			with self.subTest(rows=rows):
				export_excel(rows=rows, columns=["Sales", "Month"], charts=[CHART])
				refs = [n.text for n in self.chart_xml().findall(".//c:f", NS)]
				self.assertTrue(any(r.endswith("!$B$2:$B$3") for r in refs))
				self.assertTrue(any(r.endswith("!$A$2:$A$3") for r in refs))
		export_excel(rows=[["Month", "Sales"], ["Jan", 0], ["Feb", -4]], charts=[CHART])
		self.assertEqual(self.workbook().active["B2"].value, 0)

	def test_blanks_and_decimal_values(self):
		rows = [["Month", "Sales"], ["Jan", Decimal("1.25")], ["Feb", None], ["Mar", " "], ["Apr", 0]]
		out = export_excel(rows=rows, charts=[CHART])
		self.assertEqual(out["chart_count"], 1)
		self.assertEqual(self.workbook().active["B2"].value, 1.25)

	def test_date_categories_export_and_cached_values_match(self):
		export_excel(rows=[{"Month": datetime.date(2026, 1, 1), "Sales": 12}], charts=[CHART])
		self.assertEqual(len(self.workbook().active._charts), 1)
		if not frappe.__version__.startswith("15."):
			values = self.chart_xml().findall(".//c:val/c:numRef/c:numCache/c:pt/c:v", NS)
			self.assertEqual([v.text for v in values], ["12"])

	def test_boolean_categories_have_valid_label_caches(self):
		export_excel(rows=[{"Month": True, "Sales": 12}, {"Month": False, "Sales": 5}], charts=[CHART])
		self.assertEqual(len(self.workbook().active._charts), 1)
		if not frappe.__version__.startswith("15."):
			values = self.chart_xml().findall(".//c:cat/c:strRef/c:strCache/c:pt/c:v", NS)
			self.assertEqual([v.text for v in values], ["TRUE", "FALSE"])

	def test_bad_chart_inputs_do_not_save_a_file(self):
		for charts in [
			{},
			"chart",
			[None],
			[{}],
			[CHART] * 5,
			[{**CHART, "type": "scatter"}],
			[{**CHART, "categories": "Missing"}],
			[{**CHART, "categories": []}],
			[{**CHART, "values": []}],
			[{**CHART, "values": ["Sales"] * 7}],
			[{**CHART, "values": ["Sales", "Sales"]}],
			[{**CHART, "values": [None]}],
			[{**CHART, "values": "Sales"}],
			[{**CHART, "title": 42}],
			[{**CHART, "title": None}],
			[{**CHART, "title": "x" * 121}],
			[{**CHART, "title": "=[external.xlsx]Sheet1!A1"}],
			[{**CHART, "anchor": "Z9"}],
			[{**CHART, "type": "pie", "values": ["Sales", "Costs"]}],
		]:
			with self.subTest(charts=charts), self.assertRaises(InvalidArgumentError):
				export_excel(rows=ROWS, charts=charts)
		self.save.assert_not_called()

	def test_invalid_numeric_series_and_ambiguous_headers(self):
		for value in [True, "100", "=SUM(A1:A3)", float("nan"), float("inf"), Decimal("NaN"), "", None]:
			with self.subTest(value=value), self.assertRaises(InvalidArgumentError):
				export_excel(rows=[{"Month": "Jan", "Sales": value}], charts=[CHART])
		with self.assertRaises(InvalidArgumentError):
			export_excel(rows=[["Month", "Sales", "Sales"], ["Jan", 1, 2]], charts=[CHART])
		self.save.assert_not_called()

	def test_charts_with_sheets_or_empty_charted_sheet_are_rejected(self):
		for charts in [[], [CHART]]:
			with self.assertRaises(InvalidArgumentError):
				export_excel(sheets=[{"rows": ROWS}], charts=charts)
		with self.assertRaises(InvalidArgumentError):
			export_excel(sheets=[{"rows": ROWS}, {"rows": [], "charts": [CHART]}])
		with self.assertRaises(InvalidArgumentError):
			export_excel(sheets=[{"rows": ROWS}, {"rows": [], "charts": {}}])
		self.save.assert_not_called()

	def test_invalid_later_sheet_prevents_any_save(self):
		with self.assertRaises(InvalidArgumentError):
			export_excel(
				sheets=[
					{"rows": ROWS, "charts": [CHART]},
					{"rows": ROWS, "charts": [{**CHART, "values": ["no"]}]},
				]
			)
		self.save.assert_not_called()

	def test_plain_exports_still_have_no_charts(self):
		result = export_excel(rows=ROWS)
		self.assertEqual(result["chart_count"], 0)
		self.assertEqual(len(self.workbook().active._charts), 0)
		with zipfile.ZipFile(io.BytesIO(self.save.call_args.args[1])) as z:
			self.assertFalse(any(n.startswith("xl/charts/") for n in z.namelist()))

	@unittest.skipIf(
		frappe.__version__.startswith("15."), "develop registry has an unrelated v16-only import"
	)
	def test_registry_dispatch_preserves_charts(self):
		from jarvis.tools.registry import dispatch

		out = dispatch("export_excel", {"rows": ROWS, "charts": [CHART]})
		self.assertEqual(out["chart_count"], 1)
		self.assertEqual(len(self.workbook().active._charts), 1)
