"""export_query with a ``query`` spec: records a plain filter cannot pick.

A customer asked for "suppliers with no transactions in the last year" (2,136 rows)
as Excel. Only ``query`` can say NOT EXISTS, its result is cut to fit the chat, and
export_query took get_list filters only, so the agent tried to rebuild the file in
100-row batches through chat and stopped after one. ``spec`` exports the same query
server-side, complete, under the same permission gates.
"""

from __future__ import annotations

import contextlib
import importlib
import io
from unittest import mock

import frappe
from frappe.tests.utils import FrappeTestCase
from openpyxl import load_workbook

from jarvis.exceptions import ExportTooSlowError, InvalidArgumentError, NoDataError, PermissionDeniedError
from jarvis.tools._export import resolvers, save_export_file
from jarvis.tools.export_query import export_query

# The module, not the function of the same name that the import above binds.
export_module = importlib.import_module("jarvis.tools.export_query")

PFX = "xqs-"
TOTAL = 230  # past query's 200-row guard and the ~100-row chat cut
WITH_COMMENT = 30


def _no_comment_spec(**extra) -> dict:
	"""ToDos of this suite that nobody commented on: the NOT EXISTS shape."""
	return {
		"from": "ToDo",
		"alias": "t",
		"select": ["t.name", "t.description", "t.status"],
		"where": [
			{"field": "t.description", "op": "like", "value": f"{PFX}%"},
			{
				"op": "not exists",
				"value": {
					"from": "Comment",
					"alias": "c",
					"where": [
						{"field": "c.reference_doctype", "op": "=", "value": "ToDo"},
						{"field": "c.reference_name", "op": "=", "value": {"$field": "t.name"}},
					],
				},
			},
		],
		"order_by": [{"field": "t.description", "dir": "asc"}],
		**extra,
	}


def _sheet(env: dict) -> list[tuple]:
	content = frappe.get_doc("File", env["name"]).get_content()
	wb = load_workbook(io.BytesIO(content if isinstance(content, bytes) else content.encode()))
	return list(wb.worksheets[0].iter_rows(values_only=True))


class _Clock:
	"""Stands in for ``time`` in export_query: it only moves when a test says so."""

	def __init__(self):
		self.now = 1000.0

	def monotonic(self) -> float:
		return self.now

	def after(self, seconds: float, real):
		"""``real`` wrapped so the clock has moved ``seconds`` once it has run."""

		def run(*args, **kwargs):
			result = real(*args, **kwargs)
			self.now += seconds
			return result

		return run


def _seed() -> None:
	frappe.set_user("Administrator")
	_purge()
	for i in range(TOTAL):
		todo = frappe.get_doc({"doctype": "ToDo", "description": f"{PFX}{i:04d}"}).insert()
		if i < WITH_COMMENT:
			todo.add_comment("Comment", text="followed up")
	frappe.db.commit()


def _purge() -> None:
	names = frappe.get_all("ToDo", filters={"description": ["like", f"{PFX}%"]}, pluck="name")
	frappe.db.delete("Comment", {"reference_doctype": "ToDo", "reference_name": ["in", names or [""]]})
	frappe.db.delete("ToDo", {"description": ["like", f"{PFX}%"]})
	frappe.db.commit()


class TestExportQuerySpec(FrappeTestCase):
	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		_seed()

	@classmethod
	def tearDownClass(cls):
		_purge()
		super().tearDownClass()

	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")

	def test_a_not_exists_query_exports_every_row(self):
		env = export_query(spec=_no_comment_spec(), title="Todos nobody commented on")
		self.assertEqual(env["total"], TOTAL - WITH_COMMENT)
		rows = _sheet(env)
		self.assertEqual(rows[0], ("ID", "Description", "Status"))
		self.assertEqual(len(rows) - 1, TOTAL - WITH_COMMENT)
		self.assertEqual(rows[1][1], f"{PFX}{WITH_COMMENT:04d}")
		self.assertEqual(env["title"], "Todos nobody commented on")

	def test_csv_and_aggregates_with_named_columns(self):
		spec = {
			"from": "ToDo",
			"alias": "t",
			"select": ["t.status", {"agg": "count", "field": "t.name", "as": "todo_count"}],
			"where": [{"field": "t.description", "op": "like", "value": f"{PFX}%"}],
			"group_by": ["t.status"],
		}
		env = export_query(spec=spec, format="csv")
		content = frappe.get_doc("File", env["name"]).get_content()
		text = content.decode() if isinstance(content, bytes) else content
		self.assertTrue(text.startswith("Status,Todo Count"), text[:40])
		self.assertIn(f"Open,{TOTAL}", text)

	def test_a_small_deliberate_limit_still_applies_and_says_so(self):
		env = export_query(spec=_no_comment_spec(limit=50))
		self.assertEqual(env["total"], 50)
		# The file holds the 50 rows asked for, not the one fetched past them to find out.
		self.assertEqual(len(_sheet(env)) - 1, 50)
		self.assertEqual(env["limited_by_spec"], 50)
		self.assertIn("first 50 rows", env["note"])
		self.assertIn("without limit", env["note"])

	def test_above_the_ceiling_requires_a_chat_for_delivery(self):
		with mock.patch.object(resolvers, "_HARD_ROW_CEILING", 10):
			with self.assertRaises(InvalidArgumentError) as ctx:
				export_query(spec=_no_comment_spec())
		self.assertIn("Narrow", str(ctx.exception))
		self.assertIn("Start it from a Jarvis chat", str(ctx.exception))

	def test_shapes_an_export_refuses(self):
		with self.assertRaisesRegex(InvalidArgumentError, "drop offset"):
			export_query(spec=_no_comment_spec(offset=10))
		with self.assertRaisesRegex(InvalidArgumentError, "'as' name"):
			export_query(spec={"from": "ToDo", "select": [{"agg": "count", "field": "name"}]})
		dup = {
			"from": "ToDo",
			"alias": "t",
			"joins": [{"doctype": "User", "alias": "u", "on": {"u.name": "t.allocated_to"}, "type": "left"}],
			"select": ["t.name", "u.name"],
		}
		with self.assertRaisesRegex(InvalidArgumentError, "both named 'name'"):
			export_query(spec=dup)
		with self.assertRaisesRegex(InvalidArgumentError, "exactly one of"):
			export_query("ToDo", spec=_no_comment_spec())
		with self.assertRaisesRegex(InvalidArgumentError, "exactly one of"):
			export_query()
		with self.assertRaisesRegex(InvalidArgumentError, "in the spec itself"):
			export_query(spec=_no_comment_spec(), filters={"status": "Open"})

	def test_highlights_may_name_the_field_or_its_label(self):
		for column in ("status", "Status"):
			env = export_query(
				spec=_no_comment_spec(),
				highlights=[{"column": column, "operator": "eq", "value": "Open", "style": "warning"}],
			)
			self.assertEqual(env["highlight_count"], 1)

	def test_the_export_permission_is_required(self):
		user = "xqs-reader@example.com"
		if not frappe.db.exists("User", user):
			frappe.get_doc(
				{"doctype": "User", "email": user, "first_name": "xqs", "send_welcome_email": 0}
			).insert(ignore_permissions=True)
			frappe.db.commit()
		frappe.set_user(user)
		try:
			# ToDo is readable by everyone (role All), exportable only by System Manager.
			with self.assertRaisesRegex(PermissionDeniedError, "no export permission on ToDo"):
				export_query(spec={"from": "ToDo", "select": ["name", "description"]})
			# A table the subquery reads is held to read permission, as in query.
			with self.assertRaisesRegex(PermissionDeniedError, "referenced DocType: Comment"):
				export_query(spec=_no_comment_spec())
		finally:
			frappe.set_user("Administrator")

	def test_the_doctype_form_now_has_label_headers(self):
		env = export_query("ToDo", filters={"description": ["like", f"{PFX}%"]}, fields=["name", "status"])
		self.assertEqual(_sheet(env)[0], ("ID", "Status"))


class TestExportQuerySpecLimitsAndHeaders(FrappeTestCase):
	"""The rules beyond the happy path: an export is not held to query's model-facing
	caps, a slow query is stopped and refused, shared labels are told apart."""

	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		_seed()

	@classmethod
	def tearDownClass(cls):
		_purge()
		super().tearDownClass()

	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")

	def test_query_caps_do_not_cut_an_export(self):
		from jarvis.tools import query as query_module

		with (
			mock.patch.object(query_module, "MAX_LIMIT", 50),
			mock.patch.object(query_module, "ROW_GUARD", 20),
		):
			env = export_query(spec=_no_comment_spec())
		self.assertEqual(env["total"], TOTAL - WITH_COMMENT)

	def test_a_query_stopped_at_the_time_limit_is_refused_with_what_to_do(self):
		timed_out = frappe.db.OperationalError(
			1969, "Query execution was interrupted (max_statement_time exceeded)"
		)
		with (
			mock.patch("jarvis.api._agent_read_timeout_s", return_value=25),
			mock.patch("jarvis.api.statement_time_limit", return_value=contextlib.nullcontext()),
			mock.patch.object(frappe.db, "sql", side_effect=timed_out),
		):
			# The export's own limit (12 s) is the smaller of the two.
			with self.assertRaisesRegex(ExportTooSlowError, "ran past the 12-second limit"):
				resolvers.from_spec(_no_comment_spec())

	def test_the_select_limit_is_the_smaller_of_the_agent_and_export_limits(self):
		for agent_limit, expected in ((25, 12), (5, 5), (0, 0)):  # 0: the site switched limits off
			with (
				mock.patch("jarvis.api._agent_read_timeout_s", return_value=agent_limit),
				mock.patch("jarvis.api.statement_time_limit", return_value=contextlib.nullcontext()) as cap,
			):
				resolvers.from_spec(_no_comment_spec())
			cap.assert_called_once_with(expected)

	def test_the_select_runs_under_the_statement_limit(self):
		if frappe.db.db_type != "mariadb":
			self.skipTest("max_statement_time is MariaDB's")
		# The limit keeps a tighter one the session already has, which an earlier
		# test in the same process may have left behind: start from none.
		previous = frappe.db.sql("SELECT @@max_statement_time")[0][0]
		frappe.db.sql("SET SESSION max_statement_time = 0")
		self.addCleanup(frappe.db.sql, "SET SESSION max_statement_time = %s", (previous,))
		with (
			mock.patch("jarvis.api._agent_read_timeout_s", return_value=25),
			mock.patch.object(
				frappe.db, "set_execution_timeout", wraps=frappe.db.set_execution_timeout
			) as cap,
		):
			resolvers.from_spec(_no_comment_spec())
		cap.assert_called_with(12)

	def test_a_limit_that_cut_nothing_adds_no_note(self):
		exact = TOTAL - WITH_COMMENT
		for limit in (exact, exact + 800):
			env = export_query(spec=_no_comment_spec(limit=limit))
			self.assertEqual(env["total"], exact)
			self.assertNotIn("limited_by_spec", env)
			self.assertNotIn("note", env)

	def test_a_limit_one_below_the_matches_is_disclosed(self):
		env = export_query(spec=_no_comment_spec(limit=TOTAL - WITH_COMMENT - 1))
		self.assertEqual(env["total"], TOTAL - WITH_COMMENT - 1)
		self.assertEqual(env["limited_by_spec"], TOTAL - WITH_COMMENT - 1)

	def test_a_limit_above_the_ceiling_requires_a_chat(self):
		with mock.patch.object(resolvers, "_HARD_ROW_CEILING", 10):
			with self.assertRaises(InvalidArgumentError) as ctx:
				export_query(spec=_no_comment_spec(limit=50))
			self.assertIn("Narrow", str(ctx.exception))
			self.assertIn("Start it from a Jarvis chat", str(ctx.exception))
			# A limit equal to the ceiling is a deliberate cap, not an overflow.
			env = export_query(spec=_no_comment_spec(limit=10))
		self.assertEqual((env["total"], env["limited_by_spec"]), (10, 10))

	def test_a_value_cannot_break_out_of_its_quotes(self):
		# A backslash before a quote ends a MariaDB string early when values are
		# inlined; a spec value must stay a value (parameterised), so this matches
		# nothing instead of every row.
		spec = {
			"from": "ToDo",
			"select": ["name"],
			"where": [{"field": "description", "op": "=", "value": "a\\' OR 1=1 -- -"}],
		}
		with self.assertRaises(NoDataError):
			export_query(spec=spec)

	def test_shared_labels_are_told_apart(self):
		cols = [
			resolvers._Column("name", "Supplier", None),
			resolvers._Column("invoice", "Purchase Invoice", "ID"),
			resolvers._Column("status", "ToDo", None),
		]
		self.assertEqual(resolvers._unique_headers(cols), ["Supplier ID", "Purchase Invoice ID", "Status"])
		same_doctype = [resolvers._Column("a", None, "Total"), resolvers._Column("b", None, "Total")]
		self.assertEqual(resolvers._unique_headers(same_doctype), ["Total (a)", "Total (b)"])

	def test_a_bad_highlight_column_names_the_columns_there_are(self):
		with self.assertRaisesRegex(
			InvalidArgumentError, "Columns in this export: ID \\(name\\), Description"
		):
			export_query(
				spec=_no_comment_spec(),
				highlights=[{"column": "nope", "operator": "eq", "value": "x", "style": "warning"}],
			)


class TestExportQuerySpecExportPermission(FrappeTestCase):
	"""The Export DocPerm is needed on every table a spec export's columns come from:
	a joined table, and for a child table the parent its rows are scoped to (read
	alone is enough for query, never for a bulk file)."""

	READER = "xqs-si-reader@example.com"  # read on Sales Invoice, no export
	EXPORTER = "xqs-si-exporter@example.com"  # read + export on Sales Invoice, read on Customer
	ROLES = ("XQS SI Reader", "XQS SI Exporter")

	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		frappe.set_user("Administrator")
		for role in cls.ROLES:
			if not frappe.db.exists("Role", role):
				frappe.get_doc({"doctype": "Role", "role_name": role, "desk_access": 1}).insert()
		# A new DocPerm row has export on by default: set it explicitly both ways.
		for dt, role, export in (
			("Sales Invoice", "XQS SI Reader", 0),
			("Sales Invoice", "XQS SI Exporter", 1),
			("Customer", "XQS SI Exporter", 0),
		):
			frappe.permissions.add_permission(dt, role, 0)
			frappe.permissions.update_permission_property(dt, role, 0, "export", export)
		for user, role in ((cls.READER, "XQS SI Reader"), (cls.EXPORTER, "XQS SI Exporter")):
			if not frappe.db.exists("User", user):
				frappe.get_doc(
					{"doctype": "User", "email": user, "first_name": "xqs", "send_welcome_email": 0}
				).insert(ignore_permissions=True)
			frappe.get_doc("User", user).add_roles(role)
		frappe.db.commit()
		frappe.clear_cache()

	@classmethod
	def tearDownClass(cls):
		frappe.set_user("Administrator")
		for dt in ("Sales Invoice", "Customer"):
			frappe.db.delete("Custom DocPerm", {"parent": dt, "role": ["in", cls.ROLES]})
		frappe.db.commit()
		frappe.clear_cache()
		super().tearDownClass()

	def tearDown(self):
		frappe.set_user("Administrator")
		super().tearDown()

	def test_a_child_table_needs_export_on_its_parent(self):
		frappe.set_user(self.READER)
		self.assertTrue(frappe.has_permission("Sales Invoice", "read"))
		self.assertFalse(frappe.permissions.can_export("Sales Invoice"))
		with self.assertRaisesRegex(PermissionDeniedError, "no export permission on Sales Invoice"):
			export_query(spec={"from": "Sales Invoice Item", "select": ["item_code", "amount"]})

	def test_a_joined_table_needs_export_too(self):
		frappe.set_user(self.EXPORTER)
		spec = {
			"from": "Sales Invoice",
			"alias": "si",
			"joins": [{"doctype": "Customer", "alias": "c", "on": {"c.name": "si.customer"}}],
			"select": ["si.name", "c.customer_group"],
		}
		with self.assertRaisesRegex(PermissionDeniedError, "no export permission on Customer"):
			export_query(spec=spec)

	def test_a_child_through_an_exportable_parent_is_allowed(self):
		frappe.set_user(self.EXPORTER)
		with mock.patch.object(resolvers, "_run_bounded", return_value=[{"item_code": "X", "amount": 1}]):
			model = resolvers.from_spec({"from": "Sales Invoice Item", "select": ["item_code", "amount"]})
		self.assertEqual(model.total, 1)


class TestExportQueryDeadline(FrappeTestCase):
	"""Past the immediate deadline, hand off before saving an immediate File."""

	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		_seed()

	@classmethod
	def tearDownClass(cls):
		_purge()
		super().tearDownClass()

	def setUp(self):
		super().setUp()
		frappe.set_user("Administrator")
		self.clock = _Clock()

	def _deferred(self, patch_target: str, seconds: float, **call):
		"""Run an export whose ``patch_target`` step takes ``seconds``; return what the
		handoff did: the telemetry outcome, the File count change, and save calls."""
		real = save_export_file
		files = frappe.db.count("File")
		with (
			mock.patch.object(export_module, "time", self.clock),
			mock.patch(patch_target, side_effect=self.clock.after(seconds, self._real(patch_target))),
			mock.patch.object(export_module, "save_export_file", wraps=real) as save,
			mock.patch("jarvis.telemetry.record_export_event") as event,
			mock.patch(
				"jarvis.tools._export.jobs.enqueue",
				return_value={"export_job": "queued-test", "status": "queued"},
			) as enqueue,
		):
			self.assertEqual(export_query(**call)["status"], "queued")
			enqueue.assert_called_once()
		save.assert_not_called()
		self.assertEqual(frappe.db.count("File"), files)
		event.assert_called_once()
		return event.call_args.kwargs["outcome"]

	@staticmethod
	def _real(patch_target: str):
		module, _, name = patch_target.rpartition(".")
		return getattr(importlib.import_module(module), name)

	def test_a_slow_spec_query_is_deferred_before_a_file_is_saved(self):
		outcome = self._deferred(
			"jarvis.tools._export.resolvers.from_spec", 26, spec=_no_comment_spec(), title="slow"
		)
		self.assertEqual(outcome, "too_slow")

	def test_a_slow_doctype_query_is_deferred_too(self):
		outcome = self._deferred(
			"jarvis.tools._export.resolvers.from_query",
			26,
			doctype="ToDo",
			filters={"description": ["like", f"{PFX}%"]},
			fields=["name"],
		)
		self.assertEqual(outcome, "too_slow")

	def test_a_slow_render_is_deferred_before_a_file_is_saved(self):
		mime, render = export_module._FORMATS["xlsx"]
		slow = self.clock.after(26, render)
		with mock.patch.dict(export_module._FORMATS, {"xlsx": (mime, slow)}):
			outcome = self._deferred(
				"jarvis.tools._export.resolvers.from_spec", 0, spec=_no_comment_spec(), title="slow"
			)
		self.assertEqual(outcome, "too_slow")

	def test_an_export_inside_the_deadline_is_saved(self):
		real = resolvers.from_spec
		with (
			mock.patch.object(export_module, "time", self.clock),
			mock.patch.object(resolvers, "from_spec", side_effect=self.clock.after(24, real)),
		):
			env = export_query(spec=_no_comment_spec())
		self.assertEqual(env["total"], TOTAL - WITH_COMMENT)
		self.assertTrue(frappe.db.exists("File", env["name"]))
