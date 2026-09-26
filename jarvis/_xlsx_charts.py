"""Validated, cell-referenced charts for the two supported Excel writers."""

import datetime
import math
import re
from dataclasses import dataclass
from decimal import Decimal
from numbers import Real

from jarvis.exceptions import InvalidArgumentError


@dataclass(frozen=True)
class ExcelChart:
	type: str
	categories: int
	values: tuple[int, ...]
	title: str | None


def validate_charts(charts, data):
	"""Resolve header names before creating any workbook or private File."""
	if charts is None:
		return []
	if not isinstance(charts, list) or len(charts) > 4:
		raise InvalidArgumentError("charts must be a list of at most 4 charts per sheet")
	header = data[0]
	if charts and len(header) > 16382:
		raise InvalidArgumentError("Leave two columns of space beside the data for charts")

	def column(name):
		if not isinstance(name, str) or not name.strip() or header.count(name) != 1:
			raise InvalidArgumentError("Chart columns must name exactly one header in the exported sheet")
		return header.index(name)

	result = []
	for spec in charts:
		if not isinstance(spec, dict) or set(spec) - {"type", "categories", "values", "title"}:
			raise InvalidArgumentError("Each chart supports only type, categories, values and optional title")
		kind = spec.get("type")
		if kind not in ("column", "bar", "line", "pie"):
			raise InvalidArgumentError("Chart type must be column, bar, line or pie")
		category = column(spec.get("categories"))
		values = spec.get("values")
		if not isinstance(values, list) or not 1 <= len(values) <= 6:
			raise InvalidArgumentError("Chart values must list 1 to 6 numeric column headers")
		value_cols = tuple(column(name) for name in values)
		if len(set(value_cols)) != len(value_cols):
			raise InvalidArgumentError("Chart values must not repeat a column")
		if kind == "pie" and len(values) != 1:
			raise InvalidArgumentError("A pie chart must have exactly one value column")
		title = spec.get("title")
		if "title" in spec and (not isinstance(title, str) or len(title) > 120):
			raise InvalidArgumentError("Chart title must be text of at most 120 characters")
		# XlsxWriter interprets these titles as formulas, including external refs.
		if title and re.match(r"^=?[^!]+!\$?[A-Z]+\$?\d+", title):
			raise InvalidArgumentError("Chart title must be plain text, not a cell reference")
		for col in value_cols:
			has_number = False
			for row_number, row in enumerate(data[1:], 2):
				value = row[col] if col < len(row) else None
				if value is None or (isinstance(value, str) and not value.strip()):
					continue
				valid = isinstance(value, (Real, Decimal)) and not isinstance(value, bool)
				try:
					valid = valid and math.isfinite(value)
				except (ValueError, OverflowError):
					valid = False
				if not valid:
					raise InvalidArgumentError(
						f"Chart column {header[col]} must contain numbers or blanks (row {row_number})"
					)
				has_number = True
			if not has_number:
				raise InvalidArgumentError(f"Chart column {header[col]} needs at least one numeric value")
		result.append(ExcelChart(kind, category, value_cols, title))
	return result


def add_xlsxwriter_charts(workbook, worksheet, data, charts):
	"""Use the actual sanitized worksheet name and cache streaming cell data."""
	from frappe.utils.xlsxutils import ILLEGAL_CHARACTERS_RE, handle_html

	def cell(row, col):
		value = row[col] if col < len(row) else None
		if isinstance(value, str):
			if worksheet.get_name() not in {"Data Import Template", "Data Export"}:
				value = handle_html(value)
			value = ILLEGAL_CHARACTERS_RE.sub("", value)
		return value

	def category(row, col):
		value = cell(row, col)
		# Streaming worksheets cannot provide XlsxWriter's chart caches later.
		# Category caches need scalar labels, not Python date/bool objects.
		if isinstance(value, (datetime.date, datetime.time)):
			return value.isoformat()
		if isinstance(value, bool):
			return "TRUE" if value else "FALSE"
		return value

	name = worksheet.get_name()
	last = len(data) - 1
	for index, spec in enumerate(charts):
		chart = workbook.add_chart({"type": spec.type})
		for col in spec.values:
			chart.add_series(
				{
					"name": [name, 0, col],
					"name_data": [cell(data[0], col)],
					"categories": [name, 1, spec.categories, last, spec.categories],
					"categories_data": [category(row, spec.categories) for row in data[1:]],
					"values": [name, 1, col, last, col],
					"values_data": [
						None if isinstance(v := cell(row, col), str) and not v.strip() else v
						for row in data[1:]
					],
				}
			)
		if spec.title:
			chart.set_title({"name": spec.title})
		# Default chart height is 288px; 20 standard rows keeps charts apart.
		worksheet.insert_chart(1 + index * 20, len(data[0]) + 1, chart)


def add_openpyxl_charts(worksheet, data, charts):
	from openpyxl.chart import BarChart, LineChart, PieChart, Reference
	from openpyxl.utils import get_column_letter

	for index, spec in enumerate(charts):
		if spec.type in ("column", "bar"):
			chart = BarChart()
			chart.type = "col" if spec.type == "column" else "bar"
		else:
			chart = LineChart() if spec.type == "line" else PieChart()
		for col in spec.values:
			chart.add_data(
				Reference(worksheet, min_col=col + 1, min_row=1, max_row=len(data)), titles_from_data=True
			)
		chart.set_categories(Reference(worksheet, min_col=spec.categories + 1, min_row=2, max_row=len(data)))
		if spec.title:
			chart.title = spec.title
		worksheet.add_chart(chart, f"{get_column_letter(len(data[0]) + 2)}{2 + index * 20}")
