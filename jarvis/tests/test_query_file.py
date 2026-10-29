"""Full spreadsheet analysis: real workbook parsing and read-only SQLite."""

import io
import json
import unittest
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook

from jarvis.exceptions import InvalidArgumentError, PermissionDeniedError
from jarvis.tools._spreadsheet_query import query_spreadsheet
from jarvis.tools.query_file import query_file


def workbook(months=6, rows=1200):
	book = Workbook(write_only=True)
	for month in range(1, months + 1):
		sheet = book.create_sheet(f"Month {month}")
		sheet.append(["GSTR-2B report", None, None, None, None])
		sheet.append(["GSTIN", "Supplier", "Invoice", "Taxable Value", "IGST"])
		for row in range(1, rows + 1):
			sheet.append(["29ABCDE1234F1Z5", "Supplier A", f"M{month}-{row}", "1,23,456.78", "0.10"])
		sheet.append(["TOTAL", None, None, "do not count the total row", None])
	buffer = io.BytesIO()
	book.save(buffer)
	return buffer.getvalue()


class TestSpreadsheetQuery(unittest.TestCase):
	@classmethod
	def setUpClass(cls):
		cls.content = workbook()

	def query(self, sql=None, **kwargs):
		return query_spreadsheet(self.content, "2b.xlsx", sql=sql, **kwargs)

	def test_six_months_full_counts_and_exact_amounts(self):
		result = self.query(
			"SELECT COUNT(*) AS invoices, decimal_sum(D) AS taxable, decimal_sum(E) AS igst "
			"FROM ("
			+ " UNION ALL ".join(
				f"SELECT * FROM sheet{month} WHERE _row > 2 AND A != 'TOTAL'" for month in range(1, 7)
			)
			+ ")"
		)
		self.assertEqual(result["rows"], [[7200, "888888816.00", "720.00"]])
		self.assertEqual(result["source_rows"], 7218)
		self.assertTrue(result["coverage"]["complete"])
		self.assertFalse(result["result_truncated"])

	def test_tail_invoice_and_duplicate_detection(self):
		tail = self.query("SELECT _row, C FROM sheet6 WHERE C = 'M6-1200'")
		self.assertEqual(tail["rows"], [[1202, "M6-1200"]])
		grouped = self.query(
			"SELECT A, COUNT(*) FROM sheet1 WHERE _row > 2 AND A != 'TOTAL' GROUP BY A HAVING COUNT(*) > 1"
		)
		self.assertEqual(grouped["rows"], [["29ABCDE1234F1Z5", 1200]])

	def test_discovery_preserves_titles_headers_and_source_positions(self):
		result = self.query()
		self.assertEqual(len(result["sheets"]), 6)
		first = result["sheets"][0]
		self.assertEqual(first["table"], "sheet1")
		self.assertEqual(first["row_count"], 1203)
		self.assertEqual(first["columns"], ["A", "B", "C", "D", "E"])
		self.assertEqual(first["sample"][1], [2, "GSTIN", "Supplier", "Invoice", "Taxable Value", "IGST"])

	def test_source_coverage_is_separate_from_result_display_limit(self):
		result = self.query("SELECT C FROM sheet1 WHERE _row > 2", max_rows=3)
		self.assertEqual(result["rows"], [["M1-1"], ["M1-2"], ["M1-3"]])
		self.assertTrue(result["result_truncated"])
		self.assertEqual(result["source_rows"], 7218)
		self.assertTrue(result["coverage"]["complete"])

	def test_amount_formats_negative_notes_and_invalid_values(self):
		data = b'GSTIN,Amount\nA,0.10\nA,0.20\nB,(10.25)\nC,"1,23,456.78"\n'
		result = query_spreadsheet(data, "2b.csv", sql="SELECT decimal_sum(B) FROM sheet1 WHERE _row > 1")
		self.assertEqual(Decimal(result["rows"][0][0]), Decimal("123446.83"))
		with self.assertRaisesRegex(InvalidArgumentError, "numeric"):
			query_spreadsheet(data, "2b.csv", sql="SELECT decimal_sum(B) FROM sheet1")

	def test_only_selected_sheets_and_stable_table_names(self):
		result = self.query("SELECT COUNT(*) FROM sheet6", sheets=["Month 6"])
		self.assertEqual(result["rows"], [[1203]])
		self.assertEqual(result["source_rows"], 1203)
		self.assertEqual([s["table"] for s in result["sheets"]], ["sheet6"])
		with self.assertRaisesRegex(InvalidArgumentError, "worksheet"):
			self.query(sheets=["missing"])

	def test_sql_refuses_mutation_files_pragmas_and_multiple_statements(self):
		for sql in (
			"DELETE FROM sheet1",
			"UPDATE sheet1 SET A = 'changed'",
			"DROP TABLE sheet1",
			"ATTACH DATABASE '/tmp/query-file-should-not-exist.db' AS other",
			"PRAGMA database_list",
			"SELECT load_extension('/tmp/missing.so')",
			"SELECT * FROM sqlite_master",
			"SELECT 1; SELECT 2",
			"SELECT randomblob(1000000000)",
		):
			with self.subTest(sql=sql), self.assertRaises(InvalidArgumentError):
				query_spreadsheet(b"A\n1\n", "2b.csv", sql=sql)

	def test_formula_without_cache_cannot_be_used_as_complete_data(self):
		book = Workbook()
		book.active.append(["Tax", "=1+1"])
		buffer = io.BytesIO()
		book.save(buffer)
		with self.assertRaisesRegex(InvalidArgumentError, "formula"):
			query_spreadsheet(buffer.getvalue(), "2b.xlsx", sql="SELECT decimal_sum(B) FROM sheet1")

	def test_cached_empty_formula_remains_empty_text(self):
		from zipfile import ZIP_DEFLATED, ZipFile

		book = Workbook()
		book.active.append(["A", '=IF(1,"","")'])
		buffer = io.BytesIO()
		book.save(buffer)
		updated = io.BytesIO()
		with ZipFile(buffer) as source, ZipFile(updated, "w", ZIP_DEFLATED) as target:
			for name in source.namelist():
				data = source.read(name)
				if name.startswith("xl/worksheets/"):
					data = data.replace(b'<c r="B1">', b'<c r="B1" t="str">')
				target.writestr(name, data)
		result = query_spreadsheet(updated.getvalue(), "2b.xlsx", sql="SELECT B, B IS NULL FROM sheet1")
		self.assertEqual(result["rows"], [["", 0]])

	def test_exact_functions_grouping_blank_values_and_typed_dates(self):
		from datetime import date

		book = Workbook()
		book.active.append([date(2026, 4, 1), "A", "0.10", None])
		book.active.append([date(2026, 4, 2), "A", "0.20", None])
		book.active.append([date(2026, 5, 1), "B", "(10.00)", None])
		buffer = io.BytesIO()
		book.save(buffer)
		result = query_spreadsheet(
			buffer.getvalue(),
			"2b.xlsx",
			sql=(
				"SELECT substr(A, 1, 7), decimal_sum(C), decimal_min(C), decimal_max(C), "
				"decimal_avg(C), decimal_sub(decimal_sum(C), '0.05') FROM sheet1 GROUP BY substr(A, 1, 7) ORDER BY 1"
			),
		)
		self.assertEqual(
			result["rows"],
			[
				["2026-04", "0.30", "0.10", "0.20", "0.15", "0.25"],
				["2026-05", "-10.00", "-10.00", "-10.00", "-10.00", "-10.05"],
			],
		)
		with self.assertRaisesRegex(InvalidArgumentError, "divisor"):
			query_spreadsheet(b"A\n", "2b.csv", sql="SELECT decimal_div(1, 0)")

	def test_output_budget_retains_whole_rows_and_reports_hidden_sheets(self):
		from jarvis.tools import _spreadsheet_query as module

		with patch.object(module, "MAX_OUTPUT_CHARS", 3000):
			result = query_spreadsheet(
				("A\n" + ("x" * 1000 + "\n") * 8).encode(),
				"2b.csv",
				sql="SELECT A FROM sheet1 WHERE _row > 1",
			)
			self.assertLess(len(json.dumps(result)), 3200)
			self.assertTrue(result["result_truncated"])
			self.assertEqual(result["source_rows"], 9)
			self.assertTrue(all(row == ["x" * 1000] for row in result["rows"]))
		book = Workbook()
		book.active.append(["A", None, "B"])
		book.create_sheet("Hidden").sheet_state = "hidden"
		buffer = io.BytesIO()
		book.save(buffer)
		result = query_spreadsheet(buffer.getvalue(), "2b.xlsx", sql="SELECT B IS NULL FROM sheet1")
		self.assertEqual(result["rows"], [[1]])
		self.assertTrue(result["sheets"][1]["hidden"])
		self.assertEqual(result["sheets"][1]["row_count"], 0)

	def test_query_cannot_silently_treat_money_text_as_zero(self):
		for expression in ("sum(B)", "avg(B)", "total(B)"):
			with self.subTest(expression=expression), self.assertRaises(InvalidArgumentError):
				query_spreadsheet(b"A,B\nINV,10.00\n", "2b.csv", sql=f"SELECT {expression} FROM sheet1")

	def test_metadata_budget_applies_even_when_query_returns_no_rows(self):
		from jarvis.tools import _spreadsheet_query as module

		with patch.object(module, "MAX_OUTPUT_CHARS", 50):
			for sql in (None, "SELECT A FROM sheet1 WHERE 0"):
				with self.subTest(sql=sql), self.assertRaisesRegex(InvalidArgumentError, "metadata"):
					query_spreadsheet(b"A\n1\n", "2b.csv", sql=sql)

	def test_malformed_and_hard_limits_fail_without_partial_totals(self):
		from jarvis.tools import _spreadsheet_query as module

		with patch.object(module, "MAX_TOTAL_ROWS", 10), self.assertRaisesRegex(InvalidArgumentError, "rows"):
			self.query("SELECT COUNT(*) FROM sheet1")
		with self.assertRaisesRegex(InvalidArgumentError, "UTF-8"):
			query_spreadsheet(b"A\n\xff", "2b.csv")
		with self.assertRaises(InvalidArgumentError):
			query_spreadsheet(b'"unclosed', "2b.csv")


class TestQueryFileBoundary(unittest.TestCase):
	def test_real_api_dispatch_and_agent_budget_keep_full_source_aggregate(self):
		import frappe

		from jarvis.api import _run_tool
		from jarvis.tools._result_guard import enforce_result_budget

		doc = SimpleNamespace(name="file", file_name="2b.csv", file_url="/private/files/2b.csv")
		with (
			patch.object(frappe, "flags", frappe._dict()),
			patch.object(frappe, "conf", frappe._dict()),
			patch("jarvis.tools.query_file._resolve_file", return_value=doc),
			patch.object(frappe, "has_permission", return_value=True),
			patch("jarvis.compat.file_bytes", return_value=b"GSTIN,Tax\n" + b"A,0.10\n" * 1500),
		):
			result = _run_tool(
				"query_file",
				{
					"file_url": doc.file_url,
					"sql": "SELECT COUNT(*), decimal_sum(B) FROM sheet1 WHERE _row > 1",
				},
			)
		self.assertTrue(result["ok"], result)
		self.assertEqual(result["data"]["rows"], [[1500, "150.00"]])
		guarded, event = enforce_result_budget(result["data"], "query_file")
		self.assertIsNone(event)
		self.assertEqual(guarded["rows"], [[1500, "150.00"]])

	def test_attachment_hint_routes_large_spreadsheets_without_inlining_source(self):
		import frappe

		# The module resolves the configured admin URL at import time. No site is
		# connected for this prompt-only unit test.
		with patch.object(frappe, "conf", frappe._dict()):
			from jarvis.chat.turn_handler import _prepare_attachments

		doc = SimpleNamespace(name="file", file_name="2b.xlsx")
		with (
			patch("frappe.get_doc", return_value=doc),
			patch("frappe.has_permission", return_value=True),
			patch("jarvis.compat.file_bytes", return_value=b"PRIVATE-WORKBOOK-CONTENT"),
		):
			message, images = _prepare_attachments(
				"Total IGST?", [{"file_url": "/private/files/2b.xlsx"}], False
			)
		self.assertIn("jarvis__query_file", message)
		self.assertIn("all source rows", message)
		self.assertNotIn("PRIVATE-WORKBOOK-CONTENT", message)
		self.assertEqual(images, [])

	def test_permission_precedes_bytes_and_child_on_every_call(self):
		for version in (None, "a" * 64):
			with (
				patch("jarvis.tools.query_file._resolve_file", return_value=SimpleNamespace(name="private")),
				patch("jarvis.tools.query_file.frappe.has_permission", return_value=False),
				patch("jarvis.tools.query_file.compat.file_bytes") as read,
				patch("jarvis.tools.query_file.run_query") as child,
			):
				with self.assertRaises(PermissionDeniedError):
					query_file(file_url="/private/files/2b.xlsx", version=version)
				read.assert_not_called()
				child.assert_not_called()

	def test_changed_file_rejects_pinned_version(self):
		doc = SimpleNamespace(name="file", file_name="2b.csv", file_url="/private/files/2b.csv")
		with (
			patch("jarvis.tools.query_file._resolve_file", return_value=doc),
			patch("jarvis.tools.query_file.frappe.has_permission", return_value=True),
			patch("jarvis.tools.query_file.compat.file_bytes", side_effect=[b"A\n1", b"A\n2"]),
			patch("jarvis.tools.query_file.run_query", return_value={}) as child,
		):
			first = query_file(file_url=doc.file_url)
			with self.assertRaisesRegex(InvalidArgumentError, "changed"):
				query_file(file_url=doc.file_url, version=first["version"])
			self.assertEqual(child.call_count, 1)

	def test_real_child_process_counts_the_tail(self):
		from jarvis.tools._spreadsheet_process import run_query

		result = run_query(workbook(months=1, rows=1500), "2b.xlsx", sql="SELECT COUNT(*) FROM sheet1")
		self.assertEqual(result["rows"], [[1503]])

	def test_real_child_runs_six_months_of_financial_aggregates(self):
		from jarvis.tools._spreadsheet_process import run_query

		result = run_query(
			workbook(),
			"2b.xlsx",
			sql=(
				"SELECT COUNT(*), decimal_sum(E) FROM ("
				+ " UNION ALL ".join(
					f"SELECT * FROM sheet{month} WHERE _row > 2 AND A != 'TOTAL'" for month in range(1, 7)
				)
				+ ")"
			),
		)
		self.assertEqual(result["rows"], [[7200, "720.00"]])

	def test_real_child_wall_timeout_and_source_labels_are_not_shell_commands(self):
		import base64
		import json
		import sys

		from jarvis.tools import _spreadsheet_process as module

		with (
			patch.object(module, "WALL_SECONDS", 0.0001),
			self.assertRaisesRegex(InvalidArgumentError, "resource budget"),
		):
			module.run_query(b"A\n1\n", "2b.csv")
		label = "$(touch /tmp/should-never-execute);`whoami`\n--help.csv"
		with patch.object(
			module.subprocess, "run", return_value=SimpleNamespace(returncode=0, stdout='{"result": {}}')
		) as run:
			module.run_query(b"private data", label, sql=label)
		args, kwargs = run.call_args
		self.assertEqual(args[0], [sys.executable, "-m", "jarvis.tools._spreadsheet_worker"])
		self.assertIs(kwargs["shell"], False)
		payload = json.loads(kwargs["input"])
		self.assertEqual(base64.b64decode(payload["content"]), b"private data")
		self.assertEqual(payload["filename"], label)
		self.assertEqual(payload["sql"], label)


if __name__ == "__main__":
	unittest.main()
