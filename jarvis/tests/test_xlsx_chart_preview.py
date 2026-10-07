"""Excel charts and date cells in the file preview panel (#627)."""

from __future__ import annotations

import datetime
import io
import unittest
import zipfile

import openpyxl
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference

from jarvis.chat.xlsx_charts import (
	MAX_CHARTS_PER_SHEET,
	MAX_CHARTS_TOTAL,
	MAX_POINTS,
	extract_charts,
	preview_cell,
)


def _book(title="Data", rows=None):
	wb = openpyxl.Workbook()
	ws = wb.active
	ws.title = title
	for r in rows or [
		("Customer", "Outstanding", "Paid"),
		("Acme", 100, 5),
		("Zed", 250.5, 7),
		("Bob", 40, 9),
	]:
		ws.append(r)
	return wb, ws


def _add(ws, chart, value_cols=(2,), anchor="F2", title=None):
	for col in value_cols:
		chart.add_data(Reference(ws, min_col=col, min_row=1, max_row=ws.max_row), titles_from_data=True)
	chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
	if title:
		chart.title = title
	ws.add_chart(chart, anchor)


def _bytes(wb):
	buf = io.BytesIO()
	wb.save(buf)
	return buf.getvalue()


class TestExtractCharts(unittest.TestCase):
	def test_column_chart(self):
		wb, ws = _book()
		chart = BarChart()
		chart.type = "col"
		_add(ws, chart, title="Outstanding by customer")
		(spec,) = extract_charts(_bytes(wb))
		self.assertEqual(spec["sheet"], "Data")
		self.assertEqual(spec["type"], "bar")
		self.assertEqual(spec["title"], "Outstanding by customer")
		self.assertEqual(spec["x"], ["Acme", "Zed", "Bob"])
		self.assertEqual(spec["series"], [{"name": "Outstanding", "data": [100.0, 250.5, 40.0]}])
		self.assertFalse(spec["options"].get("horizontal"))

	def test_horizontal_bar(self):
		wb, ws = _book()
		chart = BarChart()
		chart.type = "bar"
		_add(ws, chart, value_cols=(2, 3))
		(spec,) = extract_charts(_bytes(wb))
		self.assertTrue(spec["options"]["horizontal"])
		self.assertEqual([s["name"] for s in spec["series"]], ["Outstanding", "Paid"])

	def test_line_area_pie(self):
		for cls, expected in ((LineChart, "line"), (AreaChart, "area"), (PieChart, "pie")):
			wb, ws = _book()
			_add(ws, cls())
			(spec,) = extract_charts(_bytes(wb))
			self.assertEqual(spec["type"], expected)

	def test_sheet_name_with_spaces_and_quote(self):
		wb, ws = _book(title="Q1 O'Neil")
		_add(ws, BarChart())
		(spec,) = extract_charts(_bytes(wb))
		self.assertEqual(spec["sheet"], "Q1 O'Neil")
		self.assertEqual(spec["series"][0]["data"], [100.0, 250.5, 40.0])

	def test_two_charts_on_two_sheets(self):
		wb, ws = _book()
		_add(ws, BarChart(), title="One")
		_add(ws, LineChart(), anchor="F20", title="Two")
		ws2 = wb.create_sheet("Other")
		ws2.append(("k", "v"))
		ws2.append(("a", 1))
		_add(ws2, PieChart(), title="Three")
		specs = extract_charts(_bytes(wb))
		self.assertEqual(
			[(s["sheet"], s["title"]) for s in specs], [("Data", "One"), ("Data", "Two"), ("Other", "Three")]
		)

	def test_no_charts(self):
		wb, _ = _book()
		self.assertEqual(extract_charts(_bytes(wb)), [])

	def test_not_a_zip(self):
		self.assertEqual(extract_charts(b"nope"), [])

	def test_malformed_chart_part_skipped(self):
		wb, ws = _book()
		_add(ws, BarChart(), title="Good")
		_add(ws, LineChart(), anchor="F20", title="Broken")
		src = _bytes(wb)
		out = io.BytesIO()
		with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w") as zout:
			for item in zin.infolist():
				data = zin.read(item.filename)
				if item.filename == "xl/charts/chart2.xml":
					data = b"<c:chartSpace><oops"
				zout.writestr(item, data)
		specs = extract_charts(out.getvalue())
		self.assertEqual([s["title"] for s in specs], ["Good"])

	def test_doctype_part_refused(self):
		wb, ws = _book()
		_add(ws, BarChart())
		src = _bytes(wb)
		out = io.BytesIO()
		with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w") as zout:
			for item in zin.infolist():
				data = zin.read(item.filename)
				if item.filename == "xl/charts/chart1.xml":
					data = b'<!DOCTYPE x [<!ENTITY a "b">]>' + data
				zout.writestr(item, data)
		self.assertEqual(extract_charts(out.getvalue()), [])

	def test_cache_fallback_when_cells_unavailable(self):
		wb, ws = _book()
		_add(ws, BarChart())
		src = _bytes(wb)
		out = io.BytesIO()
		cache = (
			'</f><numCache><ptCount val="2"/><pt idx="0"><v>7</v></pt><pt idx="1"><v>9</v></pt></numCache>'
		)
		with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w") as zout:
			for item in zin.infolist():
				data = zin.read(item.filename)
				if item.filename == "xl/charts/chart1.xml":
					text = data.decode().replace("'Data'!", "'Gone'!")
					text = text.replace("</f></numRef>", cache + "</numRef>")
					data = text.encode()
				zout.writestr(item, data)
		(spec,) = extract_charts(out.getvalue())
		self.assertEqual(spec["series"][0]["data"], [7.0, 9.0])

	def test_no_cells_and_no_cache_skips_chart(self):
		wb, ws = _book()
		_add(ws, BarChart())
		src = _bytes(wb)
		out = io.BytesIO()
		with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w") as zout:
			for item in zin.infolist():
				data = zin.read(item.filename)
				if item.filename == "xl/charts/chart1.xml":
					data = data.decode().replace("'Data'!", "'Gone'!").encode()
				zout.writestr(item, data)
		self.assertEqual(extract_charts(out.getvalue()), [])

	def test_bounds(self):
		rows = [("k", "v")] + [(f"r{i}", i) for i in range(MAX_POINTS + 100)]
		wb, ws = _book(rows=rows)
		_add(ws, BarChart())
		(spec,) = extract_charts(_bytes(wb))
		self.assertEqual(len(spec["x"]), MAX_POINTS)
		self.assertEqual(len(spec["series"][0]["data"]), MAX_POINTS)
		wb, ws = _book()
		for i in range(MAX_CHARTS_PER_SHEET + 3):
			_add(ws, BarChart(), anchor=f"F{2 + i * 20}")
		self.assertEqual(len(extract_charts(_bytes(wb))), MAX_CHARTS_PER_SHEET)

	def test_date_categories(self):
		rows = [("Day", "v"), (datetime.datetime(2026, 6, 1), 1), (datetime.datetime(2026, 6, 2), 2)]
		wb, ws = _book(rows=rows)
		_add(ws, LineChart())
		(spec,) = extract_charts(_bytes(wb))
		self.assertEqual(spec["x"], ["2026-06-01", "2026-06-02"])


def _rewrite(src, edits):
	"""Copy a workbook, applying {member: fn(bytes) -> bytes}."""
	out = io.BytesIO()
	with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w") as zout:
		for item in zin.infolist():
			data = zin.read(item.filename)
			if item.filename in edits:
				data = edits[item.filename](data)
			zout.writestr(item, data)
	return out.getvalue()


class TestHardening(unittest.TestCase):
	def _one(self):
		wb, ws = _book()
		_add(ws, BarChart())
		return _bytes(wb)

	def test_huge_ptcount_and_idx_are_not_allocated(self):
		cache = '</f><numCache><ptCount val="999999999"/><pt idx="0"><v>7</v></pt><pt idx="4000000000"><v>1</v></pt></numCache>'
		src = _rewrite(
			self._one(),
			{
				"xl/charts/chart1.xml": lambda d: d.decode()
				.replace("'Data'!", "'Gone'!")
				.replace("</f></numRef>", cache + "</numRef>")
				.encode()
			},
		)
		(spec,) = extract_charts(src)
		self.assertEqual(spec["series"][0]["data"], [7.0])

	def test_utf16_part_refused(self):
		src = _rewrite(
			self._one(),
			{
				"xl/charts/chart1.xml": lambda d: ('<!DOCTYPE x [<!ENTITY a "b">]>' + d.decode()).encode(
					"utf-16"
				)
			},
		)
		self.assertEqual(extract_charts(src), [])

	def test_non_utf8_declaration_refused(self):
		src = _rewrite(
			self._one(),
			{"xl/charts/chart1.xml": lambda d: b'<?xml version="1.0" encoding="UTF-16"?>' + d},
		)
		self.assertEqual(extract_charts(src), [])

	def test_wrong_root_refused(self):
		src = _rewrite(self._one(), {"xl/charts/chart1.xml": lambda d: b"<other/>"})
		self.assertEqual(extract_charts(src), [])

	def test_huge_worksheet_xml_is_not_parsed(self):
		src = _rewrite(self._one(), {"xl/worksheets/sheet1.xml": lambda d: d + b" " * 3_000_000})
		self.assertEqual(len(extract_charts(src)), 1)

	def test_per_sheet_cap_leaves_later_sheets(self):
		wb, ws = _book()
		for i in range(MAX_CHARTS_PER_SHEET + 2):
			_add(ws, BarChart(), anchor=f"F{2 + i * 20}")
		ws2 = wb.create_sheet("Later")
		ws2.append(("k", "v"))
		ws2.append(("a", 1))
		_add(ws2, PieChart(), title="Later chart")
		specs = extract_charts(_bytes(wb))
		self.assertEqual(len(specs), MAX_CHARTS_PER_SHEET + 1)
		self.assertEqual(specs[-1]["sheet"], "Later")
		self.assertLessEqual(len(specs), MAX_CHARTS_TOTAL)

	def test_rels_target_outside_xl_ignored(self):
		src = _rewrite(
			self._one(),
			{
				"xl/drawings/_rels/drawing1.xml.rels": lambda d: d.replace(
					b"/xl/charts/chart1.xml", b"../../evil.xml"
				)
			},
		)
		self.assertEqual(extract_charts(src), [])


class TestPreviewCell(unittest.TestCase):
	def test_dates(self):
		self.assertEqual(preview_cell("2026-06-01T00:00:00"), "2026-06-01")
		self.assertEqual(preview_cell("2026-06-01T14:30:05"), "2026-06-01 14:30")
		self.assertEqual(preview_cell("2026-06-01"), "2026-06-01")
		self.assertEqual(preview_cell("12:30:00"), "12:30:00")
		self.assertEqual(preview_cell("note T text here, long enough"), "note T text here, long enough")
		self.assertEqual(preview_cell(5), 5)
		self.assertEqual(preview_cell(""), "")
