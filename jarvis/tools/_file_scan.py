"""Deterministic, positional extraction windows; no jobs, models or retained state."""

from __future__ import annotations

import csv
import io
import json
import re
from itertools import islice, zip_longest
from zipfile import ZipFile

from jarvis.exceptions import InvalidArgumentError

MAX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_PAGES = 500
MAX_ROWS = 100_000  # Per worksheet.
MAX_SHEETS = 32
MAX_COLUMNS = 128
WINDOW_CHARS = 12_000
WINDOW_ROWS = 200
WINDOW_PAGES = 4
SCANNER_REVISION = "cursor-1"
CURSOR_PATTERN = r"[0-9]{1,2}:[0-9]{1,6}:[0-9]{1,8}"


def scanner_profile():
	return [
		SCANNER_REVISION,
		MAX_BYTES,
		MAX_EXPANDED_BYTES,
		MAX_PAGES,
		MAX_ROWS,
		MAX_SHEETS,
		MAX_COLUMNS,
		WINDOW_CHARS,
		WINDOW_ROWS,
		WINDOW_PAGES,
	]


def parse_cursor(cursor):
	if cursor is None:
		return 0, 1, 0
	if not isinstance(cursor, str) or not re.fullmatch(CURSOR_PATTERN, cursor):
		raise InvalidArgumentError("Invalid file cursor. Retry using the returned next_read arguments.")
	source, unit, offset = map(int, cursor.split(":"))
	if source >= MAX_SHEETS or not 1 <= unit <= MAX_ROWS:
		raise InvalidArgumentError("File cursor is outside the supported source range.")
	return source, unit, offset


def scan_file(content: bytes, filename: str, sheet: str | None = None, cursor=None) -> dict:
	from jarvis.tools.read_file import _cell

	source, start, offset = parse_cursor(cursor)
	if len(content) > MAX_BYTES:
		raise InvalidArgumentError("File exceeds the 10 MiB input limit. Split it into smaller files.")
	ext = filename.rsplit(".", 1)[-1].lower()
	out = {
		"kind": "text",
		"records": [],
		"issues": [],
		"stop_reason": None,
		"next_cursor": None,
		"source": {"filename": filename},
		"range": {"source": source, "first": start, "last": start - 1},
	}
	warnings = {}
	used = 2

	def issue(code, **where):
		# One diagnostic per class, count repetitions and retain its first location.
		# Terminal stop reasons have their OWN slot and cannot be crowded out.
		if code in warnings:
			warnings[code]["count"] += 1
		else:
			warnings[code] = {"code": code, "count": 1, **where}
			out["issues"].append(warnings[code])

	def pos(unit, char=0, src=source):
		return f"{src}:{unit}:{char}"

	def stop(code, **where):
		if out["stop_reason"] is not None:
			issue(code, **where)
		else:
			out["stop_reason"] = {"code": code, **where}

	def add(unit, value, is_row=False):
		"""Native rows normally; oversized single units resume exact serialized text."""
		nonlocal used
		char = offset if unit == start else 0
		key = "row" if is_row else "page" if ext == "pdf" else "text"
		record = {key: unit, "values" if is_row else "text": value}
		encoded = json.dumps(record, ensure_ascii=False)
		if char == 0 and used + len(encoded) + 1 <= WINDOW_CHARS:
			out["records"].append(record)
			used += len(encoded) + 1
			out["range"]["last"] = unit
			return True
		if out["records"]:
			out["next_cursor"] = pos(unit, char)
			return False
		text = json.dumps(value, ensure_ascii=False) if is_row else value
		if char >= len(text):
			raise InvalidArgumentError("File cursor offset is outside its source. Restart this file read.")
		# Bound the serialized response including quotes/control-character escaping.
		low, high = 0, min(len(text) - char, WINDOW_CHARS)
		while low < high:
			mid = (low + high + 1) // 2
			candidate = {
				key: unit,
				"offset": char,
				"fragment": text[char : char + mid],
				"encoding": "json-row" if is_row else "text",
			}
			if len(json.dumps(candidate, ensure_ascii=False)) + 3 <= WINDOW_CHARS:
				low = mid
			else:
				high = mid - 1
		out["records"].append(
			{
				key: unit,
				"offset": char,
				"fragment": text[char : char + low],
				"encoding": "json-row" if is_row else "text",
			}
		)
		out["range"]["last"] = unit
		used = WINDOW_CHARS
		if char + low < len(text):
			out["next_cursor"] = pos(unit, char + low)
			return False
		return True

	def table(rows):
		for number, pair in enumerate(rows, start):
			if number > MAX_ROWS:
				stop("row_limit", source=source, next_row=number)
				return
			if number >= start + WINDOW_ROWS:
				out["next_cursor"] = pos(number)
				return
			row, formulas = pair
			values = [_cell(v) for v in row[:MAX_COLUMNS]]
			finished = add(number, values, True)
			if out["range"]["last"] == number:
				if len(row) > MAX_COLUMNS:
					issue("column_limit", row=number, next_column=MAX_COLUMNS + 1)
				for column, formula in enumerate(formulas[:MAX_COLUMNS], 1):
					if (
						isinstance(formula, str)
						and formula.startswith("=")
						and (column > len(row) or row[column - 1] is None)
					):
						issue("formula_value_missing", row=number, column=column)
			if not finished:
				return

	if ext == "pdf":
		from pypdf import PdfReader

		if source != 0:
			raise InvalidArgumentError("PDF cursor must use source 0.")
		reader = PdfReader(io.BytesIO(content))
		if reader.is_encrypted:
			raise InvalidArgumentError("PDF is encrypted. Supply an unlocked copy to read it.")
		total = len(reader.pages)
		if start > max(total, 1):
			raise InvalidArgumentError("PDF cursor page is outside the file.")
		out.update(
			kind="pdf", source={"filename": filename, "pages_total": total, "scope": "extractable text only"}
		)
		for number in range(start, min(total, MAX_PAGES) + 1):
			if number >= start + WINDOW_PAGES:
				out["next_cursor"] = pos(number)
				break
			failed = False
			try:
				text = reader.pages[number - 1].extract_text() or ""
			except Exception:
				failed = True
				text = ""
			finished = add(number, text)
			if out["range"]["last"] == number:
				if failed:
					issue("page_extraction_failed", page=number)
				if not text.strip():
					issue("no_extractable_text", page=number)
			if not finished:
				break
		if not out["next_cursor"] and total > MAX_PAGES:
			stop("page_limit", next_page=MAX_PAGES + 1)
	elif ext == "xlsx":
		import openpyxl

		if content.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
			raise InvalidArgumentError("Workbook is encrypted or legacy XLS. Supply an unlocked XLSX copy.")
		with ZipFile(io.BytesIO(content)) as archive:
			if sum(z.file_size for z in archive.infolist()) > MAX_EXPANDED_BYTES:
				raise InvalidArgumentError(
					"Expanded workbook exceeds 64 MiB. Split the workbook before reading."
				)
		wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
		formula_wb = None
		try:
			if any(len(name) > 128 for name in wb.sheetnames):
				raise InvalidArgumentError("Workbook contains a worksheet name longer than 128 characters.")
			if sheet is not None and sheet not in wb.sheetnames:
				raise InvalidArgumentError("Requested worksheet does not exist. Check the worksheet name.")
			names = [sheet] if sheet is not None else wb.sheetnames
			if source >= len(names):
				raise InvalidArgumentError("Workbook cursor worksheet is outside the file.")
			name = names[source]
			out.update(kind="table", source={"filename": filename, "sheet": name, "sheets_total": len(names)})
			formula_wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
			ws, fs = wb[name], formula_wb[name]
			ws.reset_dimensions()
			fs.reset_dimensions()
			table(
				zip_longest(
					ws.iter_rows(min_row=start, values_only=True),
					fs.iter_rows(min_row=start, values_only=True),
					fillvalue=(),
				)
			)
			if not out["next_cursor"] and source + 1 < len(names):
				if source + 1 < MAX_SHEETS:
					out["next_cursor"] = pos(1, src=source + 1)
				else:
					stop("sheet_limit", next_sheet=source + 1)
		finally:
			wb.close()
			if formula_wb:
				formula_wb.close()
	else:
		if source != 0:
			raise InvalidArgumentError("Text/CSV cursor must use source 0.")
		try:
			text = content.decode("utf-8-sig")
		except UnicodeDecodeError:
			text = content.decode("utf-8-sig", "replace")
			issue("encoding_replacements")
		if ext in {"csv", "tsv"}:
			out["kind"] = "table"
			rows = csv.reader(io.StringIO(text), delimiter="\t" if ext == "tsv" else ",")
			table((row, ()) for row in islice(rows, start - 1, None))
		else:
			if start != 1:
				raise InvalidArgumentError("Plain text cursor must use unit 1.")
			add(1, text)
	if (
		cursor is not None
		and (start != 1 or offset != 0)
		and not out["records"]
		and not out["stop_reason"]
		and not out["next_cursor"]
	):
		raise InvalidArgumentError("File cursor is outside its source. Restart this file read.")
	return out
