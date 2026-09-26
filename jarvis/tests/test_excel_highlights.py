"""Native conditional formatting, inspected in real workbooks from both engines."""

import datetime
import io
import unittest
import zipfile
from decimal import Decimal
from unittest.mock import patch
from xml.etree import ElementTree as ET

import frappe
from openpyxl import load_workbook

from jarvis import compat
from jarvis.exceptions import InvalidArgumentError
from jarvis.tests.test_compat_frappe_versions import _builders
from jarvis.tools._export import renderers
from jarvis.tools._export.model import ExportModel
from jarvis.tools.export_excel import export_excel
from jarvis.tools.export_query import export_query

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
ROWS = [["Customer", "Balance", "Status"], ["North", 100, "Overdue"], ["West", -20, "Paid"]]
RULE = {"column": "Balance", "operator": "gt", "value": 0}


class TestExcelHighlights(unittest.TestCase):
	def setUp(self):
		import frappe.utils.xlsxutils as xu

		p = patch("frappe.utils.file_manager.save_file")
		self.save = p.start()
		self.addCleanup(p.stop)
		self.save.return_value = frappe._dict(
			file_url="/private/files/highlights.xlsx", file_name="highlights.xlsx", file_size=0, name="MOCK"
		)
		if not hasattr(xu, "get_excel_date_format"):
			p = patch.object(
				xu, "get_excel_date_format", create=True, return_value=("dd-mm-yyyy", "hh:mm:ss")
			)
			p.start()
			self.addCleanup(p.stop)

	def workbook(self):
		return load_workbook(io.BytesIO(self.save.call_args.args[1]))

	def rules(self, ws):
		return sorted(
			[
				(str(key.sqref), rule)
				for key in ws.conditional_formatting
				for rule in ws.conditional_formatting[key]
			],
			key=lambda pair: pair[1].priority,
		)

	def fill_color(self, rule, engine):
		# XlsxWriter emits differential fills as bgColor with no pattern;
		# openpyxl uses a solid fgColor. Both are native conditional fills.
		fill = rule.dxf.fill
		return (fill.bgColor if engine == "xlsxwriter" else fill.fgColor).rgb[-6:]

	def test_both_writers_emit_guarded_native_rules_with_first_match_priority(self):
		for engine, builder in _builders().items():
			with self.subTest(engine=engine), patch("jarvis.tools.export_excel._workbook_bytes", builder):
				out = export_excel(
					rows=ROWS,
					highlights=[
						{**RULE, "style": "danger", "scope": "row"},
						{**RULE, "operator": "lt", "style": "success"},
						{"column": "Status", "operator": "eq", "value": "Paid", "style": "info"},
					],
				)
				self.assertEqual(out["highlight_count"], 3)
				self.assertEqual(self.save.call_args.kwargs["is_private"], 1)
				ws = self.workbook().active
				self.assertEqual(list(ws.values), [tuple(row) for row in ROWS])
				self.assertEqual(ws.freeze_panes, "A2")
				self.assertTrue(ws["A1"].font.bold)
				rules = self.rules(ws)
				self.assertEqual([ref for ref, _ in rules], ["A2:C3", "B2:B3", "C2:C3"])
				self.assertEqual([rule.priority for _, rule in rules], [1, 2, 3])
				self.assertTrue(all(rule.type == "expression" and rule.stopIfTrue for _, rule in rules))
				self.assertEqual(rules[0][1].formula, ['AND(ISNUMBER($B2),$B2<>"",$B2>0)'])
				self.assertEqual(rules[1][1].formula, ['AND(ISNUMBER($B2),$B2<>"",$B2<0)'])
				self.assertEqual(rules[2][1].formula, ['AND(ISTEXT($C2),LEN(TRIM($C2))>0,EXACT($C2,"Paid"))'])
				self.assertEqual(self.fill_color(rules[0][1], engine), "FCE4D6")
				self.assertFalse(ws["B2"].fill.patternType)  # conditional, not a static fill

	def test_competing_cell_and_row_rules_keep_cell_precedence(self):
		for engine, builder in _builders().items():
			with self.subTest(engine=engine), patch("jarvis.tools.export_excel._workbook_bytes", builder):
				export_excel(
					rows=ROWS,
					highlights=[
						{**RULE, "style": "danger"},
						{"column": "Status", "operator": "eq", "value": "Overdue", "scope": "row"},
					],
				)
				ws = self.workbook().active
				self.assertGreater(ws["B2"].value, 0)
				self.assertEqual(ws["C2"].value, "Overdue")
				cell, row = self.rules(ws)
				self.assertEqual((cell[0], row[0]), ("B2:B3", "A2:C3"))
				self.assertEqual((cell[1].priority, row[1].priority), (1, 2))
				self.assertTrue(cell[1].stopIfTrue)
				self.assertEqual(self.fill_color(cell[1], engine), "FCE4D6")
				self.assertEqual(self.fill_color(row[1], engine), "FFF2CC")

	def test_all_numeric_operators_and_defaults(self):
		for engine, builder in _builders().items():
			with self.subTest(engine=engine), patch("jarvis.tools.export_excel._workbook_bytes", builder):
				export_excel(
					rows=ROWS,
					highlights=[
						{**RULE, "operator": op, "value": Decimal("1.25")}
						for op in ("gt", "gte", "lt", "lte", "eq", "ne")
					],
				)
				rules = self.rules(self.workbook().active)
				self.assertEqual(
					[r.formula[0] for _, r in rules],
					[f'AND(ISNUMBER($B2),$B2<>"",$B2{op}1.25)' for op in (">", ">=", "<", "<=", "=", "<>")],
				)
				self.assertTrue(all(ref == "B2:B3" for ref, _ in rules))
				self.assertTrue(all(self.fill_color(r, engine) == "FFF2CC" for _, r in rules))

	def test_literal_text_escaping_case_sensitive_not_equals_and_blanks(self):
		value = '=IF(A1="*?",1,0)'
		rows = [["Status"], [value], ["paid"], [None], [" "], [0], [False]]
		for engine, builder in _builders().items():
			with self.subTest(engine=engine), patch("jarvis.tools.export_excel._workbook_bytes", builder):
				export_excel(rows=rows, highlights=[{"column": "Status", "operator": "ne", "value": value}])
				formula = self.rules(self.workbook().active)[0][1].formula
				self.assertEqual(
					formula, ['AND(ISTEXT($A2),LEN(TRIM($A2))>0,NOT(EXACT($A2,"=IF(A1=""*?"",1,0)")))']
				)

	def test_numeric_guards_leave_mixed_values_unchanged(self):
		rows = [["Balance"], [0], ["0"], [False], [None], [" "], [-1], [1.5]]
		for engine, builder in _builders().items():
			with self.subTest(engine=engine), patch("jarvis.tools.export_excel._workbook_bytes", builder):
				export_excel(rows=rows, highlights=[{**RULE, "operator": "eq"}])
				ws = self.workbook().active
				self.assertEqual(
					[ws[f"A{r}"].value for r in range(2, 9)], [0, "0", False, None, " ", -1, 1.5]
				)
				self.assertEqual(self.rules(ws)[0][0], "A2:A8")
				self.assertIn("ISNUMBER($A2)", self.rules(ws)[0][1].formula[0])

	def test_invalid_rules_and_unknown_headers_never_save(self):
		for highlights in [
			{},
			"rules",
			[None],
			[{}],
			[RULE] * 13,
			[{**RULE, "column": "Missing"}],
			[{**RULE, "column": []}],
			[{**RULE, "operator": "contains"}],
			[{**RULE, "operator": []}],
			*[
				[{**RULE, "value": v}]
				for v in [None, True, [], {}, float("nan"), float("inf"), Decimal("NaN"), 10**1000]
			],
			*[
				[{**RULE, "operator": "eq", "value": v}]
				for v in ["", " ", "x" * 121, "a\nb", "a\x00b", "a\x7fb", "\ud800", "\ufffe"]
			],
			[{**RULE, "value": "100"}],
			[{**RULE, "style": None}],
			[{**RULE, "scope": []}],
			[{**RULE, "formula": "A1>0"}],
		]:
			with self.subTest(highlights=highlights), self.assertRaises(InvalidArgumentError):
				export_excel(rows=ROWS, highlights=highlights)
		with self.assertRaises(InvalidArgumentError):
			export_excel(rows=[["Balance", "Balance"], [1, 2]], highlights=[RULE])
		self.save.assert_not_called()

	def test_dates_are_not_numeric_targets_including_iso_conversion(self):
		for value in [datetime.date(2026, 1, 1), datetime.datetime(2026, 1, 1, 4), "2026-01-01"]:
			with self.subTest(value=value), self.assertRaises(InvalidArgumentError):
				export_excel(rows=[{"Balance": value}], highlights=[RULE])
		self.save.assert_not_called()

	def test_empty_highlighted_sheet_and_oversized_ranges_are_rejected(self):
		from jarvis._xlsx_highlights import validate_highlights

		for rows in [[], [["Balance"]], [{"Balance": None}]]:
			with self.subTest(rows=rows), self.assertRaises(InvalidArgumentError):
				export_excel(rows=rows, highlights=[RULE])
		for data in [[["Balance"]], [["Balance"] + ["x"] * 16384, [1]]]:
			with self.assertRaises(InvalidArgumentError):
				validate_highlights([RULE], data)
		self.save.assert_not_called()

	def test_exported_headers_are_resolved_after_html_and_control_cleanup(self):
		for engine, builder in _builders().items():
			for header in ["<p>Balance</p>", "Balance\x00"]:
				with (
					self.subTest(engine=engine, header=header),
					patch("jarvis.tools.export_excel._workbook_bytes", builder),
				):
					export_excel(rows=[[header], [5]], highlights=[RULE])
					self.assertEqual(self.workbook().active["A1"].value, "Balance")
					self.assertEqual(self.rules(self.workbook().active)[0][0], "A2")
		self.save.reset_mock()
		with self.assertRaises(InvalidArgumentError):
			export_excel(rows=[["Balance", "<p>Balance</p>"], [1, 2]], highlights=[RULE])
		self.save.assert_not_called()

	def test_multi_sheet_alignment_and_late_invalid_rule(self):
		for engine, builder in _builders().items():
			with self.subTest(engine=engine), patch("jarvis.tools.export_excel._workbook_bytes", builder):
				out = export_excel(
					sheets=[
						{"rows": []},
						{"title": "Follow up", "rows": ROWS, "highlights": [RULE]},
						{"title": "Raw", "rows": ROWS},
					]
				)
				self.assertEqual(out["highlight_count"], 1)
				wb = self.workbook()
				self.assertEqual(wb.sheetnames, ["Follow up", "Raw"])
				self.assertEqual([len(ws.conditional_formatting) for ws in wb], [1, 0])
		self.save.reset_mock()
		for tail in [
			{"rows": [], "highlights": [RULE]},
			{"rows": ROWS, "highlights": [{**RULE, "column": "Missing"}]},
		]:
			with self.assertRaises(InvalidArgumentError):
				export_excel(sheets=[{"rows": ROWS}, tail])
		for top in [[], [RULE]]:
			with self.assertRaises(InvalidArgumentError):
				export_excel(sheets=[{"rows": ROWS}], highlights=top)
		self.save.assert_not_called()

	def test_plain_exports_and_empty_rules_create_no_conditional_formatting(self):
		for highlights in [None, []]:
			out = export_excel(rows=ROWS, highlights=highlights)
			self.assertEqual(out["highlight_count"], 0)
			with zipfile.ZipFile(io.BytesIO(self.save.call_args.args[1])) as z:
				root = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
				self.assertEqual(root.findall("s:conditionalFormatting", NS), [])

	def test_query_renderer_preserves_sanitization_and_validates_exported_headers(self):
		for engine, builder in _builders().items():
			with self.subTest(engine=engine), patch.object(compat, "xlsx_bytes", builder):
				model = ExportModel(
					columns=["<p>Status</p>", "Balance"],
					rows=[['<p>=HYPERLINK("x","x")</p>', 5], ["<p>Paid</p>", -1]],
					total=2,
				)
				content = renderers.xlsx(
					model, highlights=[{"column": "Status", "operator": "eq", "value": "Paid"}]
				)
				ws = load_workbook(io.BytesIO(content)).active
				self.assertEqual(ws["A1"].value, "Status")
				self.assertEqual(ws["A2"].value, '\'=HYPERLINK("x","x")')
				self.assertEqual(ws["A3"].value, "Paid")
				self.assertEqual(model.meta["highlight_count"], 1)
				self.assertEqual(self.rules(ws)[0][0], "A2:A3")

	def test_query_forwards_permission_arguments_and_returns_rule_count(self):
		model = ExportModel(columns=ROWS[0], rows=ROWS[1:], total=2)
		with (
			patch("jarvis.tools.export_query.resolvers.from_query", return_value=model) as resolve,
			patch("jarvis.tools.export_query.save_export_file", return_value={}) as save,
			patch("jarvis.telemetry.record_export_event"),
		):
			out = export_query(
				"ToDo",
				filters={"status": "Open"},
				fields=ROWS[0],
				order_by="modified desc",
				parent_doctype="User",
				highlights=[RULE],
			)
		resolve.assert_called_once_with(
			"ToDo",
			filters={"status": "Open"},
			fields=ROWS[0],
			order_by="modified desc",
			parent_doctype="User",
		)
		self.assertEqual(out["total"], 2)
		self.assertEqual(out["highlight_count"], 1)
		self.assertEqual(
			len(load_workbook(io.BytesIO(save.call_args.args[1])).active.conditional_formatting), 1
		)

	def test_query_rejects_csv_highlights_and_invalid_xlsx_before_save(self):
		model = ExportModel(columns=ROWS[0], rows=ROWS[1:], total=2)
		with (
			patch("jarvis.tools.export_query.resolvers.from_query", return_value=model),
			patch("jarvis.tools.export_query.save_export_file") as save,
			patch("jarvis.telemetry.record_export_event"),
		):
			for highlights in [[], [RULE], {}]:
				with self.subTest(highlights=highlights), self.assertRaises(InvalidArgumentError):
					export_query("ToDo", format="csv", highlights=highlights)
			with self.assertRaises(InvalidArgumentError):
				export_query("ToDo", highlights=[{**RULE, "column": "Missing"}])
			save.assert_not_called()

	@unittest.skipIf(
		frappe.__version__.startswith("15."), "develop registry has an unrelated v16-only import"
	)
	def test_registry_dispatch_preserves_highlights(self):
		from jarvis.tools.registry import dispatch

		out = dispatch("export_excel", {"rows": ROWS, "highlights": [RULE]})
		self.assertEqual(out["highlight_count"], 1)
		self.assertEqual(len(self.workbook().active.conditional_formatting), 1)
