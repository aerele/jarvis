"""Tests for export_excel's "no data" guard.

A workbook with no columns, no data rows, or only blank cells is empty in
every sense — export_excel must raise NoDataError (which the agent relays as
"No data to prepare for Excel.") rather than hand the user a blank .xlsx to
open and discover is empty.
"""

from __future__ import annotations

from unittest.mock import patch

import frappe
from frappe.tests.utils import FrappeTestCase

from jarvis.exceptions import InvalidArgumentError, NoDataError
from jarvis.tools.export_excel import export_excel


class TestExportExcelInputIntegrity(FrappeTestCase):
	def setUp(self):
		import frappe.utils.xlsxutils as xu

		if not hasattr(xu, "get_excel_date_format"):
			p = patch.object(
				xu, "get_excel_date_format", create=True, return_value=("dd-mm-yyyy", "hh:mm:ss")
			)
			p.start()
			self.addCleanup(p.stop)

	def test_malformed_rows_are_rejected_before_saving(self):
		for rows in (
			{"Item": "Widget"},
			"Widget",
			[{"Item": "Widget"}, ["Other item"]],
			[{"Item": "Widget"}, None],
			[["Item"], {"Item": "Widget"}],
			[["Item"], "Widget"],
		):
			with self.subTest(rows=rows), patch("frappe.utils.file_manager.save_file") as save:
				with self.assertRaises(InvalidArgumentError) as error:
					export_excel(rows=rows)
				self.assertNotIsInstance(error.exception, NoDataError)
				save.assert_not_called()

	def test_malformed_later_sheet_cannot_be_silently_skipped(self):
		with patch("frappe.utils.file_manager.save_file") as save:
			with self.assertRaises(InvalidArgumentError):
				export_excel(
					sheets=[
						{"title": "Valid", "rows": [{"Amount": 10}]},
						{"title": "Invalid", "rows": {"Amount": 20}},
					]
				)
			save.assert_not_called()

	def test_explicit_columns_do_not_allow_mixed_row_shapes(self):
		with patch("frappe.utils.file_manager.save_file") as save:
			with self.assertRaises(InvalidArgumentError):
				export_excel(columns=["Item"], rows=[["Widget"], {"Item": "Other"}])
			save.assert_not_called()

	def test_sheet_names_and_chart_references_match_on_both_builders(self):
		from jarvis.tests.test_compat_frappe_versions import _builders

		rows = [{"Month": "Jan", "Sales": 12}]
		chart = {"type": "line", "categories": "Month", "values": ["Sales"]}
		for engine, builder in _builders().items():
			with (
				self.subTest(engine=engine),
				patch("jarvis.tools.export_excel._workbook_bytes", side_effect=builder),
				patch("frappe.utils.file_manager.save_file") as save,
			):
				export_excel(
					sheets=[
						{"title": title, "rows": rows, "charts": [chart]}
						for title in ("North/South", "north:south", "'Quarter'", "///", "x" * 30 + "'tail")
					]
				)
				wb = _load(_saved_bytes(save))
				self.assertEqual(
					wb.sheetnames, ["North South", "north south-2", "Quarter", "Sheet", "x" * 30]
				)
				for ws in wb:
					self.assertEqual(ws["B2"].value, 12)
					self.assertEqual(
						ws._charts[0].series[0].val.numRef.f.replace("'", ""), f"{ws.title}!$B$2"
					)

	def test_single_sheet_title_is_sanitized(self):
		with patch("frappe.utils.file_manager.save_file") as save:
			export_excel(title="'Revenue/Costs'", rows=[{"Amount": 12}])
			self.assertEqual(_load(_saved_bytes(save)).sheetnames, ["Revenue Costs"])


class TestExportExcelNoData(FrappeTestCase):
	def test_rejects_empty_list(self):
		with self.assertRaises(NoDataError):
			export_excel([])

	def test_rejects_list_of_empty_dicts(self):
		# No keys → no columns → a contentless workbook. Must be rejected.
		with self.assertRaises(NoDataError):
			export_excel([{}])
		with self.assertRaises(NoDataError):
			export_excel([{}, {}])

	def test_rejects_all_blank_dict_rows(self):
		# Columns exist but every value is empty → nothing to export.
		with self.assertRaises(NoDataError):
			export_excel([{"a": None, "b": None}])
		with self.assertRaises(NoDataError):
			export_excel([{"a": "", "b": "   "}])

	def test_rejects_header_only_list(self):
		with self.assertRaises(NoDataError):
			export_excel([["Name", "Qty"]])

	def test_rejects_all_blank_list_rows(self):
		with self.assertRaises(NoDataError):
			export_excel([["Name", "Qty"], ["", None]])

	def test_generates_when_there_is_real_data(self):
		# Regression guard: a row with at least one non-empty cell exports.
		with patch("frappe.utils.file_manager.save_file") as m:
			m.return_value = frappe._dict(
				file_url="/private/files/ok.xlsx", file_name="ok.xlsx", file_size=42, name="F-OK"
			)
			out = export_excel([{"a": 1, "b": None}], title="ok")
		self.assertEqual(out["file_url"], "/private/files/ok.xlsx")

	def test_generates_list_of_lists_with_data(self):
		with patch("frappe.utils.file_manager.save_file") as m:
			m.return_value = frappe._dict(
				file_url="/private/files/ll.xlsx", file_name="ll.xlsx", file_size=42, name="F-LL"
			)
			out = export_excel([["Name", "Qty"], ["Widget", 5]])
		self.assertEqual(out["file_url"], "/private/files/ll.xlsx")


def _saved_bytes(mock):
	"""The content bytes export_excel handed save_file (2nd positional arg)."""
	return mock.call_args.args[1]


def _load(content):
	from io import BytesIO

	from openpyxl import load_workbook

	return load_workbook(BytesIO(content))


class TestExportExcelMultiSheet(FrappeTestCase):
	def _mock_save(self):
		m = patch("frappe.utils.file_manager.save_file").start()
		self.addCleanup(patch.stopall)
		m.return_value = frappe._dict(
			file_url="/private/files/wb.xlsx", file_name="wb.xlsx", file_size=99, name="F-WB"
		)
		return m

	def test_builds_named_tabs_in_order(self):
		m = self._mock_save()
		out = export_excel(
			title="MIS",
			sheets=[
				{"title": "Financials", "rows": [{"metric": "Revenue", "value": 100}]},
				{"title": "Customers", "rows": [["Name", "Due"], ["Acme", 50]]},
			],
		)
		self.assertEqual(out["file_url"], "/private/files/wb.xlsx")
		wb = _load(_saved_bytes(m))
		self.assertEqual(wb.sheetnames, ["Financials", "Customers"])
		self.assertEqual(wb["Financials"]["A1"].value, "metric")  # header from dict keys
		self.assertEqual(wb["Financials"]["A2"].value, "Revenue")
		self.assertEqual(wb["Customers"]["B2"].value, 50)  # list-of-lists body

	def test_empty_tabs_are_skipped_but_the_rest_ship(self):
		m = self._mock_save()
		export_excel(
			title="w",
			sheets=[
				{"title": "HR", "rows": []},  # empty → skipped
				{"title": "Ops", "rows": [{"a": 1}]},  # kept
				{"title": "Blank", "rows": [{"a": None}]},  # all-blank → skipped
			],
		)
		self.assertEqual(_load(_saved_bytes(m)).sheetnames, ["Ops"])

	def test_all_empty_raises_no_data(self):
		with self.assertRaises(NoDataError):
			export_excel(sheets=[{"title": "A", "rows": []}, {"title": "B", "rows": [{}]}])

	def test_empty_sheets_list_is_invalid(self):
		with self.assertRaises(InvalidArgumentError):
			export_excel(sheets=[])

	def test_duplicate_titles_get_unique_sheet_names(self):
		m = self._mock_save()
		export_excel(
			sheets=[
				{"title": "Summary", "rows": [{"a": 1}]},
				{"title": "Summary", "rows": [{"a": 2}]},
			],
		)
		self.assertEqual(_load(_saved_bytes(m)).sheetnames, ["Summary", "Summary-2"])


class TestExportExcelIsoDates(FrappeTestCase):
	"""Issue #598: dates reach the tool as ISO text (rows go through JSON) and
	used to land in Excel as text, which cannot sort or filter as dates. Only
	strict ISO columns convert, all or nothing, so codes and names never do."""

	def _body(self, rows):
		from jarvis.tools.export_excel import _normalize

		return _normalize(rows, None)[1:]

	def test_iso_date_and_datetime_columns_become_real_dates(self):
		import datetime

		body = self._body(
			[
				{"d": "2026-06-01", "t": "2026-06-01 09:30:00", "blank": "2026-06-02"},
				{"d": "2026-10-11", "t": "2026-10-11T18:05", "blank": ""},
			]
		)
		self.assertEqual(body[0][0], datetime.date(2026, 6, 1))
		self.assertEqual(body[1][1], datetime.datetime(2026, 10, 11, 18, 5))
		self.assertEqual(body[0][2], datetime.date(2026, 6, 2))
		self.assertEqual(body[1][2], "")  # blanks stay blank, the column still converts

	def test_look_alikes_and_mixed_columns_stay_text(self):
		body = self._body(
			[
				{"name": "2026-00012", "code": "00123", "bad": "2026-13-40", "mixed": "2026-06-01"},
				{"name": "2026-00013", "code": "04", "bad": "2026-02-01", "mixed": "soon"},
			]
		)
		self.assertEqual(body[0], ["2026-00012", "00123", "2026-13-40", "2026-06-01"])
		self.assertEqual(body[1], ["2026-00013", "04", "2026-02-01", "soon"])
